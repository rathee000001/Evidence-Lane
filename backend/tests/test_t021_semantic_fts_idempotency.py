from __future__ import annotations

import sqlite3
from pathlib import Path

from sqlite_brain_builder.runtime import stable_runtime_v53 as runtime


def _counts(db: Path, tables: list[str]) -> dict[str, int]:
    with sqlite3.connect(db) as connection:
        return {
            table: connection.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0]
            for table in tables
        }


def test_chat_lineage_twelve_identical_ingestions_have_zero_row_growth(tmp_path: Path) -> None:
    source = tmp_path / "chat.txt"
    source.write_text(
        "User: Build the project.\nAssistant: Registered.\n"
        "User: Mandatory correction delta.\nAssistant: Hard gate recorded.\n",
        encoding="utf-8",
    )
    database = tmp_path / "chat_lineage.sqlite"
    payload = {
        "source_id": "source-chat-fixed",
        "lane_key": "chat_lineage",
        "display_name": source.name,
        "path": str(source),
        "source_type": "Chat Lineage",
    }

    runtime.build_semantic_lane(database, "chat_lineage", payload, None, 0)
    baseline = _counts(database, ["lineage_source", "lineage_turn", "turn_fts"])
    repeated = []
    for _ in range(11):
        runtime.build_semantic_lane(database, "chat_lineage", payload, None, 0)
        repeated.append(_counts(database, ["lineage_source", "lineage_turn", "turn_fts"]))

    assert repeated == [baseline] * 11


def test_ppt_twelve_identical_ingestions_have_zero_fts_growth(tmp_path: Path) -> None:
    import pptx

    source = tmp_path / "deck.pptx"
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Evidence OS"
    slide.placeholders[1].text = "Deterministic presentation indexing"
    presentation.save(source)

    database = tmp_path / "ppt.sqlite"
    payload = {
        "source_id": "source-ppt-fixed",
        "lane_key": "ppt_presentation",
        "display_name": source.name,
        "path": str(source),
        "source_type": "PPT / Presentation",
    }

    runtime.build_semantic_lane(database, "ppt_presentation", payload, None, 0)
    baseline = _counts(database, ["ppt_file", "ppt_slide", "ppt_text_block", "ppt_fts"])
    repeated = []
    for _ in range(11):
        runtime.build_semantic_lane(database, "ppt_presentation", payload, None, 0)
        repeated.append(_counts(database, ["ppt_file", "ppt_slide", "ppt_text_block", "ppt_fts"]))

    assert repeated == [baseline] * 11
