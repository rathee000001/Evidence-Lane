from __future__ import annotations

import json
import sqlite3
from pathlib import Path

from PIL import Image, ImageDraw

from sqlite_brain_builder.runtime import ocr_ingest
from sqlite_brain_builder.runtime.canonical_lanes import get_lane
from sqlite_brain_builder.runtime.env15_project_schema import resolve_env15_sector
from sqlite_brain_builder.runtime.stable_runtime_v53 import build_brain


def _fixtures(tmp_path: Path) -> dict[str, Path]:
    from docx import Document
    import fitz
    import openpyxl
    from pptx import Presentation

    docx_path = tmp_path / "evidence.docx"
    document = Document()
    document.add_heading("Evidence", level=1)
    document.add_paragraph("Structure-linked paragraph.")
    document.save(docx_path)

    xlsx_path = tmp_path / "evidence.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(["Value", "Double"])
    worksheet.append([5, "=A2*2"])
    workbook.save(xlsx_path)
    workbook.close()

    csv_path = tmp_path / "evidence.csv"
    csv_path.write_text("name,status\nbackend,pass\n", encoding="utf-8")
    json_path = tmp_path / "evidence.json"
    json_path.write_text(json.dumps({"lane": "data", "status": "pass"}), encoding="utf-8")

    pptx_path = tmp_path / "evidence.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Evidence OS"
    presentation.save(pptx_path)

    image_path = tmp_path / "evidence.png"
    image = Image.new("RGB", (240, 100), "white")
    ImageDraw.Draw(image).text((20, 35), "OCR evidence", fill="black")
    image.save(image_path)

    pdf_path = tmp_path / "evidence.pdf"
    pdf = fitz.open()
    page = pdf.new_page(width=320, height=180)
    page.insert_image(fitz.Rect(20, 30, 300, 150), filename=str(image_path))
    pdf.save(pdf_path)
    pdf.close()
    return {
        "docs": docx_path,
        "xlsx": xlsx_path,
        "csv": csv_path,
        "json": json_path,
        "ppt": pptx_path,
        "pdf_ocr": pdf_path,
        "images_ocr": image_path,
    }


def _fake_ocr(_image):
    item = {"id": "stable-ocr", "order": 1, "text": "OCR evidence", "confidence": 93.0, "bbox": [1, 2, 3, 4]}
    return {
        "status": "OCR_COMPLETE", "error": "", "blocks": [item],
        "lines": [{**item, "id": "stable-line"}], "regions": [],
        "confidence": 93.0, "review_required": False,
    }


def test_live_env15_section_c_structures_and_zero_growth(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setattr(ocr_ingest, "_ocr_pil_image", _fake_ocr)
    fixtures = _fixtures(tmp_path)
    sources = [
        {"source_id": "doc", "lane_key": "docs", "path": str(fixtures["docs"])},
        {"source_id": "xlsx", "lane_key": "data_excel", "path": str(fixtures["xlsx"])},
        {"source_id": "csv", "lane_key": "data_excel", "path": str(fixtures["csv"])},
        {"source_id": "json", "lane_key": "data_excel", "path": str(fixtures["json"])},
        {"source_id": "ppt", "lane_key": "ppt", "path": str(fixtures["ppt"])},
        {"source_id": "pdf", "lane_key": "pdf_ocr", "path": str(fixtures["pdf_ocr"])},
        {"source_id": "image", "lane_key": "images_ocr", "path": str(fixtures["images_ocr"])},
    ]
    first = build_brain(str(tmp_path), "Section C", sources, generate_mmd=False)
    root = Path(first["brain_root"])
    before = {}
    expected_types = {
        "docs": {"docs:doc_structure"},
        "data_excel": {"data_excel:workbook_metadata", "data_excel:delimited_header", "data_excel:data_structure"},
        "ppt": {"ppt:presentation_structure"},
        "pdf_ocr": {"pdf_ocr:pdf_page_structure"},
        "images_ocr": {"images_ocr:image_ocr_structure"},
    }
    for lane, expected in expected_types.items():
        _, database = resolve_env15_sector(root, lane)
        with sqlite3.connect(database) as connection:
            chunk_types = {row[0] for row in connection.execute("SELECT DISTINCT chunk_type FROM chunk_index")}
            assert expected <= chunk_types
            fts_table = get_lane(lane).fts_table
            before[lane] = (
                connection.execute("SELECT COUNT(*) FROM source_registry").fetchone()[0],
                connection.execute("SELECT COUNT(*) FROM chunk_index").fetchone()[0],
                connection.execute(f'SELECT COUNT(*) FROM "{fts_table}"').fetchone()[0],
            )

    second = build_brain(str(tmp_path), "Section C", sources, generate_mmd=False)
    assert second["incremental"]["sources_reused"] == len(sources)
    for lane, counts in before.items():
        _, database = resolve_env15_sector(root, lane)
        with sqlite3.connect(database) as connection:
            fts_table = get_lane(lane).fts_table
            assert (
                connection.execute("SELECT COUNT(*) FROM source_registry").fetchone()[0],
                connection.execute("SELECT COUNT(*) FROM chunk_index").fetchone()[0],
                connection.execute(f'SELECT COUNT(*) FROM "{fts_table}"').fetchone()[0],
            ) == counts
