from __future__ import annotations

import sqlite3
from pathlib import Path

from sqlite_brain_builder.runtime.office_ingest import parse_docx, parse_pptx, parse_workbook
from sqlite_brain_builder.runtime.stable_runtime_v53 import build_semantic_lane


def _table_counts(database: Path, tables: list[str]) -> dict[str, int]:
    with sqlite3.connect(database) as connection:
        return {table: connection.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0] for table in tables}


def _docx_fixture(path: Path) -> None:
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches
    from PIL import Image

    document = Document()
    document.core_properties.title = "Evidence document"
    document.add_heading("Backend correction", level=1)
    document.add_paragraph("First structured paragraph.")
    link_paragraph = document.add_paragraph("Reference: ")
    relationship_id = link_paragraph.part.relate_to("https://example.invalid/evidence", RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Evidence link"
    run.append(text)
    hyperlink.append(run)
    link_paragraph._p.append(hyperlink)
    document.add_paragraph("A list item", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header A"
    table.cell(0, 1).text = "Header B"
    table.cell(1, 0).text = "Merged value"
    table.cell(1, 0).merge(table.cell(1, 1))
    document.sections[0].header.paragraphs[0].text = "Evidence header"
    document.sections[0].footer.paragraphs[0].text = "Evidence footer"
    image_path = path.with_suffix(".png")
    Image.new("RGB", (20, 20), "blue").save(image_path)
    document.add_picture(str(image_path), width=Inches(0.25))
    document.save(path)


def _xlsx_fixture(path: Path) -> None:
    import openpyxl
    from openpyxl.chart import BarChart, Reference
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.workbook.defined_name import DefinedName

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet.append(["Item", "Value", "Double", "Merged", None])
    worksheet.append(["A", 10, "=B2*2", "Project", None])
    worksheet.append(["B", 20, "=SUM(B2:B3)", None, None])
    table = Table(displayName="EvidenceTable", ref="A1:C3")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    worksheet.add_table(table)
    worksheet.merge_cells("D1:E1")
    worksheet.row_dimensions[3].hidden = True
    worksheet.column_dimensions["E"].hidden = True
    validation = DataValidation(type="list", formula1='"Open,Closed"')
    validation.add("F2:F3")
    worksheet.add_data_validation(validation)
    worksheet.conditional_formatting.add("B2:B3", CellIsRule(operator="greaterThan", formula=["0"]))
    chart = BarChart()
    chart.add_data(Reference(worksheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(worksheet, min_col=1, min_row=2, max_row=3))
    worksheet.add_chart(chart, "H2")
    workbook.defined_names.add(DefinedName("EvidenceRange", attr_text="'Data'!$A$2:$A$3"))
    workbook.save(path)
    workbook.close()


def _pptx_fixture(path: Path) -> None:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    from PIL import Image

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Evidence OS"
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1.5), Inches(2), Inches(1))
    shape.text = "Canonical backend"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1.5)).table
    table.cell(0, 0).text = "Gate"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Structure"
    table.cell(1, 1).text = "PASS"
    slide.notes_slide.notes_text_frame.text = "Presenter notes preserved."
    image_path = path.with_suffix(".png")
    Image.new("RGB", (24, 24), "cyan").save(image_path)
    slide.shapes.add_picture(str(image_path), Inches(5.5), Inches(1), Inches(0.5), Inches(0.5))
    presentation.save(path)


def test_docx_captures_headings_styles_tables_merges_headers_and_chunks(tmp_path: Path) -> None:
    source = tmp_path / "fixture.docx"
    _docx_fixture(source)

    first = parse_docx(source)
    second = parse_docx(source)

    assert first == second
    assert first["metadata"]["title"] == "Evidence document"
    assert first["headings"][0]["level"] == 1
    assert any(item["is_list"] for item in first["paragraphs"])
    assert first["tables"][0]["row_count"] == 2
    assert first["tables"][0]["merged_cell_groups"]
    assert first["hyperlinks"][0]["target"] == "https://example.invalid/evidence"
    assert first["embedded_images"][0]["byte_size"] > 0
    assert {item["kind"] for item in first["headers_footers"]} == {"header", "footer"}
    assert first["chunks"]


def test_xlsx_captures_tables_charts_formulas_dependencies_ranges_and_hidden_state(tmp_path: Path) -> None:
    source = tmp_path / "fixture.xlsx"
    _xlsx_fixture(source)

    parsed = parse_workbook(source)

    assert parsed["sheets"][0]["used_range"]
    assert parsed["tables"][0]["name"] == "EvidenceTable"
    assert parsed["charts"][0]["source_ranges"]
    assert len(parsed["formulas"]) == 2
    assert parsed["formula_dependencies"]
    assert parsed["named_ranges"][0]["name"] == "EvidenceRange"
    assert parsed["merged_cells"][0]["range"] == "D1:E1"
    assert {item["kind"] for item in parsed["hidden_dimensions"]} == {"row", "column"}
    assert parsed["validations"][0]["range"] == "F2:F3"
    assert parsed["chunks"]


def test_pptx_captures_shapes_tables_notes_relationships_and_chunks(tmp_path: Path) -> None:
    source = tmp_path / "fixture.pptx"
    _pptx_fixture(source)

    parsed = parse_pptx(source)
    slide = parsed["slides"][0]

    assert len(slide["shapes"]) >= 3
    assert slide["tables"][0]["rows"][1] == ["Structure", "PASS"]
    assert "Presenter notes preserved" in slide["notes"]
    assert slide["relationships"]
    assert slide["images"][0]["byte_size"] > 0
    assert parsed["chunks"]


def test_data_dispatch_uses_actual_type_and_reports_legacy_or_optional_support(tmp_path: Path) -> None:
    from sqlite_brain_builder.runtime.office_ingest import parse_data_source

    tsv = tmp_path / "fixture.tsv"
    tsv.write_text("name\tstatus\nbackend\tpass\n", encoding="utf-8")
    json_path = tmp_path / "fixture.json"
    json_path.write_text('{"status":"pass"}', encoding="utf-8")
    jsonl = tmp_path / "fixture.jsonl"
    jsonl.write_text('{"row":1}\n{"row":2}\n', encoding="utf-8")
    legacy = tmp_path / "fixture.xls"
    legacy.write_bytes(b"legacy-xls-placeholder")

    assert parse_data_source(tsv)["source_kind"] == "delimited"
    assert parse_data_source(tsv)["payload"]["delimiter"] == "\t"
    assert parse_data_source(json_path)["source_kind"] == "json"
    assert parse_data_source(jsonl)["payload"] == [{"row": 1}, {"row": 2}]
    assert parse_data_source(legacy)["status"] == "UNSUPPORTED_CONVERTER_REQUIRED"

    parquet_path = tmp_path / "fixture.parquet"
    try:
        import pyarrow as pa
        import pyarrow.parquet as parquet
    except ImportError:
        parquet_path.write_bytes(b"PAR1")
        assert parse_data_source(parquet_path)["status"] == "UNSUPPORTED_PYARROW_NOT_INSTALLED"
    else:
        parquet.write_table(pa.table({"name": ["backend"], "status": ["pass"]}), parquet_path)
        parsed = parse_data_source(parquet_path)
        assert parsed["status"] == "SUPPORTED"
        assert parsed["payload"]["columns"][0]["name"] == "name"
        assert parsed["payload"]["chunks"][0]["rows"][0]["status"] == "pass"


def test_office_second_ingestion_adds_zero_rows_and_csv_is_not_a_workbook(tmp_path: Path) -> None:
    document = tmp_path / "fixture.docx"
    workbook = tmp_path / "fixture.xlsx"
    presentation = tmp_path / "fixture.pptx"
    csv_source = tmp_path / "fixture.csv"
    _docx_fixture(document)
    _xlsx_fixture(workbook)
    _pptx_fixture(presentation)
    csv_source.write_text("name,status\ncanonical,pass\n", encoding="utf-8")

    fixtures = [
        ("docs", document, ["doc_heading", "doc_paragraph", "doc_table_extract", "doc_chunk", "doc_fts"]),
        ("data_excel", workbook, ["sheet_workbook", "sheet_table", "sheet_formula", "sheet_formula_dependency_edge", "sheet_chart_metadata", "data_chunk", "data_fts"]),
        ("ppt", presentation, ["ppt_slide", "ppt_shape", "ppt_table", "ppt_notes", "ppt_slide_relationship", "ppt_chunk", "ppt_fts"]),
    ]
    for lane, source, tables in fixtures:
        database = tmp_path / f"{lane}.sqlite"
        payload = {"source_id": f"source-{lane}", "lane_key": lane, "path": str(source), "display_name": source.name}
        build_semantic_lane(database, lane, payload, None, 0)
        first = _table_counts(database, tables)
        build_semantic_lane(database, lane, payload, None, 0)
        assert _table_counts(database, tables) == first
        assert all(count > 0 for count in first.values())

    csv_database = tmp_path / "csv.sqlite"
    build_semantic_lane(
        csv_database,
        "data_excel",
        {"source_id": "source-csv", "lane_key": "data_excel", "path": str(csv_source), "display_name": csv_source.name},
        None,
        0,
    )
    counts = _table_counts(csv_database, ["data_source", "sheet_workbook", "csv_header", "csv_row_sample"])
    assert counts == {"data_source": 1, "sheet_workbook": 0, "csv_header": 1, "csv_row_sample": 1}
