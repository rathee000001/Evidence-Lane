from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import sqlite3
import tempfile
import time
import zipfile
from pathlib import Path

from sqlite_brain_builder.runtime import stable_runtime_v53 as base

APP_VERSION = "V5.6_ROLLBACK_V53_DELTA_REPAIR"
LANE_DEFS = base.LANE_DEFS
TAB_ORDER = base.TAB_ORDER
normalize_workspace_dir = base.normalize_workspace_dir
brain_output_dir = base.brain_output_dir
scan_tools = base.scan_tools
install_missing_dependencies = base.install_missing_dependencies
FTS_SHADOW_SUFFIXES = ("_data", "_idx", "_content", "_docsize", "_config")


def _emit(progress, stage, task, file="", percent=0, done=0, total=0):
    if progress:
        progress({"stage": stage, "task": task, "file": str(file), "percent": int(percent), "done": int(done), "total": int(total), "eta_seconds": "--", "finish_epoch": "--"})


def _hash_text(s: str) -> str:
    return hashlib.sha256((s or "").encode("utf-8", "ignore")).hexdigest()


def _table_exists(con, table: str) -> bool:
    try:
        return con.execute("SELECT COUNT(*) FROM sqlite_master WHERE type IN ('table','view') AND name=?", (table,)).fetchone()[0] > 0
    except Exception:
        return False


def _row_count(con, table: str) -> int:
    if not _table_exists(con, table):
        return 0
    try:
        return int(con.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0])
    except Exception:
        return 0


def _cols(con, table: str) -> list[str]:
    try:
        return [r[1] for r in con.execute(f'PRAGMA table_info("{table}")')]
    except Exception:
        return []


def _ensure_generic(con, table: str):
    if table.endswith("_fts"):
        con.execute(f'CREATE VIRTUAL TABLE IF NOT EXISTS "{table}" USING fts5(entity_id, text)')
    else:
        con.execute(f'CREATE TABLE IF NOT EXISTS "{table}"(id TEXT PRIMARY KEY, source_id TEXT, name TEXT, path TEXT, value TEXT, metadata_json TEXT, created_at TEXT)')


def _ginsert(con, table: str, source_id: str, name: str, path: str, value: str = "", metadata=None):
    _ensure_generic(con, table)
    metadata_json = json.dumps(metadata or {}, ensure_ascii=False, default=str)
    rid = table + "_" + _hash_text(str(source_id) + str(name) + str(path) + str(value)[:500] + metadata_json)[:18]
    c = set(_cols(con, table))
    if {"id", "source_id", "name", "path", "value", "metadata_json", "created_at"}.issubset(c):
        con.execute(f'INSERT OR REPLACE INTO "{table}"(id,source_id,name,path,value,metadata_json,created_at) VALUES(?,?,?,?,?,?,?)', (rid, source_id, str(name), str(path), str(value), metadata_json, base.now()))
    elif {"entity_id", "text"}.issubset(c):
        con.execute(f'INSERT INTO "{table}"(entity_id,text) VALUES(?,?)', (rid, str(value)))
    return rid


def _fts(con, table: str, entity_id: str, text: str):
    if not text:
        return
    _ensure_generic(con, table)
    try:
        con.execute(f'INSERT INTO "{table}"(entity_id,text) VALUES(?,?)', (entity_id, text))
    except Exception:
        pass


def _insert_dependency_manifest(con, manifest_id, file_id, manifest_type, ecosystem, path, sha256):
    if not _table_exists(con, "dependency_manifest"):
        return
    c = set(_cols(con, "dependency_manifest"))
    if "manifest_id" in c:
        con.execute('INSERT OR REPLACE INTO dependency_manifest(manifest_id,file_id,manifest_type,ecosystem,path,sha256) VALUES(?,?,?,?,?,?)', (manifest_id, file_id, manifest_type, ecosystem, path, sha256))
    else:
        _ginsert(con, "dependency_manifest", file_id, manifest_type, path, sha256, {"ecosystem": ecosystem})


def _insert_dependency_item(con, dependency_id, manifest_id, package_name, version_spec, ecosystem):
    if not _table_exists(con, "dependency_item"):
        return
    c = set(_cols(con, "dependency_item"))
    if "dependency_id" in c:
        con.execute('INSERT OR REPLACE INTO dependency_item(dependency_id,manifest_id,package_name,version_spec,ecosystem) VALUES(?,?,?,?,?)', (dependency_id, manifest_id, package_name, version_spec, ecosystem))
    else:
        _ginsert(con, "dependency_item", manifest_id, package_name, ecosystem, version_spec, {"ecosystem": ecosystem})


def _parse_dependencies_from_file(path: Path):
    out = []
    try:
        if path.name == "package.json":
            data = json.loads(path.read_text(encoding="utf-8", errors="replace"))
            for section in ["dependencies", "devDependencies", "peerDependencies", "optionalDependencies"]:
                obj = data.get(section) or {}
                if isinstance(obj, dict):
                    for name, version in obj.items():
                        out.append((name, str(version), "node", section))
        elif path.name in {"requirements.txt", "requirements-dev.txt"}:
            for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
                s = line.strip()
                if not s or s.startswith("#"):
                    continue
                pkg = re.split(r"[<>=!~;\s]", s, 1)[0]
                if pkg:
                    out.append((pkg, s[len(pkg):].strip() or "*", "python", "runtime"))
        elif path.name == "pyproject.toml":
            text = path.read_text(encoding="utf-8", errors="replace")
            for m in re.finditer(r'^[ \t]*["\']?([A-Za-z0-9_.-]+)["\']?[ \t]*=[ \t]*["\']([^"\']+)', text, re.M):
                out.append((m.group(1), m.group(2), "python", "pyproject"))
        elif path.name in {"environment.yml", "environment.yaml"}:
            for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
                s = line.strip()
                if s.startswith("-") and not s.startswith("- pip"):
                    name = s[1:].strip().split("=", 1)[0].strip()
                    if name:
                        out.append((name, s[1:].strip(), "conda", "runtime"))
    except Exception:
        pass
    return out


def _repair_dependencies_and_workflow(db: Path, source: dict, project_root: Path, progress=None):
    folder = Path(source.get("path") or "")
    if not folder.exists():
        return
    con = sqlite3.connect(str(db), timeout=60)
    try:
        if _table_exists(con, "code_file") and _row_count(con, "code_file") == 0:
            con.close()
            base.build_code_sector(db, source, project_root, progress, source.get("lane_key") or "local_code", time.time())
            con = sqlite3.connect(str(db), timeout=60)

        dep_files = []
        for name in ["package.json", "requirements.txt", "requirements-dev.txt", "pyproject.toml", "environment.yml", "environment.yaml"]:
            dep_files.extend(folder.rglob(name))
        seen = set()
        for f in dep_files:
            rel = f.relative_to(folder).as_posix()
            digest = base.sha256_file(f)
            fid = "depfile_" + digest[:16]
            mid = "manifest_" + digest[:16]
            ecosystem = "node" if f.name == "package.json" else ("python" if f.name in {"requirements.txt", "requirements-dev.txt", "pyproject.toml"} else "conda")
            _insert_dependency_manifest(con, mid, fid, f.name, ecosystem, rel, digest)
            for name, version, eco, section in _parse_dependencies_from_file(f):
                key = (name, version, eco, mid)
                if key in seen:
                    continue
                seen.add(key)
                did = "dep_" + _hash_text("|".join(key))[:18]
                _insert_dependency_item(con, did, mid, name, version, eco)
                _ginsert(con, "code_dependency_edge", mid, name, rel, f"{f.name} -> {name}", {"dependency_id": did, "section": section})

        if _table_exists(con, "workflow_node"):
            node_specs = [("routes", "Routes / pages", _row_count(con, "app_route")), ("code", "Code files / chunks / symbols", _row_count(con, "code_file")), ("dependencies", "Tools + dependencies", _row_count(con, "dependency_item")), ("artifacts", "Read-only artifacts", _row_count(con, "project_artifact"))]
            for nid, name, count in node_specs:
                _ginsert(con, "workflow_node", "code_workflow", nid, "local_code", name, {"count": count})
            for frm, to, rel in [("routes", "code", "route_uses_code"), ("code", "dependencies", "code_uses_dependency"), ("code", "artifacts", "code_preserves_artifact")]:
                _ginsert(con, "workflow_edge", "code_workflow", f"{frm}->{to}", "local_code", rel, {"from": frm, "to": to})

        if _table_exists(con, "artifact_relation_edge") and _row_count(con, "artifact_relation_edge") == 0 and _table_exists(con, "project_artifact"):
            c = set(_cols(con, "artifact_relation_edge"))
            for row in con.execute('SELECT artifact_id FROM project_artifact LIMIT 500').fetchall():
                aid = row[0]
                eid = "artifact_edge_" + _hash_text(str(aid))[:18]
                if "edge_id" in c:
                    con.execute('INSERT OR REPLACE INTO artifact_relation_edge(edge_id,artifact_id,related_entity_type,related_entity_id,relation_type,confidence) VALUES(?,?,?,?,?,?)', (eid, aid, "CODE_PROJECT", "local_code", "preserved_with_code_brain", "deterministic"))
                else:
                    _ginsert(con, "artifact_relation_edge", "local_code", str(aid), "local_code", "preserved_with_code_brain")
        con.commit()
    finally:
        try:
            con.commit(); con.execute("VACUUM")
        except Exception:
            pass
        con.close()


def _repair_xlsx(db: Path, source: dict):
    path = Path(source.get("path") or "")
    if not path.exists() or path.suffix.lower() not in {".xlsx", ".xls"}:
        return
    sid = source.get("source_id") or "xlsx_" + _hash_text(str(path))[:12]
    con = sqlite3.connect(str(db), timeout=60)
    try:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(path), data_only=False, read_only=False)
            digest = base.sha256_file(path)
            _ginsert(con, "sheet_workbook", sid, path.name, str(path), digest, {"sheet_count": len(wb.worksheets)})
            for ws in wb.worksheets:
                maxr = ws.max_row or 0; maxc = ws.max_column or 0
                rng = f"{ws.title}!A1:{ws.cell(maxr, maxc).coordinate if maxr and maxc else 'A1'}"
                _ginsert(con, "sheet_tab", sid, ws.title, str(path), f"rows={maxr} cols={maxc}")
                _ginsert(con, "sheet_range", sid, ws.title, str(path), rng, {"rows": maxr, "cols": maxc})
                _ginsert(con, "sheet_table", sid, ws.title, str(path), rng, {"table_type": "used_range"})
                formula_cells = []
                for row in ws.iter_rows(min_row=1, max_row=min(maxr, 5000), max_col=min(maxc, 200), values_only=False):
                    vals = []
                    for cell in row:
                        vals.append(cell.value)
                        if isinstance(cell.value, str) and cell.value.startswith("="):
                            formula_cells.append((cell.coordinate, cell.value))
                    if row and row[0].row <= 500:
                        _ginsert(con, "sheet_cell_sample", sid, f"{ws.title}!row_{row[0].row}", str(path), json.dumps([str(v) if v is not None else None for v in vals], ensure_ascii=False))
                for coord, formula in formula_cells[:5000]:
                    _ginsert(con, "sheet_formula", sid, f"{ws.title}!{coord}", str(path), formula)
                    for ref in re.findall(r"\b([A-Z]{1,3}\d{1,7})\b", formula):
                        _ginsert(con, "sheet_formula_dependency_edge", sid, f"{ws.title}!{coord}->{ref}", str(path), formula, {"from_cell": coord, "to_cell": ref})
                chart_count = len(getattr(ws, "_charts", []) or [])
                _ginsert(con, "sheet_chart_metadata", sid, ws.title, str(path), f"charts={chart_count}", {"chart_count": chart_count})
            _ginsert(con, "data_structure_signature", sid, "XLSX_SIGNATURE", str(path), f"sheets={len(wb.worksheets)}", {})
        except Exception:
            with zipfile.ZipFile(path) as z:
                sheets = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")]
                charts = [n for n in z.namelist() if n.startswith("xl/charts/") and n.endswith(".xml")]
            _ginsert(con, "sheet_workbook", sid, path.name, str(path), base.sha256_file(path), {"sheet_count": len(sheets), "zip_fallback": True})
            for n in sheets:
                name = Path(n).stem
                _ginsert(con, "sheet_tab", sid, name, str(path), n)
                _ginsert(con, "sheet_range", sid, name, str(path), "unknown_range_zip_fallback")
                _ginsert(con, "sheet_table", sid, name, str(path), "unknown_table_zip_fallback")
            _ginsert(con, "sheet_chart_metadata", sid, "charts", str(path), f"charts={len(charts)}", {"charts": charts[:50]})
            _ginsert(con, "data_structure_signature", sid, "XLSX_SIGNATURE", str(path), f"sheets={len(sheets)} charts={len(charts)} zip_fallback=true")
        _fts(con, "data_fts", sid, f"{path.name} {path.suffix} spreadsheet workbook parsed")
        con.commit()
    finally:
        try:
            con.commit(); con.execute("VACUUM")
        except Exception:
            pass
        con.close()


def _repair_csv(db: Path, source: dict):
    path = Path(source.get("path") or "")
    if not path.exists() or path.suffix.lower() not in {".csv", ".tsv"}:
        return
    sid = source.get("source_id") or "csv_" + _hash_text(str(path))[:12]
    delim = "\t" if path.suffix.lower() == ".tsv" else ","
    con = sqlite3.connect(str(db), timeout=60)
    try:
        with open(path, "r", encoding="utf-8", errors="replace", newline="") as f:
            reader = csv.reader(f, delimiter=delim)
            header = next(reader, [])
            _ginsert(con, "csv_header", sid, "header", str(path), json.dumps(header, ensure_ascii=False), {"columns": len(header)})
            count = 0
            for count, row in enumerate(reader, start=1):
                _ginsert(con, "csv_row_sample", sid, f"row_{count}", str(path), json.dumps(row, ensure_ascii=False))
                if count >= 5000:
                    break
        _ginsert(con, "data_structure_signature", sid, "CSV_SIGNATURE", str(path), f"columns={len(header)} sampled_rows={count}")
        _fts(con, "data_fts", sid, " ".join(map(str, header)))
        con.commit()
    finally:
        try:
            con.commit(); con.execute("VACUUM")
        except Exception:
            pass
        con.close()


def _repair_docs(db: Path, source: dict):
    path = Path(source.get("path") or "")
    if not path.exists() or path.suffix.lower() not in {".docx", ".md", ".txt", ".html", ".xml"}:
        return
    sid = source.get("source_id") or "doc_" + _hash_text(str(path))[:12]
    con = sqlite3.connect(str(db), timeout=60)
    try:
        paras, tables = [], []
        if path.suffix.lower() == ".docx":
            paras, tables = base.extract_docx_text(path)
        else:
            paras = path.read_text(encoding="utf-8", errors="replace").splitlines()
        _ginsert(con, "doc_file", sid, path.name, str(path), base.sha256_file(path), {"paragraphs": len(paras), "tables": len(tables)})
        _ginsert(con, "doc_structure", sid, "document_structure", str(path), f"paragraphs={len(paras)} tables={len(tables)}")
        for i, p in enumerate(paras, start=1):
            if not str(p).strip():
                continue
            table = "doc_heading" if re.match(r"^#{1,6}\s+", str(p).strip()) or (len(str(p).strip()) < 80 and i <= 10) else "doc_paragraph"
            rid = _ginsert(con, table, sid, f"block_{i}", str(path), str(p))
            if i % 20 == 0:
                _fts(con, "doc_fts", rid, "\n".join(map(str, paras[max(0, i-20):i])))
        for i in range(0, len(paras), 80):
            chunk = "\n".join(map(str, paras[i:i+80])).strip()
            if chunk:
                cid = _ginsert(con, "doc_chunk", sid, f"chunk_{i//80+1}", str(path), chunk, {"start": i+1})
                _fts(con, "doc_fts", cid, chunk)
        for ti, tbl in enumerate(tables, start=1):
            _ginsert(con, "doc_table_extract", sid, f"table_{ti}", str(path), json.dumps(tbl, ensure_ascii=False))
        _ginsert(con, "source_structure_signature", sid, "DOC_SIGNATURE", str(path), f"paragraphs={len(paras)} tables={len(tables)}")
        con.commit()
    finally:
        try:
            con.commit(); con.execute("VACUUM")
        except Exception:
            pass
        con.close()


def _repair_selected_non_code(root: Path, sources: list[dict], progress=None):
    project = root / "project"
    for source in sources:
        if not source.get("active", True):
            continue
        lane = source.get("lane_key")
        if lane == "data_excel_csv":
            db = project / "sectors" / lane / f"{lane}_sector_v001.sqlite"
            _repair_csv(db, source); _repair_xlsx(db, source)
        elif lane == "docs":
            db = project / "sectors" / lane / f"{lane}_sector_v001.sqlite"
            _repair_docs(db, source)
    _emit(progress, "repair", "selected non-code lane repair complete", str(root), 88, 1, 1)


def _repair_code_lanes(root: Path, sources: list[dict], progress=None):
    project = root / "project"
    for source in sources:
        if not source.get("active", True):
            continue
        lane = source.get("lane_key")
        if lane not in {"local_code", "github", "code"}:
            continue
        lane = "local_code" if lane == "code" else lane
        db = project / "sectors" / lane / f"{lane}_sector_v001.sqlite"
        _repair_dependencies_and_workflow(db, {**source, "lane_key": lane}, project, progress)
    _emit(progress, "repair", "code lane dependency/workflow repair complete", str(root), 92, 1, 1)


def build_brain(workspace_dir: str, brain_name: str, sources: list[dict], progress=None):
    result = base.build_brain(workspace_dir, brain_name, sources, progress)
    root = Path(result.get("brain_root") or base.brain_output_dir(workspace_dir, brain_name))
    _repair_code_lanes(root, sources, progress)
    _repair_selected_non_code(root, sources, progress)
    write_code_mmd(root)
    base.write_project_master_mmd(root)
    _emit(progress, "done", "brain build complete with V5.3 rollback delta repairs", str(root), 100, 1, 1)
    return {**result, "runtime": APP_VERSION}


def _mmd_safe(s: str, limit=80):
    s = str(s or "").replace('"', "'").replace("[", " ").replace("]", " ").replace("{", " ").replace("}", " ").replace("|", "/")
    s = re.sub(r"\s+", " ", s).strip()
    return ("..." + s[-limit:]) if len(s) > limit else (s or "none")


def _rows(con, sql, args=()):
    try:
        return con.execute(sql, args).fetchall()
    except Exception:
        return []


def write_code_mmd(root: Path):
    db = root / "project" / "sectors" / "local_code" / "local_code_sector_v001.sqlite"
    top = root / "project" / "topology"
    top.mkdir(parents=True, exist_ok=True)
    out = top / "local_code_lane.mmd"
    if not db.exists():
        out.write_text('flowchart LR\n  ROOT["coded project workflow"]\n', encoding="utf-8")
        return
    con = sqlite3.connect(str(db))
    routes = _rows(con, 'SELECT route_path, route_type, file_id FROM app_route ORDER BY route_path LIMIT 100')
    deps = _rows(con, 'SELECT package_name, version_spec FROM dependency_item ORDER BY package_name LIMIT 80')
    arts = _rows(con, 'SELECT artifact_type, COUNT(*), COALESCE(SUM(size_bytes),0) FROM project_artifact GROUP BY artifact_type ORDER BY COUNT(*) DESC LIMIT 40')
    file_lookup = {r[0]: r[1] for r in _rows(con, 'SELECT file_id, canonical_path FROM code_file')}
    lines = [
        "flowchart LR",
        "  classDef route fill:#fff4d6,stroke:#8a5a00,color:#111;",
        "  classDef code fill:#eaffea,stroke:#275,color:#111;",
        "  classDef dep fill:#fef3c7,stroke:#92400e,color:#111;",
        "  classDef artifact fill:#f3e8ff,stroke:#635,color:#111;",
        "  classDef root fill:#111827,stroke:#111827,color:#fff;",
        '  ROOT["coded project workflow"]:::root',
        '  ROUTES["routes / pages"]:::route',
        '  CODE["code files / line snapshots / chunks"]:::code',
        '  DEPS["tools + dependencies"]:::dep',
        '  ARTS["read-only artifacts"]:::artifact',
        "  ROOT --> ROUTES", "  ROUTES --> CODE", "  CODE --> DEPS", "  CODE --> ARTS",
        "  subgraph R[route/page to code path map]", "    direction TB",
    ]
    prev = "ROUTES"
    for i, (route, rtype, fid) in enumerate(routes):
        rn = f"R{i}"; fn = f"RF{i}"; path = file_lookup.get(fid, fid)
        lines.append(f'    {rn}["{_mmd_safe(route,70)}\\n{_mmd_safe(rtype,35)}"]:::route')
        lines.append(f'    {fn}["{_mmd_safe(path,95)}"]:::code')
        lines.append(f"    {prev} --> {rn}"); lines.append(f"    {rn} --> {fn}"); lines.append(f"    {fn} --> CODE"); prev = rn
    lines += ["  end", "  subgraph D[dependency map]", "    direction TB"]
    prev = "DEPS"
    for i, (name, ver) in enumerate(deps):
        dn = f"D{i}"; lines.append(f'    {dn}["{_mmd_safe(name,60)}\\n{_mmd_safe(ver,40)}"]:::dep'); lines.append(f"    {prev} --> {dn}"); prev = dn
    lines += ["  end", "  subgraph A[artifact payload groups]", "    direction TB"]
    prev = "ARTS"
    for i, (atype, count, size) in enumerate(arts):
        an = f"A{i}"; lines.append(f'    {an}["{_mmd_safe(atype,50)}\\ncount={count} bytes={size}"]:::artifact'); lines.append(f"    {prev} --> {an}"); prev = an
    lines.append("  end")
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")
    con.close()


def render_topology(workspace_dir: str, brain_name: str, progress=None):
    root = base.brain_output_dir(workspace_dir, brain_name)
    write_code_mmd(root); base.write_project_master_mmd(root)
    return base.render_topology(workspace_dir, brain_name, progress)


def export_one_upload_package(workspace_dir: str, brain_name: str, progress=None):
    root = base.brain_output_dir(workspace_dir, brain_name)
    write_code_mmd(root); base.write_project_master_mmd(root)
    return base.export_one_upload_package(workspace_dir, brain_name, progress)


EXACT10 = ["GEMINI_FLASH_PROMPT.txt", "UEPC_UNIVERSAL_POINTER.json", "UEPC_DOT_LOCKS_AND_READ_ORDER.txt", "PUBLIC_ENV_UOP_PROJECT_LOCKED.zip", "GENERATED_PROJECT_CONJOINED.sqlite", "PROJECT_GITLIKE_FILE_LEDGER.jsonl", "PROJECT_TOPOLOGY.mmd", "PROJECT_TOPOLOGY.svg", "PROJECT_TOPOLOGY.png", "MANIFEST_RECEIPTS_RECOVERY.json"]


def _latest_normal_zip(root: Path, brain_name: str) -> Path:
    packages = root / "packages"; zips = []
    if packages.exists():
        for p in packages.glob("*one_upload_package_v*.zip"):
            if "gemini" not in p.name.lower(): zips.append(p)
    if not zips: raise RuntimeError("NORMAL_ONE_UPLOAD_ZIP_NOT_FOUND - run Export One-Upload Package first")
    return sorted(zips, key=lambda p: p.stat().st_mtime, reverse=True)[0]


def _zip_read(z, options):
    for n in options:
        try: return z.read(n)
        except KeyError: pass
    return b""


def _public_locked_env_zip(src_zip: Path, dst_zip: zipfile.ZipFile):
    tmp = Path(tempfile.mkdtemp()) / "PUBLIC_ENV_UOP_PROJECT_LOCKED.zip"
    with zipfile.ZipFile(src_zip, "r") as z, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        for name in z.namelist():
            low = name.lower()
            if name in {".uepc_env", ".uepc_profile", ".uepc_project", "FLASH_ME_FIRST_SINGLE_PROMPT.txt", "README_NEXT_PROMPT.txt"} or low.startswith(("env/", "uop/", "project_template_locked/")) or low in {"project/project_template.sqlite", "project/project_topology_template.mmd", "project/project_topology_template.svg", "project/project_topology_template.png"}:
                out.writestr(name, z.read(name))
    dst_zip.write(tmp, "PUBLIC_ENV_UOP_PROJECT_LOCKED.zip"); shutil.rmtree(tmp.parent, ignore_errors=True)


def _conjoin_project_db(normal_zip: Path) -> Path:
    tmpdir = Path(tempfile.mkdtemp(prefix="uepc_gemini_conjoin_")); out = tmpdir / "GENERATED_PROJECT_CONJOINED.sqlite"
    dest = sqlite3.connect(str(out))
    dest.execute('CREATE TABLE uepc_source_db_index(entry_path TEXT PRIMARY KEY, sha256 TEXT, size_bytes INTEGER, table_count INTEGER, created_at TEXT)')
    dest.execute('CREATE TABLE uepc_table_inventory(entry_path TEXT, table_name TEXT, row_count INTEGER)')
    dest.execute('CREATE TABLE uepc_conjoin_row(entry_path TEXT, table_name TEXT, primary_key TEXT, row_json TEXT)')
    dest.execute('CREATE TABLE uepc_universal_pointer(key TEXT PRIMARY KEY, value TEXT)')
    with zipfile.ZipFile(normal_zip, "r") as z:
        members = [n for n in z.namelist() if n == "project/project_router.sqlite" or (n.startswith("project/sectors/") and n.endswith(".sqlite") and not n.endswith(("-wal", "-shm")))]
        for name in members:
            data = z.read(name); sha = hashlib.sha256(data).hexdigest(); tmpdb = tmpdir / re.sub(r"[^A-Za-z0-9_.-]+", "_", name); tmpdb.write_bytes(data)
            try:
                src = sqlite3.connect(str(tmpdb)); src.row_factory = sqlite3.Row
                tables = [r[0] for r in src.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name")]
                dest.execute('INSERT OR REPLACE INTO uepc_source_db_index VALUES(?,?,?,?,?)', (name, sha, len(data), len(tables), base.now()))
                for t in tables:
                    if t.endswith(FTS_SHADOW_SUFFIXES): continue
                    try:
                        rc = src.execute(f'SELECT COUNT(*) FROM "{t}"').fetchone()[0]
                        dest.execute('INSERT INTO uepc_table_inventory VALUES(?,?,?)', (name, t, rc))
                        for row in src.execute(f'SELECT * FROM "{t}"'):
                            d = dict(row); pk = d.get("id") or d.get("file_id") or d.get("route_id") or d.get("artifact_id") or d.get("commit_sha") or d.get("pointer_id") or ""
                            dest.execute('INSERT INTO uepc_conjoin_row VALUES(?,?,?,?)', (name, t, str(pk), json.dumps(d, ensure_ascii=False, default=str)))
                    except Exception:
                        dest.execute('INSERT INTO uepc_table_inventory VALUES(?,?,?)', (name, t, -1))
                src.close()
            finally:
                tmpdb.unlink(missing_ok=True)
    for k, v in {"created_at": base.now(), "mode": "GEMINI_EXACT10_PROJECT_CONJOINED_DB", "normal_package": str(normal_zip), "locked_env_file": "PUBLIC_ENV_UOP_PROJECT_LOCKED.zip"}.items():
        dest.execute('INSERT OR REPLACE INTO uepc_universal_pointer VALUES(?,?)', (k, v))
    dest.commit(); dest.execute("VACUUM"); dest.close(); return out


def _gitlike_ledger(normal_zip: Path) -> bytes:
    lines = []
    with zipfile.ZipFile(normal_zip, "r") as z:
        for name in z.namelist():
            if name.endswith("/"): continue
            if name.startswith(("project/", "manifests/", "receipts/", "recovery/")) or name in {"FLASH_ME_FIRST_SINGLE_PROMPT.txt", ".uepc_env", ".uepc_project", ".uepc_profile"}:
                data = z.read(name); lines.append(json.dumps({"path": name, "sha256": hashlib.sha256(data).hexdigest(), "size": len(data)}, ensure_ascii=False))
    return ("\n".join(lines) + "\n").encode("utf-8")


def _gemini_prompt() -> bytes:
    return ("UEPC-GEMINI-EXACT10-PROJECT-CONJOINED-BOOT-001 | Mode: pointer_first_validation | Category: exact-10 Gemini package\n\n"
            "CHAT_NAME:\n[GOLDV3]\n\n"
            "This is a Gemini-compatible exact-10 package, not the normal one-upload package. Read the 10 root files only.\n\n"
            "READ ORDER:\n1. GEMINI_FLASH_PROMPT.txt\n2. UEPC_UNIVERSAL_POINTER.json\n3. UEPC_DOT_LOCKS_AND_READ_ORDER.txt\n4. PUBLIC_ENV_UOP_PROJECT_LOCKED.zip\n5. GENERATED_PROJECT_CONJOINED.sqlite\n6. PROJECT_GITLIKE_FILE_LEDGER.jsonl\n7. PROJECT_TOPOLOGY.mmd\n8. PROJECT_TOPOLOGY.svg\n9. PROJECT_TOPOLOGY.png\n10. MANIFEST_RECEIPTS_RECOVERY.json\n\n"
            "Env/UOP/project-template inside PUBLIC_ENV_UOP_PROJECT_LOCKED.zip are locked/read-only. GENERATED_PROJECT_CONJOINED.sqlite is the joined generated-project brain view from project_router + all sector DBs.\n\n"
            "Run compact entry slip: package_seen= exact10_file_count= universal_pointer_seen= locked_env_package_seen= conjoined_project_db_seen= gitlike_ledger_seen= topology_seen= blocker_if_any=\n").encode("utf-8")


def export_gemini_exact10(workspace_dir: str, brain_name: str, progress=None):
    root = base.brain_output_dir(workspace_dir, brain_name); normal_zip = _latest_normal_zip(root, brain_name)
    packages = root / "packages"; packages.mkdir(parents=True, exist_ok=True); slug = base.slugify_name(brain_name); out_zip = packages / f"{slug}_gemini_exact10_v056.zip"
    if out_zip.exists(): out_zip.unlink()
    conjoined = _conjoin_project_db(normal_zip)
    pointer = {"created_at": base.now(), "gemini_file_count": 10, "normal_source_zip": str(normal_zip), "generated_project_db": "GENERATED_PROJECT_CONJOINED.sqlite", "locked_env": "PUBLIC_ENV_UOP_PROJECT_LOCKED.zip", "files": EXACT10}
    locks = "UEPC dot-locks/read-order for Gemini exact10\n.uepc_env=inside PUBLIC_ENV_UOP_PROJECT_LOCKED.zip\n.uepc_project=inside PUBLIC_ENV_UOP_PROJECT_LOCKED.zip + GENERATED_PROJECT_CONJOINED.sqlite\n.uepc_profile=inside PUBLIC_ENV_UOP_PROJECT_LOCKED.zip\n"
    with zipfile.ZipFile(normal_zip, "r") as z:
        topo_mmd = _zip_read(z, ["project/topology/local_code_lane.mmd", "project/topology/project_master_topology.mmd", "project/project_topology_template.mmd"])
        topo_svg = _zip_read(z, ["project/topology/local_code_lane.svg", "project/topology/project_master_topology.svg", "project/project_topology_template.svg"])
        topo_png = _zip_read(z, ["project/topology/local_code_lane_CRYSTAL.png", "project/topology/local_code_lane_4K.png", "project/topology/local_code_lane.png", "project/project_topology_template.png"])
    manifest = {"created_at": base.now(), "normal_zip_sha256": base.sha256_file(normal_zip), "purpose": "Gemini exact10 only; normal export untouched.", "files": EXACT10}
    with zipfile.ZipFile(out_zip, "w", compression=zipfile.ZIP_DEFLATED) as oz:
        oz.writestr("GEMINI_FLASH_PROMPT.txt", _gemini_prompt()); oz.writestr("UEPC_UNIVERSAL_POINTER.json", json.dumps(pointer, indent=2).encode("utf-8")); oz.writestr("UEPC_DOT_LOCKS_AND_READ_ORDER.txt", locks.encode("utf-8")); _public_locked_env_zip(normal_zip, oz); oz.write(conjoined, "GENERATED_PROJECT_CONJOINED.sqlite"); oz.writestr("PROJECT_GITLIKE_FILE_LEDGER.jsonl", _gitlike_ledger(normal_zip)); oz.writestr("PROJECT_TOPOLOGY.mmd", topo_mmd or b"flowchart TD\n  A[No topology found]\n"); oz.writestr("PROJECT_TOPOLOGY.svg", topo_svg or b"<svg xmlns='http://www.w3.org/2000/svg'><text x='10' y='20'>No SVG topology found</text></svg>"); oz.writestr("PROJECT_TOPOLOGY.png", topo_png or b""); oz.writestr("MANIFEST_RECEIPTS_RECOVERY.json", json.dumps(manifest, indent=2).encode("utf-8"))
    shutil.rmtree(conjoined.parent, ignore_errors=True)
    with zipfile.ZipFile(out_zip, "r") as z:
        names = z.namelist()
        if len(names) != 10: raise RuntimeError(f"GEMINI_EXACT10_FAILED: {len(names)} files: {names}")
        missing = [n for n in EXACT10 if n not in names]
        if missing: raise RuntimeError(f"GEMINI_EXACT10_MISSING: {missing}")
    _emit(progress, "done", "Gemini exact10 project-conjoined package ready", str(out_zip), 100, 1, 1)
    return {"gemini_package_zip": str(out_zip), "file_count": 10, "status": "GEMINI_EXACT10_PROJECT_CONJOINED_READY", "sha256": base.sha256_file(out_zip)}


def export_gemini_compatible_package(workspace_dir: str, brain_name: str, progress=None):
    return export_gemini_exact10(workspace_dir, brain_name, progress)
