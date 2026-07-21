from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import shlex
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path
from typing import Callable, Optional
from xml.etree import ElementTree as ET

from sqlite_brain_builder.ingest.lane_contracts import parse_fields
from sqlite_brain_builder.runtime.canonical_lanes import (
    LANE_REGISTRY,
    UnknownLaneAliasError,
    resolve_lane_id,
    validate_primary_code_mode,
)
from sqlite_brain_builder.runtime.chatlineage_state_travel_v57 import (
    GEMINI_FLASH_APPEND,
    patch_package_folder,
)
from sqlite_brain_builder.runtime.env15_resource import (
    find_env15_resource_root,
    validate_env15_resource,
)
from sqlite_brain_builder.runtime.env15_ingestion import (
    UNIVERSAL_INGESTION_LANES,
    ingest_env15_universal_source,
)
from sqlite_brain_builder.runtime.env15_code_ingestion import ingest_env15_code_source
from sqlite_brain_builder.runtime.env15_chat_lineage import append_env15_chat_lineage_source
from sqlite_brain_builder.runtime.env15_research import (
    append_env15_research_source,
    ensure_env15_research_append_schema,
)
from sqlite_brain_builder.runtime.env15_project_schema import (
    governed_sector_mutation,
    initialize_env15_brain_root,
    resolve_env15_sector,
    write_env15_runtime_sector_state,
)
from sqlite_brain_builder.runtime.env15_topology import (
    generate_env15_topologies,
    render_env15_topologies,
)
from sqlite_brain_builder.runtime.office_ingest import (
    parse_data_source,
    parse_docx,
    parse_pptx,
    stable_record_id,
)
from sqlite_brain_builder.runtime.ocr_ingest import parse_image_ocr, parse_pdf_ocr
from sqlite_brain_builder.runtime.package_parity import validate_active_project_package_parity
from sqlite_brain_builder.runtime.package_validation import validate_chatgpt_package
from sqlite_brain_builder.runtime.provider_package_projection import (
    CHATGPT_CODE_PROJECTION_MIN_BYTES,
    CHATGPT_PACKAGE_MAX_BYTES,
    GEMINI_PACKAGE_MAX_BYTES,
    clear_provider_skip_marker,
    compact_chatgpt_package_stage,
    enforce_provider_package_size,
    reconcile_project_router_authority,
    remove_provider_stage,
    retire_provider_outputs,
    reseal_chatgpt_package_authority,
    write_provider_skip_marker,
)
from sqlite_brain_builder.runtime.provider_readable_corpus import (
    export_gemini_provider_readable,
    materialize_provider_readable_tree,
)
from sqlite_brain_builder.runtime.universal_lane_authority import (
    find_forbidden_provider_project_lock_references,
    prune_provider_project_lock_artifacts,
    write_universal_lane_authority,
)
from sqlite_brain_builder.runtime.plan_delta_governance import (
    ensure_plan_delta_sector,
    register_initial_plan_delta,
)
from sqlite_brain_builder.runtime.structural_ingestion import (
    ingest_brain_loader,
    ingest_project_engulf,
    ingest_sqlite_brain,
)

ProgressCallback = Optional[Callable[[dict], None]]

STRUCTURAL_LANE_INGESTERS = {
    "brain_loader": ingest_brain_loader,
    "project_engulf": ingest_project_engulf,
    "sqlite_brain": ingest_sqlite_brain,
}

_FTS_SHADOW_SUFFIXES = ("_config", "_content", "_data", "_docsize", "_idx")


def _project_structural_lane_tables(
    root: Path,
    lane_id: str,
    source: dict,
    universal_result: object,
) -> object:
    """Project a specialized parser DB into its governed Env15 sector."""
    status = str(getattr(universal_result, "status", "PASS"))
    if status.startswith("SKIPPED"):
        return universal_result
    source_path = Path(str(source.get("path") or "")).expanduser().resolve(strict=True)
    prefix = {
        "brain_loader": "brain_loader",
        "research": "research",
        "project_engulf": "project_engulf",
        "sqlite_brain": "loaded_sqlite_brain",
    }[lane_id]
    with tempfile.TemporaryDirectory(prefix=f"evidenceos_{lane_id}_projection_", dir=root / "receipts") as temporary:
        staging = Path(temporary) / f"{lane_id}_structural.sqlite"
        structural_receipt = STRUCTURAL_LANE_INGESTERS[lane_id](source_path, staging)
        source_connection = sqlite3.connect(staging)
        source_connection.row_factory = sqlite3.Row
        try:
            table_rows = list(source_connection.execute(
                "SELECT name,sql,rowid FROM sqlite_master WHERE type='table' AND name LIKE ? ORDER BY rowid",
                (prefix + "%",),
            ))
            tables = [
                row for row in table_rows
                if row["sql"]
                and not re.search(r"\busing\s+fts[345]\b", str(row["sql"]), flags=re.IGNORECASE)
                and not any(row["name"].endswith(suffix) for suffix in _FTS_SHADOW_SUFFIXES)
            ]
            payloads = {}
            for row in tables:
                name = row["name"]
                columns = [item[1] for item in source_connection.execute(f'PRAGMA table_info("{name}")')]
                payloads[name] = (columns, [tuple(item) for item in source_connection.execute(f'SELECT * FROM "{name}"')])
        finally:
            source_connection.close()

        turn_id = "struct_" + sha256_bytes(f"{lane_id}|{source_path}|{structural_receipt.source_sha256}".encode())[:40]

        def mutate(connection: sqlite3.Connection) -> dict[str, int]:
            for row in tables:
                schema_sql = re.sub(
                    r"^CREATE\s+TABLE\s+", "CREATE TABLE IF NOT EXISTS ",
                    str(row["sql"]), count=1, flags=re.IGNORECASE,
                )
                connection.execute(schema_sql)
            inserted = 0
            for name, (columns, rows_) in payloads.items():
                if not rows_:
                    continue
                quoted_columns = ",".join('"' + column.replace('"', '""') + '"' for column in columns)
                placeholders = ",".join("?" for _ in columns)
                before = connection.total_changes
                connection.executemany(
                    f'INSERT OR IGNORE INTO "{name}"({quoted_columns}) VALUES({placeholders})', rows_
                )
                inserted += connection.total_changes - before
            return {
                "structural_projection_rows": inserted,
                "structural_projection_tables": len(payloads),
                "structural_source_preserved": int(structural_receipt.source_preserved),
            }

        governed_sector_mutation(
            root,
            lane_id,
            actor="Evidence OS SQLite Builder",
            reason=f"project specialized {lane_id} structural tables",
            turn_id=turn_id,
            mutate=mutate,
        )
    return universal_result


def ingest_env15_structural_source(root: Path, lane_id: str, source: dict) -> object:
    if lane_id == "research":
        return append_env15_research_source(root, source)
    universal_result = ingest_env15_universal_source(
        root,
        lane_id,
        source,
        actor="Evidence OS SQLite Builder",
        reason=f"explicit {lane_id} source intake",
    )
    return _project_structural_lane_tables(root, lane_id, source, universal_result)


_CANONICAL_SEMANTIC_PROJECTION_LANES = {
    "discussion", "analysis", "plan", "mode", "docs", "data_excel",
    "ppt", "pdf_ocr", "images_ocr", "artifacts", "custom",
}


def _canonical_semantic_projection_mutator(
    root: Path,
    lane_id: str,
    source: dict,
) -> Callable[[sqlite3.Connection], dict[str, int]] | None:
    """Stage rich lane rows, then project them in the governed mutation."""

    if lane_id not in _CANONICAL_SEMANTIC_PROJECTION_LANES:
        return None
    source_id = str(source.get("source_id") or "")
    if not source_id:
        raise RuntimeError(f"CANONICAL_PROJECTION_SOURCE_ID_REQUIRED:{lane_id}")
    user_projection_tables: set[str] = set()
    if lane_id == "custom":
        raw_contract = source.get("schema_contract") or source.get("schemaContract") or ""
        contract_text = (
            "\n".join(str(value) for value in raw_contract)
            if isinstance(raw_contract, (list, tuple))
            else str(raw_contract)
        )
        user_projection_tables = set(parse_fields(contract_text))
    with tempfile.TemporaryDirectory(
        prefix=f"evidenceos_{lane_id}_semantic_projection_",
        dir=root / "receipts",
    ) as temporary:
        staging = Path(temporary) / f"{lane_id}_semantic.sqlite"
        build_semantic_lane(staging, lane_id, source, None, time.time())
        source_connection = sqlite3.connect(staging)
        source_connection.row_factory = sqlite3.Row
        try:
            table_sql: dict[str, str] = {}
            payloads: dict[str, tuple[list[str], list[tuple]]] = {}
            for table in LANE_DEFS[lane_id]["schema"]:
                if table.endswith("_fts"):
                    continue
                row = source_connection.execute(
                    "SELECT name,sql,rowid FROM sqlite_master WHERE type='table' AND name=?",
                    (table,),
                ).fetchone()
                if row is None or not row["sql"]:
                    continue
                table_sql[table] = str(row["sql"])
                columns = [item[1] for item in source_connection.execute(f'PRAGMA table_info("{table}")')]
                payloads[table] = (
                    columns,
                    [tuple(item) for item in source_connection.execute(f'SELECT * FROM "{table}"')],
                )
        finally:
            source_connection.close()

    def mutate(connection: sqlite3.Connection) -> dict[str, int]:
        created = 0
        incompatible_preserved = 0
        user_projection_tables_preserved = 0
        for name, create_sql in table_sql.items():
            exists = connection.execute(
                "SELECT 1 FROM sqlite_schema WHERE type='table' AND name=?", (name,)
            ).fetchone()
            if not exists:
                connection.execute(create_sql)
                created += 1
        inserted = 0
        projected_tables = 0
        for name, (columns, rows_) in payloads.items():
            if name in user_projection_tables:
                user_projection_tables_preserved += 1
                projected_tables += 1
                continue
            target_columns = [item[1] for item in connection.execute(f'PRAGMA table_info("{name}")')]
            if target_columns != columns:
                incompatible_preserved += 1
                continue
            if "source_id" in columns:
                connection.execute(f'DELETE FROM "{name}" WHERE source_id=?', (source_id,))
            if rows_:
                quoted_columns = ",".join('"' + column.replace('"', '""') + '"' for column in columns)
                placeholders = ",".join("?" for _ in columns)
                before = connection.total_changes
                connection.executemany(
                    f'INSERT OR REPLACE INTO "{name}"({quoted_columns}) VALUES({placeholders})',
                    rows_,
                )
                inserted += connection.total_changes - before
            projected_tables += 1
        return {
            "canonical_projection_rows": inserted,
            "canonical_projection_tables": projected_tables,
            "canonical_projection_tables_created": created,
            "canonical_projection_incompatible_tables_preserved": incompatible_preserved,
            "canonical_projection_user_tables_preserved": user_projection_tables_preserved,
        }

    return mutate


def ingest_env15_universal_with_projection(root: Path, lane_id: str, source: dict) -> object:
    return ingest_env15_universal_source(
        root,
        lane_id,
        source,
        actor="Evidence OS SQLite Builder",
        reason=f"explicit {lane_id} source intake",
        mutation_extension=_canonical_semantic_projection_mutator(root, lane_id, source),
    )

APP_VERSION = "V5.3_CLEAN_COMPACT_REAL_LANES"

if os.name == "nt":
    _HIDDEN_SI = subprocess.STARTUPINFO()
    _HIDDEN_SI.dwFlags |= subprocess.STARTF_USESHOWWINDOW
    _HIDDEN_SI.wShowWindow = 0
    _HIDDEN_FLAGS = subprocess.CREATE_NO_WINDOW
else:
    _HIDDEN_SI = None
    _HIDDEN_FLAGS = 0


def _hidden_kwargs(kwargs):
    if os.name == "nt":
        kwargs.setdefault("startupinfo", _HIDDEN_SI)
        kwargs.setdefault("creationflags", _HIDDEN_FLAGS)
        kwargs.setdefault("stdin", subprocess.DEVNULL)
    return kwargs


def _git_longpaths(cmd):
    if isinstance(cmd, (list, tuple)) and cmd and cmd[0] == "git":
        return ["git", "-c", "core.longpaths=true", *cmd[1:]]
    return cmd


def run_hidden(cmd, **kwargs):
    cmd = _git_longpaths(cmd)
    return subprocess.run(cmd, **_hidden_kwargs(kwargs))


def popen_hidden(cmd, **kwargs):
    cmd = _git_longpaths(cmd)
    return subprocess.Popen(cmd, **_hidden_kwargs(kwargs))


def check_output_hidden(cmd, **kwargs):
    cmd = _git_longpaths(cmd)
    return subprocess.check_output(cmd, **_hidden_kwargs(kwargs))


def check_call_hidden(cmd, **kwargs):
    cmd = _git_longpaths(cmd)
    return subprocess.check_call(cmd, **_hidden_kwargs(kwargs))


LOCKED_FLASH_PROMPT = """EVIDENCE LANE ENV/UOP + LIVE PROJECT BOOT V5.9

CHAT_NAME:
BRIEF_NATURE_OF_CHAT:

READ ORDER:
1. FLASH_ME_FIRST_SINGLE_PROMPT.txt
2. .uepc_env
3. .uepc_profile
4. .uepc_project
5. env/env_law.md and env/env_sqlite.sqlite
6. uop/uop_law.md and uop/uop_sqlite.sqlite
7. project/project_router.sqlite
8. project/pointers/INDEX.json
9. the selected universal lane pointer and its exact live Project sector database
10. manifests and receipts

AUTHORITY:
Env and UOP are the only locked read-only authorities. Project work occurs directly in the live governed sector graph. Local Code and GitHub Code are mutually exclusive intake modes for one Project database.

SOURCE-INTAKE LAW:
A registered lane pointer identifies its exact database, SHA-256, source count, build state, access, and HIL gate. An omitted lane is SKIPPED_NO_SOURCE and its prior bytes remain unchanged. Do not infer population from package size.

AUTOMATIC APPEND LAW:
Exactly two Project lanes accept automatic append-only model writeback:
- chat_lineage
- research

Every visible user/model output must carry a stable identity, actor/provider, source time when known, ingestion time, and content SHA-256. Chat Lineage reimports append only the unseen suffix of the same source. Research records the visible question, candidate finding, evidence, limitation, and HIL boundary. Neither lane promotes candidate truth.

OTHER PROJECT LANES:
Every other Project lane is read-only unless the user issues an explicit named one-turn HIL grant for that lane. Validate, receipt, hash, and relock the exact sector after an authorized write.

VISIBLE TELEMETRY:
Open with ENTRY SLIP | mode | active pointer | gates | operators | write scope.
Close with EXIT SLIP | read/write status | mutation yes/no | append receipt | gates | next pointer.
Never store hidden chain-of-thought; store only user-visible summaries and evidence.

VALIDATION BEFORE CLAIMING WRITEBACK:
- ZIP CRC and member manifest
- SQLite integrity and foreign keys for Env, UOP, router, and affected sectors
- universal pointer database/hash parity
- Research and Chat Lineage append-only receipts
- topology/hash proof when present

Do not claim a mutation, append, Refresh, Fuse, HIL approval, or promotion unless its exact receipt and revalidation are present.
"""

CODE_EXTS = {
    ".py", ".js", ".jsx", ".ts", ".tsx", ".mjs", ".cjs", ".css", ".scss", ".html", ".htm",
    ".sql", ".sh", ".ps1", ".bat", ".cmd", ".yml", ".yaml", ".toml", ".ini", ".cfg",
}
CODE_DOC_EXTS = {".md", ".txt"}
TEXT_ARTIFACT_EXTS = {".json", ".jsonl", ".csv", ".tsv", ".xml", ".svg", ".mmd", ".html", ".md", ".txt", ".ipynb"}
BINARY_ASSET_EXTS = {
    ".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".ico", ".tif", ".tiff", ".avif",
    ".glb", ".gltf", ".mp4", ".mov", ".avi", ".mkv", ".webm", ".wav", ".mp3",
}
HEAVY_SKIP_DIRS = {".git", "node_modules", ".next", "dist", "build", "__pycache__", ".venv", "venv", ".pytest_cache"}
MAX_ARTIFACT_TEXT_EXTRACT = 1_000_000
CHUNK_LINES = 120
MAX_CODE_CHUNK_TEXT = 12_000
MAX_EXACT_GIT_DOC_BYTES = 256_000
GENERATED_HISTORY_DIRS = {
    "artifacts", "cache", "catboost_info", "checkpoints", "data", "generated", "logs", "model_outputs",
    "outputs", "predictions", "reports", "results", "runs", "snapshots", "temp", "tmp",
}
GENERATED_HISTORY_EXTS = {".csv", ".jsonl", ".parquet", ".tsv", ".xls", ".xlsx"}
SMALL_HUMAN_DOC_EXTS = {".adoc", ".html", ".md", ".mmd", ".rst", ".svg", ".txt", ".xml"}
CODE_INDEX_POLICY_VERSION = "tiered-git-history-v3-author-committer-exact-lines"
CONFIG_JSON_NAMES = {
    "manifest.json", "package.json", "settings.json", "tsconfig.json", "vercel.json",
}
DEPENDENCY_MANIFEST_NAMES = {
    "package.json",
    "requirements.txt",
    "pyproject.toml",
    "package-lock.json",
    "pnpm-lock.yaml",
    "yarn.lock",
    "poetry.lock",
}
SPECIAL_CODE_FILENAMES = {"dockerfile", "makefile", "justfile", "procfile"}

_LANE_TABS = {
    "github_code": "GitHub",
    "local_code": "Local Code",
    "chat_lineage": "Chat Lineage",
    "discussion": "Discussion",
    "analysis": "Analysis",
    "plan": "Plan",
    "mode": "Mode",
    "docs": "Docs",
    "data_excel": "Data",
    "ppt": "PPT",
    "pdf_ocr": "PDF/OCR",
    "images_ocr": "Images/OCR",
    "artifacts": "Artifacts",
    "custom": "Custom",
    "brain_loader": "Brain Loader",
    "research": "Research",
    "project_engulf": "Project Engulf",
    "sqlite_brain": "SQLite Brain",
}


def _lane_filetypes(label: str, extensions: tuple[str, ...]):
    patterns = " ".join(f"*{extension}" for extension in extensions)
    return [(f"{label} files", patterns)]


# Compatibility payload for the existing runtime.  Identity, aliases, sector
# destination, schema, parser, FTS, MMD, and mutation policy all come from the
# single canonical registry; the legacy runtime-specific `tab`/`filetypes`
# fields are derived here rather than maintained as a second lane registry.
LANE_DEFS = {
    lane_id: {
        "canonical_lane_id": lane.canonical_lane_id,
        "label": lane.display_label,
        "tab": _LANE_TABS[lane_id],
        "source_types": list(lane.source_types),
        "extensions": list(lane.extensions),
        "filetypes": _lane_filetypes(lane.display_label, lane.extensions),
        "schema": list(lane.schema_contract),
        "frontend_aliases": list(lane.frontend_aliases),
        "backend_aliases": list(lane.backend_aliases),
        "sector_folder": lane.sector_folder,
        "sqlite_filename": lane.sqlite_filename,
        "parser_id": lane.parser_id,
        "chunker_version": lane.chunker_version,
        "fts_table": lane.fts_table,
        "mmd_node_id": lane.mmd_node_id,
        "mutation_policy": lane.mutation_policy,
    }
    for lane_id, lane in LANE_REGISTRY.items()
}

TAB_ORDER = [*_LANE_TABS.values(), "Packages", "Receipts"]


def now():
    return time.strftime("%Y-%m-%dT%H:%M:%S")


def slugify_name(name: str) -> str:
    s = re.sub(r"[^A-Za-z0-9._-]+", "_", (name or "new_brain").strip()).strip("._-")
    return (s or "new_brain").lower()


def normalize_workspace_dir(workspace_dir: str) -> Path:
    p = Path(workspace_dir).expanduser()
    # No hidden brains/brains nesting. User workspace root is the direct parent of brain folders.
    while p.name.lower() == "brains":
        p = p.parent
    return p


def brain_output_dir(workspace_dir: str, brain_name: str) -> Path:
    root = normalize_workspace_dir(workspace_dir)
    slug = slugify_name(brain_name)
    output_slug = f"{slug}_output"
    if root.name.lower() in {slug.lower(), output_slug.lower()}:
        return root
    return root / output_slug


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def file_id_for(rel: str, digest: str) -> str:
    return "file_" + hashlib.sha256(rel.encode("utf-8", "ignore")).hexdigest()[:16]


def emit(cb: ProgressCallback, stage: str, task: str, file: str = "", percent: int = 0, done: int = 0, total: int = 0, started: Optional[float]=None):
    elapsed = int(time.time() - started) if started else 0
    eta = "--"
    finish_epoch = "--"
    if started and done and total and done > 0:
        rate = elapsed / max(done, 1)
        rem = max(0, int(rate * (total - done)))
        eta = rem
        finish_epoch = int(time.time() + rem)
    if cb:
        cb({"stage": stage, "task": task, "file": str(file), "percent": int(percent), "done": done, "total": total, "elapsed_seconds": elapsed, "eta_seconds": eta, "finish_epoch": finish_epoch})


def connect(db: Path) -> sqlite3.Connection:
    db.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(str(db), timeout=60)
    con.execute("PRAGMA foreign_keys=ON")
    con.execute("PRAGMA journal_mode=DELETE")
    con.execute("PRAGMA synchronous=NORMAL")
    return con


def vacuum_close(con: sqlite3.Connection):
    try:
        con.commit()
        con.execute("PRAGMA wal_checkpoint(TRUNCATE)")
    except Exception:
        pass
    try:
        con.execute("VACUUM")
    except Exception:
        pass
    con.close()


GENERIC_TABLE_COLUMNS = {
    "id": "TEXT",
    "source_id": "TEXT",
    "name": "TEXT",
    "path": "TEXT",
    "value": "TEXT",
    "metadata_json": "TEXT",
    "created_at": "TEXT",
}


def create_generic_table(con: sqlite3.Connection, table: str):
    if table.endswith("_fts"):
        con.execute(f"CREATE VIRTUAL TABLE IF NOT EXISTS {table} USING fts5(entity_id, text)")
    else:
        con.execute(f"""
        CREATE TABLE IF NOT EXISTS {table}(
            id TEXT PRIMARY KEY,
            source_id TEXT,
            name TEXT,
            path TEXT,
            value TEXT,
            metadata_json TEXT,
            created_at TEXT
        )
        """)
        ensure_columns(con, table, GENERIC_TABLE_COLUMNS)


def ensure_columns(conn, table_name: str, required_columns: dict[str, str]) -> None:
    cols = {r[1] for r in conn.execute(f'PRAGMA table_info("{table_name}")')}
    if not cols:
        return
    for name, col_type in required_columns.items():
        if name not in cols:
            conn.execute(f'ALTER TABLE "{table_name}" ADD COLUMN "{name}" {col_type}')
            cols.add(name)


def safe_export_name(brain_name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", str(brain_name or "brain")).strip("_") or "brain"


PROJECT_MUTATION_LANES = [
    ("intake", "Project mutation intake", "project_mutation_intake", "append public-model request packets"),
    ("task", "Project mutation task queue", "project_mutation_task", "split requests into explicit project work items"),
    ("patch", "Project patch proposal", "project_mutation_patch", "store proposed file/table changes and hash evidence"),
    ("validation", "Project validation gate", "project_mutation_validation", "record test, render, SQLite, and export gate results"),
    ("receipt", "Project mutation receipt", "project_mutation_receipt", "close the loop with changed files/tables and final status"),
    ("writeback", "Project model writeback", "project_model_writeback", "append structured model output to generated project truth"),
]

PROJECT_MUTATION_TABLE_NAMES = tuple(lane[2] for lane in PROJECT_MUTATION_LANES) + (
    "project_mutation_lane_registry",
    "project_public_ai_write_contract",
)


def ensure_project_mutation_schema(con: sqlite3.Connection) -> None:
    con.executescript("""
    CREATE TABLE IF NOT EXISTS project_mutation_lane_registry(
        lane_id TEXT PRIMARY KEY,
        lane_label TEXT NOT NULL,
        target_table TEXT NOT NULL,
        write_scope TEXT NOT NULL,
        lane_status TEXT NOT NULL,
        mutation_allowed INTEGER NOT NULL DEFAULT 1,
        notes TEXT,
        created_at TEXT
    );
    CREATE TABLE IF NOT EXISTS project_mutation_intake(
        request_id TEXT PRIMARY KEY,
        source_model TEXT,
        request_title TEXT,
        request_summary TEXT,
        requested_at TEXT,
        mutation_scope TEXT,
        target_lane TEXT,
        status TEXT,
        payload_json TEXT,
        created_at TEXT
    );
    CREATE TABLE IF NOT EXISTS project_mutation_task(
        task_id TEXT PRIMARY KEY,
        request_id TEXT,
        lane_id TEXT,
        target_path TEXT,
        target_table TEXT,
        action_type TEXT,
        status TEXT,
        priority INTEGER,
        payload_json TEXT,
        created_at TEXT
    );
    CREATE TABLE IF NOT EXISTS project_mutation_patch(
        patch_id TEXT PRIMARY KEY,
        task_id TEXT,
        target_path TEXT,
        target_table TEXT,
        patch_kind TEXT,
        before_hash TEXT,
        after_hash TEXT,
        patch_text TEXT,
        metadata_json TEXT,
        created_at TEXT
    );
    CREATE TABLE IF NOT EXISTS project_mutation_validation(
        validation_id TEXT PRIMARY KEY,
        patch_id TEXT,
        gate_name TEXT,
        status TEXT,
        details TEXT,
        evidence_path TEXT,
        created_at TEXT
    );
    CREATE TABLE IF NOT EXISTS project_mutation_receipt(
        receipt_id TEXT PRIMARY KEY,
        request_id TEXT,
        final_status TEXT,
        changed_files_json TEXT,
        changed_tables_json TEXT,
        receipt_path TEXT,
        created_at TEXT
    );
    CREATE TABLE IF NOT EXISTS project_model_writeback(
        writeback_id TEXT PRIMARY KEY,
        model_name TEXT,
        request_id TEXT,
        table_name TEXT,
        row_pk TEXT,
        write_policy TEXT,
        payload_json TEXT,
        created_at TEXT
    );
    CREATE TABLE IF NOT EXISTS project_public_ai_write_contract(
        contract_id TEXT PRIMARY KEY,
        allowed_table TEXT NOT NULL,
        allowed_action TEXT NOT NULL,
        denied_scope TEXT NOT NULL,
        required_gate TEXT NOT NULL,
        notes TEXT,
        created_at TEXT
    );
    """)
    for table_name, columns in {
        "project_mutation_lane_registry": {
            "lane_id": "TEXT", "lane_label": "TEXT", "target_table": "TEXT", "write_scope": "TEXT",
            "lane_status": "TEXT", "mutation_allowed": "INTEGER", "notes": "TEXT", "created_at": "TEXT",
        },
        "project_mutation_intake": {
            "request_id": "TEXT", "source_model": "TEXT", "request_title": "TEXT", "request_summary": "TEXT",
            "requested_at": "TEXT", "mutation_scope": "TEXT", "target_lane": "TEXT", "status": "TEXT",
            "payload_json": "TEXT", "created_at": "TEXT",
        },
        "project_mutation_task": {
            "task_id": "TEXT", "request_id": "TEXT", "lane_id": "TEXT", "target_path": "TEXT",
            "target_table": "TEXT", "action_type": "TEXT", "status": "TEXT", "priority": "INTEGER",
            "payload_json": "TEXT", "created_at": "TEXT",
        },
        "project_mutation_patch": {
            "patch_id": "TEXT", "task_id": "TEXT", "target_path": "TEXT", "target_table": "TEXT",
            "patch_kind": "TEXT", "before_hash": "TEXT", "after_hash": "TEXT", "patch_text": "TEXT",
            "metadata_json": "TEXT", "created_at": "TEXT",
        },
        "project_mutation_validation": {
            "validation_id": "TEXT", "patch_id": "TEXT", "gate_name": "TEXT", "status": "TEXT",
            "details": "TEXT", "evidence_path": "TEXT", "created_at": "TEXT",
        },
        "project_mutation_receipt": {
            "receipt_id": "TEXT", "request_id": "TEXT", "final_status": "TEXT",
            "changed_files_json": "TEXT", "changed_tables_json": "TEXT", "receipt_path": "TEXT",
            "created_at": "TEXT",
        },
        "project_model_writeback": {
            "writeback_id": "TEXT", "model_name": "TEXT", "request_id": "TEXT", "table_name": "TEXT",
            "row_pk": "TEXT", "write_policy": "TEXT", "payload_json": "TEXT", "created_at": "TEXT",
        },
        "project_public_ai_write_contract": {
            "contract_id": "TEXT", "allowed_table": "TEXT", "allowed_action": "TEXT", "denied_scope": "TEXT",
            "required_gate": "TEXT", "notes": "TEXT", "created_at": "TEXT",
        },
    }.items():
        ensure_columns(con, table_name, columns)
    created_at = now()
    for lane_id, lane_label, target_table, notes in PROJECT_MUTATION_LANES:
        con.execute(
            """
            INSERT OR IGNORE INTO project_mutation_lane_registry(
                lane_id, lane_label, target_table, write_scope, lane_status, mutation_allowed, notes, created_at
            ) VALUES(?,?,?,?,?,?,?,?)
            """,
            (lane_id, lane_label, target_table, "GENERATED_PROJECT_MUTABLE", "SCHEMA_READY", 1, notes, created_at),
        )
        con.execute(
            """
            INSERT OR IGNORE INTO project_public_ai_write_contract(
                contract_id, allowed_table, allowed_action, denied_scope, required_gate, notes, created_at
            ) VALUES(?,?,?,?,?,?,?)
            """,
            (
                "contract_" + lane_id,
                target_table,
                "INSERT_OR_APPEND_ONLY",
                "ENV_UOP_AND_PUBLIC_TEMPLATE_LOCKED",
                "PROJECT_MUTATION_REQUEST",
                "Public models may write structured rows here when the user explicitly requests a project mutation.",
                created_at,
            ),
        )



def _migrate_brain_manifest_legacy(con):
    cols = [r[1] for r in con.execute("PRAGMA table_info(brain_manifest)")]
    if cols and "builder_version" not in cols:
        con.execute("ALTER TABLE brain_manifest ADD COLUMN builder_version TEXT")
        con.commit()

def ensure_router(router: Path, brain_name: str):
    con = connect(router)
    con.executescript("""
    CREATE TABLE IF NOT EXISTS brain_manifest(brain_id TEXT PRIMARY KEY, brain_name TEXT, brain_slug TEXT, created_at TEXT, status TEXT, builder_version TEXT);
    CREATE TABLE IF NOT EXISTS sector_registry(sector_id TEXT PRIMARY KEY, lane_key TEXT UNIQUE, lane_label TEXT, sector_db_path TEXT, active_bool INTEGER, version TEXT, sector_hash TEXT, created_at TEXT);
    CREATE TABLE IF NOT EXISTS source_registry(source_id TEXT PRIMARY KEY, lane_key TEXT, lane_label TEXT, source_type TEXT, display_name TEXT, path TEXT, source_hash TEXT, active_bool INTEGER, created_at TEXT);
    CREATE TABLE IF NOT EXISTS package_manifest(package_id TEXT PRIMARY KEY, package_path TEXT, created_at TEXT, package_hash TEXT);
    CREATE TABLE IF NOT EXISTS active_export_manifest(id TEXT PRIMARY KEY, lane_key TEXT, sector_db_path TEXT, active_bool INTEGER, created_at TEXT);
    """)
    ensure_columns(con, "brain_manifest", {
        "brain_id": "TEXT",
        "brain_name": "TEXT",
        "brain_slug": "TEXT",
        "created_at": "TEXT",
        "status": "TEXT",
        "builder_version": "TEXT",
    })
    ensure_columns(con, "sector_registry", {
        "sector_id": "TEXT",
        "lane_key": "TEXT",
        "lane_label": "TEXT",
        "sector_db_path": "TEXT",
        "active_bool": "INTEGER",
        "version": "TEXT",
        "sector_hash": "TEXT",
        "created_at": "TEXT",
    })
    ensure_columns(con, "source_registry", {
        "source_id": "TEXT",
        "lane_key": "TEXT",
        "lane_label": "TEXT",
        "source_type": "TEXT",
        "display_name": "TEXT",
        "path": "TEXT",
        "source_hash": "TEXT",
        "active_bool": "INTEGER",
        "created_at": "TEXT",
    })
    ensure_project_mutation_schema(con)
    con.execute(
        "INSERT OR REPLACE INTO brain_manifest(brain_id, brain_name, brain_slug, created_at, status, builder_version) VALUES(?,?,?,?,?,?)",
        ("brain_" + slugify_name(brain_name), brain_name, slugify_name(brain_name), now(), "ACTIVE", APP_VERSION),
    )
    con.commit(); con.close()


def ensure_sector(db: Path, lane_key: str):
    lane_key = resolve_lane_id(lane_key, scope="any")
    con = connect(db)
    con.execute("CREATE TABLE IF NOT EXISTS sector_manifest(sector_id TEXT PRIMARY KEY, lane_key TEXT, lane_label TEXT, version TEXT, status TEXT, schema_tables_json TEXT, fill_policy TEXT, created_at TEXT)")
    con.execute("CREATE TABLE IF NOT EXISTS sector_pointer(pointer_id TEXT PRIMARY KEY, lane_key TEXT, sector_db_path TEXT, active_bool INTEGER, status TEXT, mmd_required INTEGER, created_at TEXT)")
    con.execute("CREATE TABLE IF NOT EXISTS source_active_state(source_id TEXT PRIMARY KEY, active_bool INTEGER, lane_key TEXT, display_name TEXT, reason TEXT, created_at TEXT)")
    con.execute("""
    CREATE TABLE IF NOT EXISTS ingestion_state(
        source_id TEXT PRIMARY KEY,
        source_hash TEXT NOT NULL,
        parser_id TEXT NOT NULL,
        chunker_version TEXT NOT NULL,
        completed_at TEXT NOT NULL
    )
    """)
    ensure_columns(con, "sector_manifest", {
        "sector_id": "TEXT",
        "lane_key": "TEXT",
        "lane_label": "TEXT",
        "version": "TEXT",
        "status": "TEXT",
        "schema_tables_json": "TEXT",
        "fill_policy": "TEXT",
        "created_at": "TEXT",
    })
    ensure_columns(con, "sector_pointer", {
        "pointer_id": "TEXT",
        "lane_key": "TEXT",
        "sector_db_path": "TEXT",
        "active_bool": "INTEGER",
        "status": "TEXT",
        "mmd_required": "INTEGER",
        "created_at": "TEXT",
    })
    if lane_key in {"local_code", "github_code"}:
        create_code_schema(con)
    elif lane_key not in STRUCTURAL_LANE_INGESTERS:
        for table in LANE_DEFS.get(lane_key, LANE_DEFS["custom"])["schema"]:
            create_generic_table(con, table)
        if lane_key == "chat_lineage":
            # Read compatibility for pre-T021 brains.  New identity and writes
            # remain canonical; this table prevents older consumers from
            # breaking while state-travel migration is completed.
            create_generic_table(con, "lineage_source")
    con.execute(
        "INSERT OR REPLACE INTO sector_manifest(sector_id, lane_key, lane_label, version, status, schema_tables_json, fill_policy, created_at) VALUES(?,?,?,?,?,?,?,?)",
        ("sector_"+lane_key, lane_key, LANE_DEFS.get(lane_key, {}).get("label", lane_key), "v001", "SCHEMA_READY", json.dumps(LANE_DEFS.get(lane_key, LANE_DEFS["custom"])["schema"]), "Unlocked sector DB fillable only by explicit user command; env/uop/template remain locked.", now()),
    )
    con.execute(
        "INSERT OR REPLACE INTO sector_pointer(pointer_id, lane_key, sector_db_path, active_bool, status, mmd_required, created_at) VALUES(?,?,?,?,?,?,?)",
        ("pointer_"+lane_key, lane_key, str(db), 1, "SCHEMA_READY", 1 if lane_key in ("local_code","github_code") else 0, now()),
    )
    con.commit(); con.close()


def init_brain_layout(workspace_dir: str, brain_name: str):
    root = brain_output_dir(workspace_dir, brain_name)
    initialize_env15_brain_root(root)
    return root


def finalize_project_pointers_and_hashes(
    root: Path,
    *,
    active_lane_ids: list[str] | tuple[str, ...] = (),
    ingestion_results: list[dict] | tuple[dict, ...] = (),
    brain_name: str | None = None,
    package_use_mode: str = "CANONICAL_FLASHABLE",
):
    """Refresh every canonical sector hash and its pointer before packaging."""

    if (root / ".uepc_env").is_file():
        runtime_state = write_env15_runtime_sector_state(root)
        universal = write_universal_lane_authority(
            root,
            active_lane_ids=active_lane_ids,
            ingestion_results=ingestion_results,
            brain_name=brain_name,
            package_use_mode=package_use_mode,
        )
        return {"runtime_state": runtime_state, "universal_lane_authority": universal}

    project = root / "project"
    router = project / "project_router.sqlite"
    index = []
    con = connect(router)
    try:
        for lane_key, lane in LANE_DEFS.items():
            db = project / "sectors" / lane["sector_folder"] / lane["sqlite_filename"]
            if not db.is_file():
                raise RuntimeError(f"CANONICAL_SECTOR_MISSING:{lane_key}:{db}")
            digest = sha256_file(db)
            con.execute(
                "UPDATE sector_registry SET sector_hash=?, sector_db_path=?, active_bool=1 WHERE lane_key=?",
                (digest, str(db), lane_key),
            )
            pointer = {
                "lane_key": lane_key,
                "lane_label": lane["label"],
                "sector_db": f"project/sectors/{lane['sector_folder']}/{lane['sqlite_filename']}",
                "sector_sha256": digest,
                "active": True,
                "schema_tables": lane["schema"],
                "mmd_required": lane_key in {"local_code", "github_code"},
                "status": "FINALIZED",
                "mutation_policy": lane["mutation_policy"],
            }
            (project / "pointers" / f"{lane_key}_pointer.json").write_text(
                json.dumps(pointer, indent=2), encoding="utf-8"
            )
            index.append(pointer)
        con.commit()
    finally:
        con.close()
    (project / "sector_index.json").write_text(json.dumps(index, indent=2), encoding="utf-8")
    project_pointer = {
        "router": "project/project_router.sqlite",
        "router_sha256": sha256_file(router),
        "sectors": "project/sectors/",
        "pointers": "project/pointers/",
        "sector_count": len(index),
        "canonical_lane_ids": list(LANE_DEFS),
        "status": "FINALIZED",
        "created_at": now(),
    }
    (project / "project_pointer.json").write_text(
        json.dumps(project_pointer, indent=2), encoding="utf-8"
    )
    return project_pointer


def detect_language(ext: str, path: str) -> str:
    return {".py":"python",".js":"javascript",".jsx":"jsx",".ts":"typescript",".tsx":"tsx",".css":"css",".scss":"scss",".html":"html",".sql":"sql",".md":"markdown",".txt":"text",".json":"json",".yml":"yaml",".yaml":"yaml",".ps1":"powershell"}.get(ext.lower(), ext.lower().strip(".") or "unknown")


def read_text_lossless(path: Path):
    data = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(enc), enc, data
        except Exception:
            continue
    return data.decode("utf-8", "replace"), "utf-8-replace", data


def create_code_schema(con: sqlite3.Connection):
    # The generic lane shell creates placeholder tables on first use. Replace only
    # those incompatible placeholders; a real code sector is migrated in place and
    # is never dropped/re-indexed on a later build.
    required_signatures = {
        "source_file": {"file_id", "logical_path", "sha256"},
        "source_byte_coverage": {"file_id", "coverage_status", "sha256"},
        "code_repo": {"id", "source_id", "metadata_json"},
        "code_folder": {"folder_id", "normalized_path", "source_id"},
        "git_remote": {"id", "source_id", "metadata_json"},
        "git_branch": {"id", "source_id", "metadata_json"},
        "git_commit": {"commit_sha", "commit_time", "commit_order"},
        "git_commit_parent_edge": {"edge_id", "commit_sha", "parent_sha"},
        "git_file_change": {"id", "source_id", "metadata_json"},
        "git_diff_hunk": {"hunk_id", "commit_sha", "old_path", "new_path"},
        "git_line_change": {"line_change_id", "hunk_id", "commit_sha", "path"},
        "git_rename_map": {"rename_id", "commit_sha", "old_path", "new_path"},
        "code_file": {"file_id", "canonical_path", "current_sha256"},
        "code_file_version": {"file_version_id", "file_id", "raw_file_sha256"},
        "code_line_snapshot": {"line_id", "file_version_id", "file_id"},
        "code_chunk": {"chunk_id", "file_version_id", "file_id"},
        "code_symbol": {"symbol_id", "file_version_id", "file_id"},
        "app_route": {"route_id", "route_path", "file_id"},
        "dependency_manifest": {"manifest_id", "file_id", "sha256"},
        "dependency_item": {"dependency_id", "manifest_id", "package_name"},
        "code_import_edge": {"edge_id", "from_file_id", "import_target"},
        "project_artifact": {"artifact_id", "source_file_id", "artifact_sha256"},
        "artifact_relation_edge": {"edge_id", "artifact_id", "related_entity_id"},
    }
    for table_name, required in required_signatures.items():
        existing = {row[1] for row in con.execute(f'PRAGMA table_info("{table_name}")')}
        if existing and not required.issubset(existing):
            con.execute(f'DROP TABLE "{table_name}"')
    for fts_name in ("code_fts", "line_fts", "symbol_fts", "route_fts", "commit_fts", "artifact_fts"):
        row = con.execute(
            "SELECT sql FROM sqlite_master WHERE type='table' AND name=?",
            (fts_name,),
        ).fetchone()
        if row and "VIRTUAL TABLE" not in str(row[0] or "").upper():
            con.execute(f'DROP TABLE "{fts_name}"')
    con.executescript("""
    CREATE TABLE IF NOT EXISTS source_file(file_id TEXT PRIMARY KEY, logical_path TEXT, extension TEXT, size_bytes INTEGER, sha256 TEXT, lane_status TEXT, created_at TEXT);
    CREATE TABLE IF NOT EXISTS source_byte_coverage(file_id TEXT PRIMARY KEY, coverage_status TEXT, byte_count INTEGER, sha256 TEXT, notes TEXT, created_at TEXT);
    CREATE TABLE IF NOT EXISTS code_repo(id TEXT PRIMARY KEY, source_id TEXT, name TEXT, path TEXT, value TEXT, metadata_json TEXT, created_at TEXT);
    CREATE TABLE IF NOT EXISTS code_folder(folder_id TEXT PRIMARY KEY, source_id TEXT, raw_path TEXT, normalized_path TEXT, parent_folder_id TEXT, child_folder_count INTEGER, child_file_count INTEGER, languages_json TEXT, role_signals_json TEXT, created_at TEXT);
    CREATE TABLE IF NOT EXISTS git_remote(id TEXT PRIMARY KEY, source_id TEXT, name TEXT, path TEXT, value TEXT, metadata_json TEXT, created_at TEXT);
    CREATE TABLE IF NOT EXISTS git_branch(id TEXT PRIMARY KEY, source_id TEXT, name TEXT, path TEXT, value TEXT, metadata_json TEXT, created_at TEXT);
    CREATE TABLE IF NOT EXISTS git_commit(commit_sha TEXT PRIMARY KEY, short_sha TEXT, author_name TEXT, author_email_hash TEXT, author_time TEXT, committer_name TEXT, committer_email_hash TEXT, committer_time TEXT, commit_time TEXT, message TEXT, commit_order INTEGER);
    CREATE TABLE IF NOT EXISTS git_commit_parent_edge(edge_id TEXT PRIMARY KEY, commit_sha TEXT, parent_sha TEXT, parent_order INTEGER);
    CREATE TABLE IF NOT EXISTS git_file_change(id TEXT PRIMARY KEY, source_id TEXT, name TEXT, path TEXT, value TEXT, metadata_json TEXT, created_at TEXT);
    CREATE TABLE IF NOT EXISTS git_diff_hunk(hunk_id TEXT PRIMARY KEY, commit_sha TEXT, old_path TEXT, new_path TEXT, old_start INTEGER, old_count INTEGER, new_start INTEGER, new_count INTEGER, hunk_header TEXT, patch_sha256 TEXT, history_policy TEXT, changed_line_count INTEGER, stored_line_count INTEGER);
    CREATE TABLE IF NOT EXISTS git_line_change(line_change_id TEXT PRIMARY KEY, hunk_id TEXT, commit_sha TEXT, path TEXT, old_line_number INTEGER, new_line_number INTEGER, change_type TEXT, line_text TEXT, line_sha256 TEXT);
    CREATE TABLE IF NOT EXISTS git_rename_map(rename_id TEXT PRIMARY KEY, commit_sha TEXT, old_path TEXT, new_path TEXT, similarity INTEGER);
    CREATE TABLE IF NOT EXISTS code_file(file_id TEXT PRIMARY KEY, canonical_path TEXT, language TEXT, extension TEXT, current_sha256 TEXT, is_active INTEGER);
    CREATE TABLE IF NOT EXISTS code_file_version(file_version_id TEXT PRIMARY KEY, file_id TEXT, commit_sha TEXT, raw_file_sha256 TEXT, normalized_text_sha256 TEXT, path_at_commit TEXT, language TEXT, extension TEXT, line_count INTEGER, byte_count INTEGER, is_deleted INTEGER, is_renamed INTEGER, created_at TEXT);
    CREATE TABLE IF NOT EXISTS code_line_snapshot(line_id TEXT PRIMARY KEY, file_version_id TEXT, file_id TEXT, line_number INTEGER, line_text TEXT, line_sha256 TEXT, normalized_line_sha256 TEXT, indent_level INTEGER, is_blank INTEGER, is_comment INTEGER, search_text TEXT);
    CREATE TABLE IF NOT EXISTS code_chunk(chunk_id TEXT PRIMARY KEY, file_version_id TEXT, file_id TEXT, chunk_type TEXT, language TEXT, role_name TEXT, start_line INTEGER, end_line INTEGER, chunk_text TEXT, chunk_sha256 TEXT);
    CREATE TABLE IF NOT EXISTS code_symbol(symbol_id TEXT PRIMARY KEY, symbol_name TEXT, symbol_type TEXT, file_version_id TEXT, file_id TEXT, language TEXT, start_line INTEGER, end_line INTEGER, parent_symbol_id TEXT, signature TEXT, symbol_sha256 TEXT);
    CREATE TABLE IF NOT EXISTS app_route(route_id TEXT PRIMARY KEY, route_path TEXT, route_type TEXT, file_id TEXT, file_version_id TEXT, route_sha256 TEXT);
    CREATE TABLE IF NOT EXISTS dependency_manifest(manifest_id TEXT PRIMARY KEY, file_id TEXT, manifest_type TEXT, ecosystem TEXT, path TEXT, sha256 TEXT);
    CREATE TABLE IF NOT EXISTS dependency_item(dependency_id TEXT PRIMARY KEY, manifest_id TEXT, package_name TEXT, version_spec TEXT, ecosystem TEXT);
    CREATE TABLE IF NOT EXISTS code_import_edge(edge_id TEXT PRIMARY KEY, from_file_id TEXT, from_path TEXT, import_target TEXT, import_type TEXT, line_number INTEGER);
    CREATE TABLE IF NOT EXISTS code_dependency_edge(id TEXT PRIMARY KEY, source_id TEXT, name TEXT, path TEXT, value TEXT, metadata_json TEXT, created_at TEXT);
    CREATE TABLE IF NOT EXISTS project_artifact(artifact_id TEXT PRIMARY KEY, source_file_id TEXT, artifact_type TEXT, artifact_sha256 TEXT, path TEXT, size_bytes INTEGER, metadata_json TEXT, semantic_status TEXT);
    CREATE TABLE IF NOT EXISTS artifact_relation_edge(edge_id TEXT PRIMARY KEY, artifact_id TEXT, related_entity_type TEXT, related_entity_id TEXT, relation_type TEXT, confidence TEXT);
    CREATE TABLE IF NOT EXISTS code_index_state(repo_id TEXT, commit_sha TEXT, policy_version TEXT, indexed_at TEXT, PRIMARY KEY(repo_id, commit_sha));
    CREATE TABLE IF NOT EXISTS workflow_node(id TEXT PRIMARY KEY, source_id TEXT, name TEXT, path TEXT, value TEXT, metadata_json TEXT, created_at TEXT);
    CREATE TABLE IF NOT EXISTS workflow_edge(id TEXT PRIMARY KEY, source_id TEXT, name TEXT, path TEXT, value TEXT, metadata_json TEXT, created_at TEXT);
    CREATE VIRTUAL TABLE IF NOT EXISTS code_fts USING fts5(entity_id, text);
    CREATE VIRTUAL TABLE IF NOT EXISTS line_fts USING fts5(entity_id, text);
    CREATE VIRTUAL TABLE IF NOT EXISTS symbol_fts USING fts5(entity_id, text);
    CREATE VIRTUAL TABLE IF NOT EXISTS route_fts USING fts5(entity_id, text);
    CREATE VIRTUAL TABLE IF NOT EXISTS commit_fts USING fts5(entity_id, text);
    CREATE VIRTUAL TABLE IF NOT EXISTS artifact_fts USING fts5(entity_id, text);
    """)
    ensure_columns(con, "git_diff_hunk", {
        "history_policy": "TEXT", "changed_line_count": "INTEGER", "stored_line_count": "INTEGER",
    })
    ensure_columns(con, "git_commit", {
        "author_time": "TEXT", "committer_name": "TEXT",
        "committer_email_hash": "TEXT", "committer_time": "TEXT",
    })


def create_code_query_indexes(con: sqlite3.Connection) -> None:
    # Build these only after tiered history pruning has completed.
    con.executescript("""
    CREATE INDEX IF NOT EXISTS idx_git_file_change_commit_path ON git_file_change(value, path);
    CREATE INDEX IF NOT EXISTS idx_git_diff_hunk_commit_new_path ON git_diff_hunk(commit_sha, new_path);
    CREATE INDEX IF NOT EXISTS idx_git_diff_hunk_commit_old_path ON git_diff_hunk(commit_sha, old_path);
    CREATE INDEX IF NOT EXISTS idx_git_line_change_hunk ON git_line_change(hunk_id);
    CREATE INDEX IF NOT EXISTS idx_git_line_change_commit_path ON git_line_change(commit_sha, path);
    CREATE INDEX IF NOT EXISTS idx_git_line_change_path ON git_line_change(path);
    CREATE INDEX IF NOT EXISTS idx_git_parent_commit ON git_commit_parent_edge(commit_sha);
    CREATE INDEX IF NOT EXISTS idx_git_parent_parent ON git_commit_parent_edge(parent_sha);
    CREATE INDEX IF NOT EXISTS idx_code_file_version_file_commit ON code_file_version(file_id, commit_sha);
    CREATE INDEX IF NOT EXISTS idx_code_line_snapshot_version_line ON code_line_snapshot(file_version_id, line_number);
    CREATE INDEX IF NOT EXISTS idx_code_chunk_version_start ON code_chunk(file_version_id, start_line);
    CREATE INDEX IF NOT EXISTS idx_code_index_state_repo_commit ON code_index_state(repo_id, commit_sha);
    """)


def is_ignored(rel: Path) -> bool:
    return any(part in HEAVY_SKIP_DIRS for part in rel.parts)


def is_dependency_manifest_file(path: Path) -> bool:
    name = path.name.lower()
    return name in DEPENDENCY_MANIFEST_NAMES or name.endswith((".config.js", ".config.ts", ".config.mjs", ".config.cjs"))


def classify_code_project_file(path: Path, rel: Path):
    ext = path.suffix.lower()
    name = path.name.lower()
    size = path.stat().st_size
    rel_parts = {part.lower() for part in rel.parts}
    if ext in BINARY_ASSET_EXTS:
        return "CODE_MEDIA_ASSET_HASH_ONLY"
    if rel_parts.intersection(GENERATED_HISTORY_DIRS):
        if size > MAX_ARTIFACT_TEXT_EXTRACT:
            return "CODE_ARTIFACT_METADATA_ONLY"
        return "CODE_ARTIFACT_STUDIED_READ_ONLY"
    if ext in CODE_EXTS or name in SPECIAL_CODE_FILENAMES:
        if size > 5_000_000:
            return "CODE_TEXT_TOO_LARGE_ARTIFACT_READ_ONLY"
        return "LOSSLESS_CODE_LAYERED_CHUNKED"
    if is_dependency_manifest_file(path) or ext in TEXT_ARTIFACT_EXTS or ext in CODE_DOC_EXTS:
        if size > MAX_ARTIFACT_TEXT_EXTRACT:
            return "CODE_ARTIFACT_METADATA_ONLY"
        return "CODE_ARTIFACT_STUDIED_READ_ONLY"
    if ext in {".parquet", ".xlsx", ".xls", ".pdf", ".zip"}:
        return "CODE_ARTIFACT_METADATA_ONLY"
    return "UNSUPPORTED_HASH_ONLY"


def classify_code_role(rel: str, ext: str, text: str) -> str:
    p = rel.replace("\\", "/").lower()
    name = Path(p).name
    t = (text or "")[:8000].lower()
    if name in DEPENDENCY_MANIFEST_NAMES or name in SPECIAL_CODE_FILENAMES or ext in {".toml", ".yaml", ".yml", ".ini", ".cfg"}:
        return "CONFIG_BUILD_TOOLING"
    path_parts = [part for part in p.strip("/").split("/") if part]
    test_directory = any(part in {"test", "tests", "spec", "specs", "__tests__"} for part in path_parts[:-1])
    test_filename = (
        Path(name).stem.startswith(("test_", "spec_"))
        or Path(name).stem.endswith(("_test", "_tests", "_spec", "_specs", ".test", ".spec"))
        or p.endswith((".test.py", ".spec.py", ".test.ts", ".spec.ts", ".test.tsx", ".spec.tsx", ".test.js", ".spec.js"))
    )
    if test_directory or test_filename:
        return "TEST_QA"
    if p.startswith("src/app/") or "/pages/" in p:
        if "/api/" in p or name in {"route.ts", "route.js", "route.py"}:
            return "API_BACKEND_ROUTE"
        return "UI_PAGE_ROUTE"
    if "/components/" in p or "/ui/" in p or ext in {".tsx", ".jsx"}:
        if any(token in t for token in ("use client", "onclick", "classname", "<button", "return <")):
            return "UI_UX_COMPONENT"
    if "/api/" in p or "/server/" in p or "fastapi" in t or "flask" in t or "express()" in t:
        return "BACKEND_PROCESS"
    if "/lib/" in p or "/utils/" in p or "/services/" in p or "/core/" in p:
        return "SERVICE_OR_UTILITY"
    if "/data/" in p or "/db/" in p or ext == ".sql" or "sqlite" in t or "select " in t:
        return "DATA_DB_LAYER"
    if ext in {".md", ".txt"}:
        return "DOCS_OR_NOTES"
    if ext in {".css", ".scss"}:
        return "UI_STYLE_THEME"
    return "GENERAL_CODE"


def code_chunk_preview(text: str) -> str:
    if len(text) <= MAX_CODE_CHUNK_TEXT:
        return text
    return text[:MAX_CODE_CHUNK_TEXT] + "\n...[chunk truncated for compact SQLite storage]"


def insert_code_chunk(con: sqlite3.Connection, chunk_id: str, file_version_id: str, file_id: str, chunk_type: str, language: str, role_name: str, start_line: int, end_line: int, chunk_text: str, *, replace_fts: bool = True) -> None:
    preview = code_chunk_preview(chunk_text)
    csha = sha256_bytes(chunk_text.encode("utf-8", "ignore"))
    con.execute(
        """INSERT OR REPLACE INTO code_chunk(
            chunk_id, file_version_id, file_id, chunk_type, language, role_name,
            start_line, end_line, chunk_text, chunk_sha256
        ) VALUES(?,?,?,?,?,?,?,?,?,?)""",
        (chunk_id, file_version_id, file_id, chunk_type, language, role_name, start_line, end_line, preview, csha),
    )
    if preview.strip():
        if replace_fts:
            con.execute("DELETE FROM code_fts WHERE entity_id=?", (chunk_id,))
        con.execute("INSERT INTO code_fts(entity_id,text) VALUES(?,?)", (chunk_id, preview))


def folder_id_for(rel: str) -> str:
    return "folder_" + sha256_bytes(rel.encode("utf-8", "ignore"))[:16]


def insert_folder_chunks(con: sqlite3.Connection, folder: Path, files: list[Path], source_id: str, *, replace_fts: bool = True) -> None:
    directory_paths = {Path(".")}
    for candidate in folder.rglob("*"):
        rel_candidate = candidate.relative_to(folder)
        if candidate.is_dir() and not is_ignored(rel_candidate):
            directory_paths.add(rel_candidate)
    for file in files:
        rel_path = file.relative_to(folder)
        directory_paths.add(rel_path.parent)

    folder_map: dict[str, dict[str, set[str] | int | dict[str, int]]] = {}
    for rel_dir in directory_paths:
        normalized = rel_dir.as_posix() or "."
        folder_map.setdefault(normalized, {"files": set(), "folders": set(), "languages": {}, "file_count": 0})
        if normalized != ".":
            parent = rel_dir.parent.as_posix() or "."
            parent_info = folder_map.setdefault(parent, {"files": set(), "folders": set(), "languages": {}, "file_count": 0})
            parent_info["folders"].add(rel_dir.name)  # type: ignore[index]

    for file in files:
        rel_path = file.relative_to(folder)
        rel_dir = rel_path.parent.as_posix() or "."
        ext = file.suffix.lower()
        language = detect_language(ext, file.as_posix())
        info = folder_map.setdefault(rel_dir, {"files": set(), "folders": set(), "languages": {}, "file_count": 0})
        info["files"].add(rel_path.name)  # type: ignore[index]
        info["file_count"] = int(info["file_count"]) + 1
        langs = info["languages"]  # type: ignore[assignment]
        langs[language] = langs.get(language, 0) + 1

    for folder_rel, info in sorted(folder_map.items()):
        raw_rel = "." if folder_rel == "." else str(Path(folder_rel))
        languages = dict(sorted((info["languages"]).items()))  # type: ignore[union-attr]
        parent_folder_id = None if folder_rel == "." else folder_id_for(Path(folder_rel).parent.as_posix() or ".")
        payload = {
            "folder_path": raw_rel,
            "normalized_path": folder_rel,
            "parent_folder_id": parent_folder_id,
            "child_files": sorted(info["files"])[:200],  # type: ignore[index]
            "child_folders": sorted(info["folders"])[:200],  # type: ignore[index]
            "file_count": int(info["file_count"]),
            "languages": languages,
            "role_signals": sorted(k for k, v in languages.items() if v),
        }
        node_id = folder_id_for(folder_rel)
        text = json.dumps(payload, ensure_ascii=False, sort_keys=True)
        con.execute(
            """INSERT OR REPLACE INTO code_folder(
                folder_id, source_id, raw_path, normalized_path, parent_folder_id,
                child_folder_count, child_file_count, languages_json, role_signals_json, created_at
            ) VALUES(?,?,?,?,?,?,?,?,?,?)""",
            (
                node_id,
                source_id,
                raw_rel,
                folder_rel,
                parent_folder_id,
                len(info["folders"]),  # type: ignore[arg-type]
                int(info["file_count"]),
                json.dumps(languages, sort_keys=True),
                json.dumps(payload["role_signals"], sort_keys=True),
                now(),
            ),
        )
        con.execute(
            """INSERT OR REPLACE INTO workflow_node(
                id, source_id, name, path, value, metadata_json, created_at
            ) VALUES(?,?,?,?,?,?,?)""",
            (node_id, source_id, "FOLDER_LAYOUT", folder_rel, text, json.dumps({"chunk_type": "FOLDER_CHUNK"}, sort_keys=True), now()),
        )
        insert_code_chunk(con, "chunk_" + node_id, "", node_id, "FOLDER_CHUNK", "folder", "PROJECT_FOLDER_LAYOUT", 0, 0, text, replace_fts=replace_fts)


def insert_source_registry(router: Path, source: dict, lane_key: str, digest=""):
    con = connect(router)
    lane = LANE_DEFS.get(lane_key, LANE_DEFS["custom"])
    sid = source.get("source_id") or "source_" + hashlib.sha256((source.get("path") or source.get("display_name") or lane_key).encode()).hexdigest()[:12]
    ensure_columns(con, "source_registry", {
        "source_id": "TEXT",
        "lane_key": "TEXT",
        "lane_label": "TEXT",
        "source_type": "TEXT",
        "display_name": "TEXT",
        "path": "TEXT",
        "source_hash": "TEXT",
        "active_bool": "INTEGER",
        "created_at": "TEXT",
    })
    con.execute(
        "INSERT OR REPLACE INTO source_registry(source_id, lane_key, lane_label, source_type, display_name, path, source_hash, active_bool, created_at) VALUES(?,?,?,?,?,?,?,?,?)",
        (sid, lane_key, lane["label"], source.get("source_type",""), source.get("display_name") or source.get("path") or lane["label"], source.get("path",""), digest, 1 if source.get("active", True) else 0, now()),
    )
    con.commit(); con.close()


def parse_imports(text: str, language: str):
    out = []
    for i, line in enumerate(text.splitlines(), start=1):
        s = line.strip()
        target = None; typ = None
        if language == "python":
            m = re.match(r"import\s+([A-Za-z0-9_\.]+)", s)
            if m: target=m.group(1); typ="python_import"
            m = re.match(r"from\s+([A-Za-z0-9_\.]+)\s+import", s)
            if m: target=m.group(1); typ="python_from_import"
        elif language in {"javascript","typescript","tsx","jsx"}:
            m = re.search(r"from\s+['\"]([^'\"]+)", s)
            if m: target=m.group(1); typ="js_import"
            m = re.match(r"import\(['\"]([^'\"]+)", s)
            if m: target=m.group(1); typ="js_dynamic_import"
            m = re.search(r"require\(['\"]([^'\"]+)['\"]\)", s)
            if m: target=m.group(1); typ="commonjs_require"
            navigation = re.search(
                r"(?:window\.)?location(?:\.href)?\s*=\s*['\"]([^'\"]+)['\"]"
                r"|(?:window\.)?location\.(?:assign|replace)\(\s*['\"]([^'\"]+)['\"]",
                s,
            )
            if navigation:
                out.append((i, navigation.group(1) or navigation.group(2), "js_navigation"))
        if target:
            out.append((i, target, typ))
    if language == "html":
        tag_contracts = (
            ("script", "src", "html_script"),
            ("link", "href", "html_stylesheet"),
            ("a", "href", "html_navigation"),
            ("area", "href", "html_navigation"),
            ("img", "src", "html_asset"),
            ("source", "src", "html_asset"),
            ("video", "src", "html_asset"),
            ("audio", "src", "html_asset"),
            ("form", "action", "html_form_action"),
        )
        for tag, attribute, import_type in tag_contracts:
            pattern = re.compile(
                rf"<{tag}\b[^>]*?\b{attribute}\s*=\s*(?:['\"]([^'\"]+)['\"]|([^\s>]+))",
                flags=re.IGNORECASE,
            )
            for match in pattern.finditer(text):
                target = (match.group(1) or match.group(2) or "").strip()
                if target and not target.casefold().startswith(("#", "data:", "javascript:", "mailto:", "tel:")):
                    out.append((text.count("\n", 0, match.start()) + 1, target, import_type))
    elif language in {"css", "scss"}:
        for match in re.finditer(r"url\(\s*(?:['\"]([^'\"]+)['\"]|([^)'\"\s]+))\s*\)", text, flags=re.IGNORECASE):
            target = (match.group(1) or match.group(2) or "").strip()
            if target and not target.casefold().startswith(("data:", "#")):
                out.append((text.count("\n", 0, match.start()) + 1, target, "css_asset"))
        for match in re.finditer(r"@import\s+(?:url\(\s*)?['\"]([^'\"]+)['\"]", text, flags=re.IGNORECASE):
            out.append((text.count("\n", 0, match.start()) + 1, match.group(1).strip(), "css_import"))
    deduped = []
    seen = set()
    for item in out:
        if item not in seen:
            seen.add(item)
            deduped.append(item)
    return deduped


def _line_indent(line: str) -> int:
    return len(line) - len(line.lstrip(" \t"))


def detect_symbols(text: str, language: str):
    out = []
    patterns = []
    if language == "python":
        patterns = [(r"^\s*def\s+([A-Za-z_][A-Za-z0-9_]*)\s*\(", "FUNCTION"), (r"^\s*class\s+([A-Za-z_][A-Za-z0-9_]*)", "CLASS")]
    elif language in {"typescript","tsx","javascript","jsx"}:
        patterns = [(r"function\s+([A-Za-z_][A-Za-z0-9_]*)\s*\(", "FUNCTION"), (r"(?:const|let|var)\s+([A-Za-z_][A-Za-z0-9_]*)\s*=\s*(?:\([^)]*\)\s*=>|function)", "FUNCTION_OR_COMPONENT"), (r"export\s+default\s+function\s+([A-Za-z_][A-Za-z0-9_]*)?", "REACT_COMPONENT")]
    for i, line in enumerate(text.splitlines(), start=1):
        for pat, typ in patterns:
            m = re.search(pat, line)
            if m:
                out.append((i, m.group(1) or "default", typ, line.strip()[:300]))
    return out


def detect_symbol_ranges(text: str, language: str):
    lines = text.splitlines()
    starts = []
    for line_no, name, typ, sig in detect_symbols(text, language):
        indent = _line_indent(lines[line_no - 1]) if 0 < line_no <= len(lines) else 0
        starts.append({"start": line_no, "end": len(lines), "name": name, "type": typ, "signature": sig, "indent": indent})
    if language == "python":
        for item in starts:
            for line_no in range(int(item["start"]) + 1, len(lines) + 1):
                line = lines[line_no - 1]
                stripped = line.strip()
                if stripped and not stripped.startswith("#") and _line_indent(line) <= int(item["indent"]):
                    item["end"] = max(int(item["start"]), line_no - 1)
                    break
    else:
        for idx, item in enumerate(starts):
            next_start = starts[idx + 1]["start"] if idx + 1 < len(starts) else len(lines) + 1
            item["end"] = max(item["start"], int(next_start) - 1)
    return starts


def detect_route(rel: str, ext: str):
    p = rel.replace("\\", "/")
    if ext.casefold() in {".html", ".htm"}:
        route = "/" + p.lstrip("/")
        if re.search(r"(?:^|/)index\.html?$", p, flags=re.IGNORECASE):
            parent = p.rsplit("/", 1)[0] if "/" in p else ""
            route = f"/{parent}/" if parent else "/"
        return route, "STATIC_HTML_PAGE"
    if re.search(r"src/app/(.*)/page\.tsx$", p) or p == "src/app/page.tsx":
        route = "/" + p.replace("src/app", "").replace("/page.tsx", "").strip("/")
        return route or "/", "FRONTEND_PAGE"
    if re.search(r"src/app/api/.*/route\.ts$", p):
        route = "/api/" + p.split("src/app/api/",1)[1].replace("/route.ts", "")
        return route, "API_ENDPOINT"
    return None, None


def detect_routes(rel: str, ext: str, text: str):
    routes = []
    route, rtype = detect_route(rel, ext)
    if route:
        routes.append({"route": route, "type": rtype, "line": 1, "method": "", "signature": rel})
    for line_no, line in enumerate(text.splitlines(), start=1):
        stripped = line.strip()
        m = re.match(r"@(?:app|router|blueprint)\.(get|post|put|patch|delete|route)\(['\"]([^'\"]+)['\"]", stripped)
        if m:
            routes.append({"route": m.group(2), "type": "PYTHON_BACKEND_ROUTE", "line": line_no, "method": m.group(1).upper(), "signature": stripped[:300]})
            continue
        m = re.match(r"(?:app|router)\.(get|post|put|patch|delete)\(['\"]([^'\"]+)['\"]", stripped)
        if m:
            routes.append({"route": m.group(2), "type": "JS_BACKEND_ROUTE", "line": line_no, "method": m.group(1).upper(), "signature": stripped[:300]})
    deduped = []
    seen = set()
    for item in routes:
        key = (item["route"], item["type"], item["line"])
        if key not in seen:
            seen.add(key)
            deduped.append(item)
    return deduped


def parse_package_json_dependencies(text: str):
    deps = []
    try:
        data = json.loads(text)
    except Exception:
        return deps
    for section in ("dependencies", "devDependencies", "peerDependencies", "optionalDependencies"):
        for package_name, version_spec in (data.get(section) or {}).items():
            deps.append((str(package_name), str(version_spec), "node", section))
    return deps


def parse_requirements_dependencies(text: str):
    deps = []
    for raw in text.splitlines()[:2000]:
        line = raw.split("#", 1)[0].strip()
        if not line or line.startswith(("-", "--")):
            continue
        package_name = re.split(r"[=<>!~;\[\] ]+", line, maxsplit=1)[0].strip()
        if package_name:
            deps.append((package_name, line, "python", "runtime"))
    return deps


def parse_pyproject_dependencies(text: str):
    deps = []
    try:
        import tomllib  # type: ignore[import-not-found]
        data = tomllib.loads(text)
        project = data.get("project", {})
        for spec in project.get("dependencies", []) or []:
            package_name = re.split(r"[=<>!~;\[\] ]+", str(spec), maxsplit=1)[0].strip()
            if package_name:
                deps.append((package_name, str(spec), "python", "project.dependencies"))
        optional = project.get("optional-dependencies", {}) or {}
        for group, values in optional.items():
            for spec in values or []:
                package_name = re.split(r"[=<>!~;\[\] ]+", str(spec), maxsplit=1)[0].strip()
                if package_name:
                    deps.append((package_name, str(spec), "python", f"project.optional-dependencies.{group}"))
    except Exception:
        in_dependencies = False
        for raw in text.splitlines():
            stripped = raw.strip()
            if stripped.startswith("[") and stripped.endswith("]"):
                in_dependencies = stripped in {"[project]", "[tool.poetry.dependencies]"}
                continue
            if in_dependencies and "=" in stripped and not stripped.startswith("#"):
                package_name = stripped.split("=", 1)[0].strip().strip('"').strip("'")
                if package_name and package_name.lower() != "python":
                    deps.append((package_name, stripped, "python", "pyproject"))
    return deps


def insert_dependency_manifest(con: sqlite3.Connection, file: Path, fid: str, rel: str, digest: str, text: str) -> None:
    name = file.name.lower()
    if not is_dependency_manifest_file(file):
        return
    ecosystem = "node" if name in {"package.json", "package-lock.json", "pnpm-lock.yaml", "yarn.lock"} or name.endswith((".config.js", ".config.ts", ".config.mjs", ".config.cjs")) else "python"
    manifest_id = "manifest_" + sha256_bytes((fid + "|" + name).encode())[:16]
    con.execute(
        """INSERT OR REPLACE INTO dependency_manifest(
            manifest_id, file_id, manifest_type, ecosystem, path, sha256
        ) VALUES(?,?,?,?,?,?)""",
        (manifest_id, fid, name, ecosystem, rel, digest),
    )
    deps = []
    if name == "package.json":
        deps = parse_package_json_dependencies(text)
    elif name == "requirements.txt":
        deps = parse_requirements_dependencies(text)
    elif name == "pyproject.toml":
        deps = parse_pyproject_dependencies(text)
    for package_name, version_spec, item_ecosystem, group in deps:
        dep_id = "dep_" + sha256_bytes((manifest_id + "|" + package_name + "|" + group).encode())[:16]
        con.execute(
            """INSERT OR REPLACE INTO dependency_item(
                dependency_id, manifest_id, package_name, version_spec, ecosystem
            ) VALUES(?,?,?,?,?)""",
            (dep_id, manifest_id, package_name, version_spec, item_ecosystem),
        )
        edge_id = "depedge_" + sha256_bytes((dep_id + "|" + fid).encode())[:16]
        con.execute(
            """INSERT OR REPLACE INTO code_dependency_edge(
                id, source_id, name, path, value, metadata_json, created_at
            ) VALUES(?,?,?,?,?,?,?)""",
            (edge_id, fid, package_name, rel, group, json.dumps({"manifest_id": manifest_id, "version_spec": version_spec}), now()),
        )


def study_project_artifact(con: sqlite3.Connection, project_root: Path, vault: Path, file: Path, fid: str, rel: str, ext: str, size: int, digest: str, status: str, *, preserve_payload: bool = True) -> None:
    dest = preserve_artifact(vault, file, digest) if preserve_payload else None
    aid = "artifact_" + digest[:16]
    atype = ext.strip(".") or file.name.lower() or "artifact"
    meta = {
        "original_path": rel,
        "preserved_path": str(dest.relative_to(project_root)) if dest else "",
        "policy": "read_only_artifact_studied_not_code_chunked" if dest else "hash_and_text_studied_without_reference_asset_copy",
        "status": status,
    }
    con.execute(
        """INSERT OR REPLACE INTO project_artifact(
            artifact_id, source_file_id, artifact_type, artifact_sha256,
            path, size_bytes, metadata_json, semantic_status
        ) VALUES(?,?,?,?,?,?,?,?)""",
        (aid, fid, atype, digest, rel, size, json.dumps(meta, sort_keys=True), "READ_ONLY_STUDIED"),
    )
    con.execute(
        """INSERT OR REPLACE INTO artifact_relation_edge(
            edge_id, artifact_id, related_entity_type, related_entity_id, relation_type, confidence
        ) VALUES(?,?,?,?,?,?)""",
        ("artifact_edge_" + sha256_bytes((aid + "|" + fid).encode())[:16], aid, "source_file", fid, "BELONGS_TO_CODE_PROJECT", "deterministic"),
    )
    text = ""
    # The Env15 adapter imports the code/schema/dependency projections from its
    # disposable database, but deliberately does not import artifact FTS text.
    # When payload preservation is disabled, reading and indexing thousands of
    # reference-app translations/SVGs/Markdown files is pure throw-away work.
    # Dependency manifests remain parsed because their structured dependency
    # rows are part of the governed code projection.
    if (preserve_payload and ext in TEXT_ARTIFACT_EXTS) or is_dependency_manifest_file(file):
        try:
            text, _enc, _data = read_text_lossless(file)
        except Exception:
            text = ""
    if text.strip():
        con.execute("DELETE FROM artifact_fts WHERE entity_id=?", (aid,))
        con.execute("INSERT INTO artifact_fts(entity_id,text) VALUES(?,?)", (aid, text[:MAX_ARTIFACT_TEXT_EXTRACT]))
        insert_dependency_manifest(con, file, fid, rel, digest, text)


def sanitize_git_remote_url(url: str) -> str:
    return re.sub(r"(https?://)([^/@]+@)", r"\1", url or "")


def git_stdout(folder: Path, args: list[str], timeout: int = 30) -> str:
    try:
        # Git can emit raw high-bit bytes for very small media fixtures that do
        # not contain a NUL byte and are therefore classified as text. Capture
        # bytes first so Windows' locale decoder cannot terminate the reader
        # thread and silently discard the rest of a commit patch.
        result = run_hidden(["git", "-C", str(folder), *args], capture_output=True, timeout=timeout)
        if result.returncode == 0:
            return result.stdout.decode("utf-8", errors="replace").strip()
    except Exception:
        pass
    return ""


def current_git_head(folder: Path) -> str:
    if not (folder / ".git").exists():
        return ""
    return git_stdout(folder, ["rev-parse", "HEAD"], timeout=15)


def git_bytes(folder: Path, args: list[str], timeout: int = 30) -> bytes:
    try:
        result = run_hidden(["git", "-C", str(folder), *args], capture_output=True, timeout=timeout)
        if result.returncode == 0:
            return result.stdout
    except Exception:
        pass
    return b""


def normalize_git_path(value: str) -> str:
    value = (value or "").strip().strip('"')
    if value == "/dev/null":
        return ""
    if value.startswith(("a/", "b/")):
        value = value[2:]
    return value.replace("\\", "/")


def git_line_history_policy(path: str, blob_size: int = 0) -> str:
    normalized = normalize_git_path(path).lower()
    ext = Path(normalized).suffix.lower()
    name = Path(normalized).name
    parts = {part for part in normalized.split("/") if part}
    stem = Path(normalized).stem
    generated_name = any(token in stem for token in ("forecast", "prediction", "generated", "model_output"))
    if ext in BINARY_ASSET_EXTS:
        return "METADATA_HASH_ONLY"
    # Directory role wins over extension. Generated output can contain .py, .md,
    # or .json files, but those are artifacts rather than authored source code.
    if parts.intersection(GENERATED_HISTORY_DIRS) or generated_name:
        return "HUNK_SUMMARY_GENERATED_ARTIFACT"
    if ext in CODE_EXTS or name in SPECIAL_CODE_FILENAMES or is_dependency_manifest_file(Path(name)):
        return "EXACT_CODE_CONFIG"
    if name in CONFIG_JSON_NAMES or any(token in name for token in ("config", "manifest", "schema", "settings")):
        return "EXACT_CODE_CONFIG" if blob_size <= MAX_ARTIFACT_TEXT_EXTRACT else "HUNK_SUMMARY_LARGE_ARTIFACT"
    if ext in GENERATED_HISTORY_EXTS or ext in {".ipynb", ".json"}:
        return "HUNK_SUMMARY_STRUCTURED_ARTIFACT"
    if blob_size > MAX_EXACT_GIT_DOC_BYTES:
        return "HUNK_SUMMARY_LARGE_ARTIFACT"
    if ext in SMALL_HUMAN_DOC_EXTS:
        return "EXACT_SMALL_DOC"
    return "HUNK_SUMMARY_ARTIFACT"


def parse_git_diff_hunks(commit_sha: str, patch_text: str, path_policies: Optional[dict[str, str]] = None) -> list[dict]:
    path_policies = path_policies or {}
    hunks: list[dict] = []
    current_old = ""
    current_new = ""
    current: Optional[dict] = None
    old_line = 0
    new_line = 0

    def flush_current() -> None:
        nonlocal current
        if current and not current.get("skip"):
            current["patch_sha256"] = current.pop("patch_hasher").hexdigest()
            current.pop("store_exact_lines", None)
            current["stored_line_count"] = len(current["changes"])
            hunks.append(current)
        current = None

    for line in patch_text.splitlines():
        if line.startswith("diff --git "):
            flush_current()
            try:
                fields = shlex.split(line)
            except ValueError:
                fields = line.split()
            current_old = normalize_git_path(fields[2]) if len(fields) > 2 else ""
            current_new = normalize_git_path(fields[3]) if len(fields) > 3 else ""
            continue
        if line.startswith("--- "):
            current_old = normalize_git_path(line[4:].split("\t", 1)[0])
            continue
        if line.startswith("+++ "):
            current_new = normalize_git_path(line[4:].split("\t", 1)[0])
            continue
        if line.startswith("@@ "):
            flush_current()
            match = re.match(r"@@ -(\d+)(?:,(\d+))? \+(\d+)(?:,(\d+))? @@", line)
            if not match:
                continue
            old_line = int(match.group(1))
            new_line = int(match.group(3))
            compared_path = current_new or current_old
            history_policy = path_policies.get(compared_path) or git_line_history_policy(compared_path)
            patch_hasher = hashlib.sha256()
            patch_hasher.update(line.encode("utf-8", "surrogatepass"))
            current = {
                "commit_sha": commit_sha,
                "old_path": current_old,
                "new_path": current_new,
                "old_start": old_line,
                "old_count": int(match.group(2) or "1"),
                "new_start": new_line,
                "new_count": int(match.group(4) or "1"),
                "hunk_header": line,
                "changes": [],
                "patch_hasher": patch_hasher,
                "history_policy": history_policy,
                "changed_line_count": 0,
                "store_exact_lines": history_policy.startswith("EXACT_"),
                "skip": history_policy == "METADATA_HASH_ONLY",
            }
            continue
        if not current or line.startswith("\\ No newline at end of file"):
            continue
        current["patch_hasher"].update(b"\n")
        current["patch_hasher"].update(line.encode("utf-8", "surrogatepass"))
        if current.get("skip"):
            continue
        path = current_new or current_old
        if line.startswith("-") and not line.startswith("---"):
            current["changed_line_count"] += 1
            if current["store_exact_lines"]:
                current["changes"].append({"path": path, "old_line": old_line, "new_line": None, "change_type": "DELETE", "text": line[1:]})
            old_line += 1
        elif line.startswith("+") and not line.startswith("+++"):
            current["changed_line_count"] += 1
            if current["store_exact_lines"]:
                current["changes"].append({"path": path, "old_line": None, "new_line": new_line, "change_type": "ADD", "text": line[1:]})
            new_line += 1
        elif line.startswith(" "):
            old_line += 1
            new_line += 1
    flush_current()
    return hunks


def record_git_diff_details(con: sqlite3.Connection, folder: Path, commit_sha: str, path_policies: Optional[dict[str, str]] = None) -> None:
    con.execute("DELETE FROM git_line_change WHERE commit_sha=?", (commit_sha,))
    con.execute("DELETE FROM git_diff_hunk WHERE commit_sha=?", (commit_sha,))
    patch_text = git_stdout(
        folder,
        ["show", "--format=", "--find-renames", "--unified=0", "--no-ext-diff", "--no-color", commit_sha],
        timeout=120,
    )
    for hunk_order, hunk in enumerate(parse_git_diff_hunks(commit_sha, patch_text, path_policies)):
        hunk_id = "hunk_" + sha256_bytes(
            f"{commit_sha}|{hunk['old_path']}|{hunk['new_path']}|{hunk['hunk_header']}|{hunk_order}".encode()
        )[:20]
        con.execute(
            """INSERT OR REPLACE INTO git_diff_hunk(
                hunk_id, commit_sha, old_path, new_path, old_start, old_count, new_start, new_count,
                hunk_header, patch_sha256, history_policy, changed_line_count, stored_line_count
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                hunk_id,
                commit_sha,
                hunk["old_path"],
                hunk["new_path"],
                hunk["old_start"],
                hunk["old_count"],
                hunk["new_start"],
                hunk["new_count"],
                hunk["hunk_header"],
                hunk["patch_sha256"],
                hunk["history_policy"],
                hunk["changed_line_count"],
                hunk["stored_line_count"],
            ),
        )
        for line_order, change in enumerate(hunk["changes"]):
            line_change_id = "line_change_" + sha256_bytes(
                f"{hunk_id}|{line_order}|{change['change_type']}|{change['old_line']}|{change['new_line']}".encode()
            )[:20]
            line_text = str(change["text"])
            con.execute(
                """INSERT OR REPLACE INTO git_line_change(
                    line_change_id, hunk_id, commit_sha, path, old_line_number,
                    new_line_number, change_type, line_text, line_sha256
                ) VALUES(?,?,?,?,?,?,?,?,?)""",
                (
                    line_change_id,
                    hunk_id,
                    commit_sha,
                    change["path"],
                    change["old_line"],
                    change["new_line"],
                    change["change_type"],
                    line_text,
                    sha256_bytes(line_text.encode("utf-8", "surrogatepass")),
                ),
            )


def compact_existing_commit_history(con: sqlite3.Connection, commit_sha: str) -> None:
    """Migrate an already indexed commit without asking Git to index it again."""
    rows = list(
        con.execute(
            "SELECT hunk_id,old_path,new_path,COALESCE(changed_line_count,0) "
            "FROM git_diff_hunk WHERE commit_sha=?",
            (commit_sha,),
        )
    )
    for hunk_id, old_path, new_path, recorded_changed_count in rows:
        path = normalize_git_path(new_path or old_path or "")
        policy = git_line_history_policy(path)
        stored_count = con.execute(
            "SELECT COUNT(*) FROM git_line_change WHERE hunk_id=?",
            (hunk_id,),
        ).fetchone()[0]
        changed_count = max(int(recorded_changed_count or 0), int(stored_count or 0))
        if policy == "METADATA_HASH_ONLY":
            con.execute("DELETE FROM git_line_change WHERE hunk_id=?", (hunk_id,))
            con.execute("DELETE FROM git_diff_hunk WHERE hunk_id=?", (hunk_id,))
        elif not policy.startswith("EXACT_"):
            con.execute("DELETE FROM git_line_change WHERE hunk_id=?", (hunk_id,))
            con.execute(
                "UPDATE git_diff_hunk SET history_policy=?,changed_line_count=?,stored_line_count=0 WHERE hunk_id=?",
                (policy, changed_count, hunk_id),
            )
        else:
            con.execute(
                "UPDATE git_diff_hunk SET history_policy=?,changed_line_count=?,stored_line_count=? WHERE hunk_id=?",
                (policy, changed_count, stored_count, hunk_id),
            )


def git_blob_identity(folder: Path, revision: str, path: str) -> tuple[str, int]:
    if not revision or not path:
        return "", 0
    blob_oid = git_stdout(folder, ["rev-parse", "--verify", f"{revision}:{path}"], timeout=15)
    blob_size_text = git_stdout(folder, ["cat-file", "-s", blob_oid], timeout=15) if blob_oid else ""
    return blob_oid, int(blob_size_text) if blob_size_text.isdigit() else 0


def record_git_file_version(
    con: sqlite3.Connection,
    folder: Path,
    commit_sha: str,
    change_type: str,
    old_path: str,
    new_path: str,
    file_id: str,
) -> tuple[Optional[str], dict]:
    path_at_commit = old_path if change_type == "D" else new_path
    ext = Path(path_at_commit).suffix.lower()
    name = Path(path_at_commit).name.lower()
    parent_sha = git_stdout(folder, ["rev-parse", f"{commit_sha}^1"], timeout=15)
    old_blob_oid, old_blob_size = git_blob_identity(folder, parent_sha, old_path or new_path)
    new_blob_oid, new_blob_size = ("", 0) if change_type == "D" else git_blob_identity(folder, commit_sha, new_path)
    blob_oid = new_blob_oid or old_blob_oid
    blob_size = new_blob_size or old_blob_size
    line_history_policy = git_line_history_policy(path_at_commit, blob_size)
    metadata = {
        "content_policy": "CHANGE_EVENT_ONLY",
        "line_history_policy": line_history_policy,
        "blob_oid": blob_oid,
        "blob_hash": blob_oid,
        "blob_size": blob_size,
        "old_blob_oid": old_blob_oid,
        "old_blob_size": old_blob_size,
        "new_blob_oid": new_blob_oid,
        "new_blob_size": new_blob_size,
    }
    is_code = ext in CODE_EXTS or name in SPECIAL_CODE_FILENAMES
    is_media = ext in BINARY_ASSET_EXTS

    if is_media:
        metadata["content_policy"] = "MEDIA_ASSET_HASH_METADATA_ONLY"
        return None, metadata
    if not line_history_policy.startswith("EXACT_"):
        metadata["content_policy"] = "READ_ONLY_ARTIFACT_CHANGE_EVENT"
        return None, metadata
    if not is_code:
        metadata["content_policy"] = "READ_ONLY_ARTIFACT_CHANGE_EVENT"
        return None, metadata

    language = detect_language(ext, path_at_commit)
    file_version_id = "fv_git_" + sha256_bytes(f"{commit_sha}|{path_at_commit}|{change_type}".encode())[:20]
    con.execute(
        """INSERT OR IGNORE INTO code_file(
            file_id, canonical_path, language, extension, current_sha256, is_active
        ) VALUES(?,?,?,?,?,?)""",
        (file_id, path_at_commit, language, ext, "", 0),
    )
    if change_type == "D":
        con.execute(
            """INSERT OR REPLACE INTO code_file_version(
                file_version_id, file_id, commit_sha, raw_file_sha256, normalized_text_sha256,
                path_at_commit, language, extension, line_count, byte_count,
                is_deleted, is_renamed, created_at
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (file_version_id, file_id, commit_sha, "", "", path_at_commit, language, ext, 0, 0, 1, 0, now()),
        )
        metadata["content_policy"] = "CODE_VERSION_TOMBSTONE"
        metadata["file_version_id"] = file_version_id
        return file_version_id, metadata

    data = git_bytes(folder, ["show", f"{commit_sha}:{path_at_commit}"], timeout=60)
    if not data and metadata["new_blob_size"]:
        metadata["content_policy"] = "CODE_VERSION_UNREADABLE"
        return None, metadata
    text = ""
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    raw_sha256 = sha256_bytes(data)
    normalized_sha256 = sha256_bytes(text.replace("\r\n", "\n").encode("utf-8", "ignore"))
    con.execute(
        """INSERT OR REPLACE INTO code_file_version(
            file_version_id, file_id, commit_sha, raw_file_sha256, normalized_text_sha256,
            path_at_commit, language, extension, line_count, byte_count,
            is_deleted, is_renamed, created_at
        ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (
            file_version_id,
            file_id,
            commit_sha,
            raw_sha256,
            normalized_sha256,
            path_at_commit,
            language,
            ext,
            len(text.splitlines()),
            len(data),
            0,
            1 if change_type.startswith("R") else 0,
            now(),
        ),
    )
    metadata["content_policy"] = "CODE_VERSION_HASHED_NOT_CHUNKED"
    metadata["file_version_id"] = file_version_id
    metadata["raw_sha256"] = raw_sha256
    return file_version_id, metadata


def record_git_lineage(
    con: sqlite3.Connection,
    folder: Path,
    repo_id: str,
    source_id: str,
    file_ids_by_rel: dict[str, str],
    *,
    current_checkout_only: bool = False,
) -> None:
    if not (folder / ".git").exists():
        return
    branch = git_stdout(folder, ["branch", "--show-current"], timeout=15)
    remotes = git_stdout(folder, ["remote", "-v"], timeout=15)
    for line in remotes.splitlines():
        parts = line.split()
        if len(parts) >= 2 and "(fetch)" in line:
            remote_name, remote_url = parts[0], sanitize_git_remote_url(parts[1])
            remote_id = "remote_" + sha256_bytes((repo_id + "|" + remote_name + "|" + remote_url).encode())[:16]
            con.execute(
                """INSERT OR REPLACE INTO git_remote(
                    id, source_id, name, path, value, metadata_json, created_at
                ) VALUES(?,?,?,?,?,?,?)""",
                (remote_id, repo_id, remote_name, str(folder), remote_url, json.dumps({"source_id": source_id}, sort_keys=True), now()),
            )
    log_format = "%H%x1f%h%x1f%an%x1f%ae%x1f%aI%x1f%cn%x1f%ce%x1f%cI%x1f%P%x1f%s"
    primary_log = git_stdout(
        folder,
        ["log", "-1", "HEAD", f"--pretty=format:{log_format}"]
        if current_checkout_only
        else ["log", "HEAD", "--topo-order", f"--pretty=format:{log_format}"],
        timeout=30 if current_checkout_only else 120,
    )
    all_log = "" if current_checkout_only else git_stdout(
        folder, ["log", "--all", "--topo-order", f"--pretty=format:{log_format}"], timeout=120
    )
    commits = []
    seen_commits = set()
    for line in [*primary_log.splitlines(), *all_log.splitlines()]:
        commit_sha = line.split("\x1f", 1)[0]
        if line.strip() and commit_sha not in seen_commits:
            seen_commits.add(commit_sha)
            commits.append(line)
    head_sha = current_git_head(folder)
    default_branch = git_stdout(folder, ["symbolic-ref", "--short", "refs/remotes/origin/HEAD"], timeout=15)
    branch_rows = git_stdout(
        folder,
        ["for-each-ref", "--format=%(refname:short)%09%(objectname)", "refs/heads"],
        timeout=30,
    ).splitlines()
    if not branch_rows:
        branch_rows = [f"{branch or 'detached'}\t{head_sha}"]
    for branch_row in branch_rows:
        branch_parts = branch_row.split("\t", 1)
        branch_name = branch_parts[0] or "detached"
        branch_head = branch_parts[1] if len(branch_parts) > 1 else head_sha
        branch_id = "branch_" + sha256_bytes((repo_id + "|" + branch_name).encode())[:16]
        con.execute(
            """INSERT OR REPLACE INTO git_branch(
                id, source_id, name, path, value, metadata_json, created_at
            ) VALUES(?,?,?,?,?,?,?)""",
            (
                branch_id,
                repo_id,
                branch_name,
                str(folder),
                branch_head,
                json.dumps(
                    {
                        "commit_count": len(commits),
                        "default_branch": default_branch,
                        "is_current": branch_name == branch,
                        "walk_order": (
                            "CURRENT_CHECKOUT_HEAD_ONLY"
                            if current_checkout_only
                            else "CURRENT_HEAD_BACKWARD_THEN_OTHER_REFS"
                        ),
                        "provenance_scope": (
                            "CURRENT_CHECKOUT_PROVENANCE"
                            if current_checkout_only
                            else "FULL_REPOSITORY_REVERSE_HISTORY"
                        ),
                        "commit_sha": branch_head,
                    },
                    sort_keys=True,
                ),
                now(),
            ),
        )
    lineage_file_ids = dict(file_ids_by_rel)
    for order, line in enumerate(commits):
        parts = line.split("\x1f")
        if len(parts) < 10:
            continue
        (
            commit_sha, short_sha, author_name, author_email, author_time,
            committer_name, committer_email, committer_time, parents, subject,
        ) = parts[:10]
        index_state = con.execute(
            "SELECT policy_version FROM code_index_state WHERE repo_id=? AND commit_sha=?",
            (repo_id, commit_sha),
        ).fetchone()
        if index_state and index_state[0] == CODE_INDEX_POLICY_VERSION:
            continue
        existing_commit = con.execute(
            "SELECT 1 FROM git_commit WHERE commit_sha=?",
            (commit_sha,),
        ).fetchone()
        if existing_commit:
            compact_existing_commit_history(con, commit_sha)
            con.execute(
                "INSERT OR REPLACE INTO code_index_state(repo_id,commit_sha,policy_version,indexed_at) VALUES(?,?,?,?)",
                (repo_id, commit_sha, CODE_INDEX_POLICY_VERSION, now()),
            )
            continue
        email_hash = sha256_bytes(author_email.encode()) if author_email else ""
        committer_email_hash = sha256_bytes(committer_email.encode()) if committer_email else ""
        con.execute(
            """INSERT OR REPLACE INTO git_commit(
                commit_sha, short_sha, author_name, author_email_hash,
                author_time, committer_name, committer_email_hash, committer_time,
                commit_time, message, commit_order
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
            (
                commit_sha, short_sha, author_name, email_hash, author_time,
                committer_name, committer_email_hash, committer_time,
                committer_time, subject, order,
            ),
        )
        con.execute("DELETE FROM commit_fts WHERE entity_id=?", (commit_sha,))
        con.execute("INSERT INTO commit_fts(entity_id,text) VALUES(?,?)", (commit_sha, f"{short_sha} {subject} parents={parents}"))
        for parent_order, parent_sha in enumerate(p for p in parents.split() if p):
            parent_edge_id = "parent_" + sha256_bytes(f"{commit_sha}|{parent_sha}|{parent_order}".encode())[:20]
            con.execute(
                """INSERT OR REPLACE INTO git_commit_parent_edge(
                    edge_id, commit_sha, parent_sha, parent_order
                ) VALUES(?,?,?,?)""",
                (parent_edge_id, commit_sha, parent_sha, parent_order),
            )
        if current_checkout_only:
            con.execute(
                "INSERT OR REPLACE INTO code_index_state(repo_id,commit_sha,policy_version,indexed_at) VALUES(?,?,?,?)",
                (repo_id, commit_sha, "LOCAL_CURRENT_CHECKOUT_PROVENANCE_V1", now()),
            )
            continue
        show_text = git_stdout(folder, ["show", "--name-status", "--find-renames", "--format=", commit_sha], timeout=30)
        commit_path_policies: dict[str, str] = {}
        for change_order, raw_change in enumerate(line for line in show_text.splitlines() if line.strip()):
            fields = raw_change.split("\t")
            change_type = fields[0]
            old_path = ""
            new_path = normalize_git_path(fields[-1]) if len(fields) >= 2 else ""
            if change_type.startswith(("R", "C")) and len(fields) >= 3:
                old_path, new_path = normalize_git_path(fields[1]), normalize_git_path(fields[2])
            elif change_type == "D":
                old_path, new_path = new_path, ""
            linked_file_id = lineage_file_ids.get(new_path) or lineage_file_ids.get(old_path)
            if not linked_file_id:
                linked_file_id = file_id_for(new_path or old_path, "")
            if new_path:
                lineage_file_ids[new_path] = linked_file_id
            if old_path:
                lineage_file_ids[old_path] = linked_file_id
            file_version_id, version_metadata = record_git_file_version(
                con,
                folder,
                commit_sha,
                change_type,
                old_path,
                new_path,
                linked_file_id,
            )
            line_history_policy = str(version_metadata["line_history_policy"])
            if new_path:
                commit_path_policies[new_path] = line_history_policy
            if old_path:
                commit_path_policies[old_path] = line_history_policy
            if change_type.startswith("R"):
                similarity_text = change_type[1:]
                rename_id = "rename_" + sha256_bytes(f"{commit_sha}|{old_path}|{new_path}".encode())[:20]
                con.execute(
                    """INSERT OR REPLACE INTO git_rename_map(
                        rename_id, commit_sha, old_path, new_path, similarity
                    ) VALUES(?,?,?,?,?)""",
                    (rename_id, commit_sha, old_path, new_path, int(similarity_text) if similarity_text.isdigit() else 0),
                )
            change_id = "change_" + sha256_bytes((commit_sha + "|" + raw_change).encode())[:16]
            con.execute(
                """INSERT OR REPLACE INTO git_file_change(
                    id, source_id, name, path, value, metadata_json, created_at
                ) VALUES(?,?,?,?,?,?,?)""",
                (
                    change_id,
                    repo_id,
                    change_type,
                    new_path or old_path,
                    commit_sha,
                    json.dumps(
                        {
                            "old_path": old_path,
                            "new_path": new_path,
                            "file_id": linked_file_id,
                            "file_version_id": file_version_id,
                            "change_order": change_order,
                            "commit_is_change_event": True,
                            **version_metadata,
                        },
                        sort_keys=True,
                    ),
                    now(),
                ),
            )
        record_git_diff_details(con, folder, commit_sha, commit_path_policies)
        con.execute(
            "INSERT OR REPLACE INTO code_index_state(repo_id,commit_sha,policy_version,indexed_at) VALUES(?,?,?,?)",
            (repo_id, commit_sha, CODE_INDEX_POLICY_VERSION, now()),
        )


def preserve_artifact(vault: Path, file: Path, digest: str):
    vault.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", file.name)
    dest = vault / f"{digest[:16]}__{safe}"
    if not dest.exists():
        shutil.copy2(file, dest)
    return dest


def clear_live_file_projection(con: sqlite3.Connection, file_id: str) -> None:
    """Replace the live snapshot for one path without touching immutable Git versions."""
    chunk_ids = [row[0] for row in con.execute("SELECT chunk_id FROM code_chunk WHERE file_id=?", (file_id,))]
    line_ids = [row[0] for row in con.execute("SELECT line_id FROM code_line_snapshot WHERE file_id=?", (file_id,))]
    symbol_ids = [row[0] for row in con.execute("SELECT symbol_id FROM code_symbol WHERE file_id=?", (file_id,))]
    route_ids = [row[0] for row in con.execute("SELECT route_id FROM app_route WHERE file_id=?", (file_id,))]
    for table_name, identifiers in (
        ("code_fts", chunk_ids),
        ("line_fts", line_ids),
        ("symbol_fts", symbol_ids),
        ("route_fts", route_ids),
    ):
        for identifier in identifiers:
            con.execute(f'DELETE FROM "{table_name}" WHERE entity_id=?', (identifier,))
    manifests = [row[0] for row in con.execute("SELECT manifest_id FROM dependency_manifest WHERE file_id=?", (file_id,))]
    for manifest_id in manifests:
        con.execute("DELETE FROM dependency_item WHERE manifest_id=?", (manifest_id,))
    con.execute("DELETE FROM dependency_manifest WHERE file_id=?", (file_id,))
    con.execute("DELETE FROM code_dependency_edge WHERE source_id=?", (file_id,))
    con.execute("DELETE FROM code_import_edge WHERE from_file_id=?", (file_id,))
    con.execute("DELETE FROM app_route WHERE file_id=?", (file_id,))
    con.execute("DELETE FROM code_symbol WHERE file_id=?", (file_id,))
    con.execute("DELETE FROM code_chunk WHERE file_id=?", (file_id,))
    con.execute("DELETE FROM code_line_snapshot WHERE file_id=?", (file_id,))
    con.execute("UPDATE code_file SET is_active=0 WHERE file_id=?", (file_id,))


def run_code_lane_digestion_job(con: sqlite3.Connection, folder: Path, source: dict, project_root: Path, progress: ProgressCallback, lane_key: str, started: float, *, preserve_artifact_payloads: bool = True) -> None:
    vault = project_root / "artifacts" / "code_asset_vault"
    files = [p for p in folder.rglob("*") if p.is_file() and not is_ignored(p.relative_to(folder))]
    total = len(files)
    source_id = source.get("source_id", "source_code")
    repo_id = "repo_" + hashlib.sha256(str(folder).encode()).hexdigest()[:12]
    current_commit = current_git_head(folder) or None
    con.execute(
        """INSERT OR REPLACE INTO code_repo(
            id, source_id, name, path, value, metadata_json, created_at
        ) VALUES(?,?,?,?,?,?,?)""",
        (
            repo_id,
            source_id,
            folder.name,
            str(folder),
            "CODE_LANE_DIGESTION_JOB",
            json.dumps({"root": str(folder), "lane_key": lane_key, "current_commit": current_commit, "file_count": total}, sort_keys=True),
            now(),
        ),
    )
    replace_existing = bool(con.execute("SELECT 1 FROM code_file LIMIT 1").fetchone())
    insert_folder_chunks(con, folder, files, source_id, replace_fts=replace_existing)
    file_ids_by_rel: dict[str, str] = {}

    for idx, file in enumerate(files, start=1):
        rel = file.relative_to(folder).as_posix()
        raw_rel = str(file.relative_to(folder))
        ext = file.suffix.lower()
        size = file.stat().st_size
        digest = sha256_file(file)
        fid = file_id_for(rel, digest)
        file_ids_by_rel[rel] = fid
        # The Env15 adapter builds into a fresh staging database.  Avoid a
        # quadratic series of unindexed DELETE scans for files that cannot yet
        # exist; retain replacement behavior when this function is used on a
        # populated legacy database.
        if con.execute("SELECT 1 FROM code_file WHERE file_id=? LIMIT 1", (fid,)).fetchone():
            clear_live_file_projection(con, fid)
        status = classify_code_project_file(file, Path(rel))
        coverage_note = {
            "policy": "code snapshot first; Git lineage recorded as separate change events",
            "raw_path": raw_rel,
            "normalized_path": rel,
            "media_asset_study": "blocked" if status == "CODE_MEDIA_ASSET_HASH_ONLY" else "not_media_asset",
            "artifact_study": status.startswith("CODE_ARTIFACT"),
        }
        con.execute(
            "INSERT OR REPLACE INTO source_file(file_id, logical_path, extension, size_bytes, sha256, lane_status, created_at) VALUES(?,?,?,?,?,?,?)",
            (fid, rel, ext, size, digest, status, now()),
        )
        con.execute(
            """INSERT OR REPLACE INTO source_byte_coverage(
                file_id, coverage_status, byte_count, sha256, notes, created_at
            ) VALUES(?,?,?,?,?,?)""",
            (fid, status, size, digest, json.dumps(coverage_note, sort_keys=True), now()),
        )

        if status == "LOSSLESS_CODE_LAYERED_CHUNKED":
            text, enc, data = read_text_lossless(file)
            lang = detect_language(ext, rel)
            role = classify_code_role(rel, ext, text)
            norm_digest = sha256_bytes(text.replace("\r\n", "\n").encode("utf-8", "ignore"))
            fvid = "fv_" + sha256_bytes((rel + "|" + digest).encode())[:16]
            lines = text.splitlines(keepends=False)
            con.execute(
                """INSERT OR REPLACE INTO code_file(
                    file_id, canonical_path, language, extension, current_sha256, is_active
                ) VALUES(?,?,?,?,?,?)""",
                (fid, rel, lang, ext, digest, 1),
            )
            con.execute(
                """INSERT OR REPLACE INTO code_file_version(
                    file_version_id, file_id, commit_sha, raw_file_sha256, normalized_text_sha256,
                    path_at_commit, language, extension, line_count, byte_count,
                    is_deleted, is_renamed, created_at
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (fvid, fid, current_commit, digest, norm_digest, rel, lang, ext, len(lines), len(data), 0, 0, now()),
            )
            file_payload = "\n".join(
                [
                    f"path: {raw_rel}",
                    f"normalized_path: {rel}",
                    f"language: {lang}",
                    f"role: {role}",
                    f"encoding: {enc}",
                    f"line_count: {len(lines)}",
                    f"sha256: {digest}",
                    "",
                    text,
                ]
            )
            insert_code_chunk(con, "chunk_file_" + sha256_bytes((rel + "|" + digest).encode())[:16], fvid, fid, "FILE_CHUNK", lang, role, 1 if lines else 0, len(lines), file_payload, replace_fts=replace_existing)

            for ln, line in enumerate(lines, start=1):
                lsha = sha256_bytes(line.encode("utf-8", "ignore"))
                norm = line.strip()
                nsha = sha256_bytes(norm.encode("utf-8", "ignore"))
                is_blank = 1 if not norm else 0
                is_comment = 1 if norm.startswith(("#", "//", "/*", "*", "--")) else 0
                indent = len(line) - len(line.lstrip(" \t"))
                lid = f"line_{sha256_bytes(rel.encode())[:10]}_{ln}"
                con.execute(
                    """INSERT OR REPLACE INTO code_line_snapshot(
                        line_id, file_version_id, file_id, line_number, line_text,
                        line_sha256, normalized_line_sha256, indent_level,
                        is_blank, is_comment, search_text
                    ) VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                    (lid, fvid, fid, ln, line, lsha, nsha, indent, is_blank, is_comment, line),
                )
                if norm:
                    if replace_existing:
                        con.execute("DELETE FROM line_fts WHERE entity_id=?", (lid,))
                    con.execute("INSERT INTO line_fts(entity_id,text) VALUES(?,?)", (lid, line))

            imports = parse_imports(text, lang)
            for ln, target, typ in imports:
                eid = "import_" + sha256_bytes(f"{fid}|{ln}|{target}".encode())[:16]
                con.execute(
                    """INSERT OR REPLACE INTO code_import_edge(
                        edge_id, from_file_id, from_path, import_target, import_type, line_number
                    ) VALUES(?,?,?,?,?,?)""",
                    (eid, fid, rel, target, typ, ln),
                )
            if imports:
                import_text = "\n".join(f"{ln}: {typ} {target}" for ln, target, typ in imports)
                insert_code_chunk(con, "chunk_imports_" + sha256_bytes((rel + "|" + digest).encode())[:16], fvid, fid, "IMPORT_CHUNK", lang, role, imports[0][0], imports[-1][0], import_text, replace_fts=replace_existing)

            symbols = detect_symbol_ranges(text, lang)
            for item in symbols:
                item["symbol_id"] = "sym_" + sha256_bytes(
                    f"{fid}|{item['start']}|{item['name']}|{item['signature']}".encode()
                )[:16]
            for symbol_index, item in enumerate(symbols):
                start_line = int(item["start"])
                end_line = int(item["end"])
                name = str(item["name"])
                typ = str(item["type"])
                sig = str(item["signature"])
                sid = str(item["symbol_id"])
                parent_symbol_id = None
                for candidate in reversed(symbols[:symbol_index]):
                    if (
                        int(candidate["start"]) < start_line <= int(candidate["end"])
                        and int(candidate["indent"]) < int(item["indent"])
                    ):
                        parent_symbol_id = str(candidate["symbol_id"])
                        break
                con.execute(
                    """INSERT OR REPLACE INTO code_symbol(
                        symbol_id, symbol_name, symbol_type, file_version_id, file_id,
                        language, start_line, end_line, parent_symbol_id, signature, symbol_sha256
                    ) VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                    (sid, name, typ, fvid, fid, lang, start_line, end_line, parent_symbol_id, sig, sha256_bytes(sig.encode())),
                )
                if replace_existing:
                    con.execute("DELETE FROM symbol_fts WHERE entity_id=?", (sid,))
                con.execute("INSERT INTO symbol_fts(entity_id,text) VALUES(?,?)", (sid, f"{name} {typ} {sig}"))
                symbol_body = "\n".join(lines[start_line - 1:end_line])
                symbol_text = f"symbol: {name}\ntype: {typ}\nsignature: {sig}\npath: {rel}\n\n{symbol_body}"
                insert_code_chunk(con, "chunk_" + sid, fvid, fid, "SYMBOL_CHUNK", lang, typ, start_line, end_line, symbol_text, replace_fts=replace_existing)

            routes = detect_routes(rel, ext, text)
            for route_item in routes:
                route_path = str(route_item["route"])
                rtype = str(route_item["type"])
                method = str(route_item.get("method") or "")
                route_line = int(route_item.get("line") or 1)
                rid = "route_" + sha256_bytes((route_path + "|" + fid + "|" + str(route_line)).encode())[:16]
                con.execute(
                    """INSERT OR REPLACE INTO app_route(
                        route_id, route_path, route_type, file_id, file_version_id, route_sha256
                    ) VALUES(?,?,?,?,?,?)""",
                    (rid, route_path, rtype, fid, fvid, sha256_bytes((route_path + "|" + rel).encode())),
                )
                if replace_existing:
                    con.execute("DELETE FROM route_fts WHERE entity_id=?", (rid,))
                con.execute("INSERT INTO route_fts(entity_id,text) VALUES(?,?)", (rid, f"{route_path} {rtype} {method} {rel}"))
                route_text = f"route: {route_path}\ntype: {rtype}\nmethod: {method}\npath: {rel}\nline: {route_line}\nsignature: {route_item.get('signature', '')}"
                insert_code_chunk(con, "chunk_" + rid, fvid, fid, "ROUTE_CHUNK", lang, rtype, route_line, route_line, route_text, replace_fts=replace_existing)

            for start in range(1, len(lines) + 1, CHUNK_LINES):
                end = min(len(lines), start + CHUNK_LINES - 1)
                chunk_text = "\n".join(lines[start - 1:end])
                cid = f"chunk_lines_{sha256_bytes(rel.encode())[:10]}_{start}_{end}"
                insert_code_chunk(con, cid, fvid, fid, "LINE_CHUNK", lang, role, start, end, chunk_text, replace_fts=replace_existing)
            if not lines:
                cid = f"chunk_lines_{sha256_bytes(rel.encode())[:10]}_0_0"
                insert_code_chunk(con, cid, fvid, fid, "LINE_CHUNK", lang, role, 0, 0, "", replace_fts=replace_existing)

            insert_dependency_manifest(con, file, fid, rel, digest, text)

        elif status in {"CODE_ARTIFACT_STUDIED_READ_ONLY", "CODE_ARTIFACT_METADATA_ONLY", "CODE_TEXT_TOO_LARGE_ARTIFACT_READ_ONLY"}:
            study_project_artifact(con, project_root, vault, file, fid, rel, ext, size, digest, status, preserve_payload=preserve_artifact_payloads)
        # Media/image/video/3D assets intentionally stop at source_file + byte coverage rows.

        if idx % 25 == 0 or idx == total:
            con.commit()
            emit(progress, "build", "layered code lane digestion", rel, int(idx * 70 / max(total, 1)), idx, total, started)

    # Local Code keeps the current filesystem snapshot plus attributable current
    # checkout provenance.  Reverse-history patches, renames, and line-change
    # radiation remain exclusive to the separately selected GitHub lane.
    if lane_key in {"github_code", "github"}:
        record_git_lineage(con, folder, repo_id, source_id, file_ids_by_rel)
    elif lane_key in {"local_code", "local"} and (folder / ".git").exists():
        record_git_lineage(
            con,
            folder,
            repo_id,
            source_id,
            file_ids_by_rel,
            current_checkout_only=True,
        )
    create_code_query_indexes(con)
    con.commit()


def build_code_sector(db: Path, source: dict, project_root: Path, progress: ProgressCallback, lane_key: str, started: float, *, preserve_artifact_payloads: bool = True):
    folder = Path(source.get("path") or "")
    if not folder.exists():
        raise RuntimeError(f"CODE_SOURCE_NOT_FOUND: {folder}")
    con = connect(db)
    create_code_schema(con)
    try:
        run_code_lane_digestion_job(con, folder, source, project_root, progress, lane_key, started, preserve_artifact_payloads=preserve_artifact_payloads)
    finally:
        if preserve_artifact_payloads:
            vacuum_close(con)
        else:
            # Env15 consumes this database once and deletes it. VACUUMing a
            # disposable FTS-heavy adapter was the dominant OpenWebUI delay.
            con.commit()
            con.close()


def _ginsert(con, table, source_id, name, path, value, meta=None):
    create_generic_table(con, table)
    ident = table + "_" + sha256_bytes(f"{source_id}|{name}|{path}|{value[:100]}".encode("utf-8", "ignore"))[:16]
    quoted_table = '"' + table.replace('"', '""') + '"'
    con.execute(f"DELETE FROM {quoted_table} WHERE id=?", (ident,))
    con.execute(
        f"""INSERT INTO {quoted_table}(id, source_id, name, path, value, metadata_json, created_at)
        VALUES(?,?,?,?,?,?,?)""",
        (ident, source_id, name, path, value, json.dumps(meta or {}, ensure_ascii=False), now()),
    )
    return ident


def _fts_replace(con: sqlite3.Connection, table: str, entity_id: str, text: str) -> None:
    """Synchronize one generic FTS row by its stable entity identity."""
    if not text:
        return
    quoted_table = '"' + table.replace('"', '""') + '"'
    con.execute(f"DELETE FROM {quoted_table} WHERE entity_id=?", (entity_id,))
    con.execute(f"INSERT INTO {quoted_table}(entity_id,text) VALUES(?,?)", (entity_id, text))


def _clear_semantic_source_rows(con: sqlite3.Connection, lane_key: str, source_id: str) -> None:
    """Remove one changed source transactionally before writing its replacement."""

    tables = list(LANE_DEFS[lane_key]["schema"])
    if lane_key == "chat_lineage":
        tables.append("lineage_source")
    prefix = source_id + "_"
    for table in tables:
        quoted = '"' + table.replace('"', '""') + '"'
        exists = con.execute(
            "SELECT 1 FROM sqlite_master WHERE name=? AND type IN ('table','view')", (table,)
        ).fetchone()
        if not exists:
            continue
        if table.endswith("_fts"):
            con.execute(
                f"DELETE FROM {quoted} WHERE entity_id=? OR substr(entity_id,1,?)=?",
                (source_id, len(prefix), prefix),
            )
            continue
        columns = {row[1] for row in con.execute(f"PRAGMA table_info({quoted})")}
        if "source_id" in columns:
            con.execute(f"DELETE FROM {quoted} WHERE source_id=?", (source_id,))


def extract_docx_text(path: Path):
    paragraphs = []
    tables = []
    try:
        import docx
        doc = docx.Document(str(path))
        for p in doc.paragraphs:
            if p.text.strip(): paragraphs.append(p.text)
        for ti, t in enumerate(doc.tables):
            rows = []
            for r in t.rows:
                rows.append([c.text for c in r.cells])
            tables.append(rows)
    except Exception:
        try:
            with zipfile.ZipFile(path) as z:
                xml = z.read("word/document.xml")
                root = ET.fromstring(xml)
                texts = [n.text for n in root.iter() if n.tag.endswith('}t') and n.text]
                paragraphs = texts
        except Exception:
            pass
    return paragraphs, tables


def extract_pptx(path: Path):
    slides = []
    try:
        import pptx
        prs = pptx.Presentation(str(path))
        for si, slide in enumerate(prs.slides, start=1):
            texts=[]
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip(): texts.append(shape.text)
            slides.append((si, texts, ""))
    except Exception:
        try:
            with zipfile.ZipFile(path) as z:
                names = sorted([n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")])
                for si, name in enumerate(names, start=1):
                    root = ET.fromstring(z.read(name))
                    texts=[n.text for n in root.iter() if n.tag.endswith('}t') and n.text]
                    slides.append((si, texts, name))
        except Exception:
            pass
    return slides


def extract_xlsx(path: Path):
    sheets=[]
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), data_only=False, read_only=False)
        for ws in wb.worksheets:
            rows=[]; formulas=[]
            max_r = min(ws.max_row or 0, 50000)
            max_c = min(ws.max_column or 0, 200)
            for r in ws.iter_rows(min_row=1, max_row=max_r, max_col=max_c, values_only=False):
                vals=[]
                for cell in r:
                    vals.append(cell.value)
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formulas.append((cell.coordinate, cell.value))
                rows.append(vals)
            sheets.append((ws.title, ws.max_row, ws.max_column, rows, formulas))
    except Exception:
        try:
            with zipfile.ZipFile(path) as z:
                names=[n for n in z.namelist() if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")]
                for n in names:
                    sheets.append((Path(n).stem, 0, 0, [], []))
        except Exception:
            pass
    return sheets


def extract_pdf(path: Path):
    pages=[]
    try:
        import fitz
        doc = fitz.open(str(path))
        for i, page in enumerate(doc, start=1):
            pages.append((i, page.get_text("text") or "", len(page.get_images(full=True))))
        doc.close()
    except Exception:
        try:
            import pypdf
            reader = pypdf.PdfReader(str(path))
            for i, p in enumerate(reader.pages, start=1):
                pages.append((i, p.extract_text() or "", 0))
        except Exception:
            pass
    return pages


def build_semantic_lane(db: Path, lane_key: str, source: dict, progress: ProgressCallback, started: float):
    lane_key = resolve_lane_id(lane_key, scope="any")
    if lane_key == "research":
        raise RuntimeError("RESEARCH_REQUIRES_APPEND_ONLY_SERVICE")
    ensure_sector(db, lane_key)
    if lane_key in STRUCTURAL_LANE_INGESTERS:
        path_s = source.get("path") or ""
        if not path_s:
            raise RuntimeError(f"{lane_key.upper()}_SOURCE_PATH_REQUIRED")
        receipt = STRUCTURAL_LANE_INGESTERS[lane_key](Path(path_s), db)
        if receipt.status != "PASS":
            raise RuntimeError(f"{lane_key.upper()}_INGEST_FAILED:{receipt.status}")
        emit(
            progress,
            "build",
            f"{lane_key} structurally indexed",
            path_s,
            75,
            1,
            1,
            started,
        )
        return receipt
    con = connect(db)
    path_s = source.get("path") or ""
    text = source.get("text") or ""
    sid = source.get("source_id") or "source_" + sha256_bytes((path_s + text[:100]).encode())[:12]
    display = source.get("display_name") or Path(path_s).name or lane_key
    ext = Path(path_s).suffix.lower() if path_s else ".txt"
    digest = sha256_file(Path(path_s)) if path_s and Path(path_s).exists() else sha256_bytes(text.encode("utf-8", "ignore"))
    lane_contract = LANE_DEFS[lane_key]
    previous = con.execute(
        "SELECT source_hash,parser_id,chunker_version FROM ingestion_state WHERE source_id=?",
        (sid,),
    ).fetchone()
    if previous and tuple(previous) == (
        digest,
        lane_contract["parser_id"],
        lane_contract["chunker_version"],
    ):
        con.close()
        emit(progress, "build", f"{lane_key} unchanged; index-once skip", display, 75, 1, 1, started)
        return
    _clear_semantic_source_rows(con, lane_key, sid)
    _ginsert(con, f"{lane_key.split('_')[0]}_source" if lane_key in {"discussion","analysis","plan","mode"} else LANE_DEFS[lane_key]["schema"][0], sid, display, path_s, digest, {"source_type": source.get("source_type"), "schema_contract": source.get("schema_contract")})
    if lane_key == "chat_lineage":
        _ginsert(con, "lineage_source", sid, display, path_s, digest, {"compatibility": "pre_t021_read_only"})

    if lane_key in {"chat_lineage","discussion","analysis","plan","mode"}:
        if path_s and Path(path_s).exists():
            if ext == ".docx":
                paras, _ = extract_docx_text(Path(path_s)); text = "\n".join(paras)
            else:
                try: text = Path(path_s).read_text(encoding="utf-8", errors="replace")
                except Exception: text = ""
        lines = [l for l in text.splitlines()]
        prefix = {"chat_lineage":"lineage", "discussion":"discussion", "analysis":"analysis", "plan":"plan", "mode":"mode"}[lane_key]
        item_table = {"chat_lineage":"lineage_turn", "discussion":"discussion_item", "analysis":"analysis_claim", "plan":"plan_task", "mode":"mode_rule"}[lane_key]
        for i in range(0, len(lines), 80):
            chunk = "\n".join(lines[i:i+80]).strip()
            if chunk:
                _ginsert(con, item_table, sid, f"chunk_{i//80+1}", path_s, chunk, {"start_line": i+1, "end_line": i+len(lines[i:i+80])})
                fts = LANE_DEFS[lane_key]["fts_table"]
                if fts in LANE_DEFS[lane_key]["schema"]:
                    _fts_replace(con, fts, sid+f"_{i}", chunk)
    elif lane_key == "docs":
        paras=[]; tables=[]
        if ext == ".docx":
            parsed = parse_docx(Path(path_s))
            _ginsert(con, "doc_file", sid, display, path_s, digest, parsed["metadata"])
            for section in parsed["sections"]:
                _ginsert(con, "doc_structure", sid, section["id"], path_s, "section", section)
            for item in parsed["headers_footers"]:
                _ginsert(con, "doc_structure", sid, item["id"], path_s, item["text"], item)
            for item in parsed["hyperlinks"]:
                _ginsert(con, "doc_structure", sid, item["id"], path_s, item["target"], item)
            if parsed["comments_text"]:
                _ginsert(con, "doc_structure", sid, "comments", path_s, parsed["comments_text"], {"kind":"comments"})
            if parsed["footnotes_text"]:
                _ginsert(con, "doc_structure", sid, "footnotes", path_s, parsed["footnotes_text"], {"kind":"footnotes"})
            for item in parsed["paragraphs"]:
                if item["text"].strip():
                    _ginsert(con, "doc_paragraph", sid, item["id"], path_s, item["text"], item)
            for item in parsed["headings"]:
                _ginsert(con, "doc_heading", sid, item["id"], path_s, item["text"], item)
            for item in parsed["tables"]:
                _ginsert(con, "doc_table_extract", sid, item["id"], path_s, json.dumps(item["rows"], ensure_ascii=False), {key:value for key,value in item.items() if key != "rows"})
            for item in parsed["embedded_images"]:
                _ginsert(con, "doc_image_reference", sid, item["id"], path_s, item["sha256"], item)
            for item in parsed["chunks"]:
                _ginsert(con, "doc_chunk", sid, item["id"], path_s, item["text"], item)
                _fts_replace(con, "doc_fts", sid+"_"+item["id"], item["text"])
            paras=[item["text"] for item in parsed["paragraphs"]]
            tables=parsed["tables"]
            _ginsert(
                con,
                "source_structure_signature",
                sid,
                "DOCX_STRUCTURE",
                path_s,
                f"sections={len(parsed['sections'])} headings={len(parsed['headings'])} paragraphs={len(parsed['paragraphs'])} tables={len(parsed['tables'])} images={len(parsed['embedded_images'])}",
                {"parser":parsed["parser"]},
            )
        elif ext == ".pdf":
            pages = extract_pdf(Path(path_s)); paras = [p[1] for p in pages if p[1]]
        else:
            paras = Path(path_s).read_text(encoding="utf-8", errors="replace").splitlines() if path_s else text.splitlines()
        if ext != ".docx":
            _ginsert(con, "doc_file", sid, display, path_s, digest, {"paragraphs": len(paras), "tables": len(tables)})
            for i, p in enumerate(paras, start=1):
                if p.strip():
                    table = "doc_heading" if re.match(r"^#{1,6}\s+", p.strip()) else "doc_paragraph"
                    _ginsert(con, table, sid, f"block_{i}", path_s, p)
            for i in range(0, len(paras), 80):
                chunk="\n".join(paras[i:i+80]).strip()
                if chunk:
                    chunk_id=f"chunk_{i//80+1}"
                    _ginsert(con, "doc_chunk", sid, chunk_id, path_s, chunk, {"start":i+1})
                    _fts_replace(con, "doc_fts", sid+"_"+chunk_id, chunk)
            _ginsert(con, "source_structure_signature", sid, "DOC_STRUCTURE", path_s, f"paragraphs={len(paras)} tables={len(tables)}", {})
    elif lane_key == "data_excel":
        parsed_data = parse_data_source(Path(path_s))
        source_kind = parsed_data["source_kind"]
        status = parsed_data["status"]
        payload = parsed_data["payload"]
        _ginsert(con, "data_source", sid, display, path_s, digest, {"source_kind":source_kind,"status":status})
        if status != "SUPPORTED":
            _ginsert(con, "data_structure_signature", sid, "UNSUPPORTED", path_s, status, {"source_kind":source_kind})
        elif source_kind == "delimited":
            header = payload["header"]
            rows = payload["rows"]
            _ginsert(con, "csv_header", sid, "header", path_s, json.dumps(header, ensure_ascii=False), {"columns": len(header),"delimiter":payload["delimiter"]})
            for count,row in enumerate(rows, start=1):
                _ginsert(con, "csv_row_sample", sid, f"row_{count}", path_s, json.dumps(row, ensure_ascii=False))
            for offset in range(0,len(rows),250):
                group=rows[offset:offset+250]
                value=json.dumps(group,ensure_ascii=False)
                chunk_id=f"delimited_{offset+1}_{offset+len(group)}"
                _ginsert(con,"data_chunk",sid,chunk_id,path_s,value,{"source_kind":source_kind})
                _fts_replace(con,"data_fts",sid+"_"+chunk_id,value)
            _ginsert(con, "data_structure_signature", sid, "DELIMITED_SIGNATURE", path_s, f"columns={len(header)} rows={len(rows)}")
        elif source_kind == "workbook":
            workbook=payload
            _ginsert(con,"sheet_workbook",sid,display,path_s,digest,{**workbook["workbook"],"sheet_count":len(workbook["sheets"])})
            for sheet in workbook["sheets"]:
                cells=sheet.get("cells",[])
                meta={key:value for key,value in sheet.items() if key != "cells"}
                _ginsert(con,"sheet_tab",sid,sheet["id"],path_s,sheet["title"],meta)
                for cell in cells:
                    _ginsert(con,"sheet_cell_sample",sid,cell["id"],path_s,json.dumps(cell["value"],ensure_ascii=False,default=str),cell)
            for item in [*workbook["ranges"],*workbook["named_ranges"],*workbook["merged_cells"],*workbook["hidden_dimensions"],*workbook["validations"]]:
                item_id=item.get("id") or "range_"+sha256_bytes(json.dumps(item,sort_keys=True,default=str).encode())[:16]
                _ginsert(con,"sheet_range",sid,item_id,path_s,json.dumps(item,ensure_ascii=False,default=str),item)
            for item in workbook["tables"]:
                _ginsert(con,"sheet_table",sid,item["id"],path_s,item["range"],item)
            for item in workbook["formulas"]:
                _ginsert(con,"sheet_formula",sid,item["id"],path_s,item["formula"],item)
            for item in workbook["formula_dependencies"]:
                _ginsert(con,"sheet_formula_dependency_edge",sid,item["id"],path_s,item["to"],item)
            for item in workbook["charts"]:
                _ginsert(con,"sheet_chart_metadata",sid,item["id"],path_s,json.dumps(item["source_ranges"],ensure_ascii=False),item)
            for item in workbook["chunks"]:
                value=json.dumps(item["cells"],ensure_ascii=False,default=str)
                _ginsert(con,"data_chunk",sid,item["id"],path_s,value,{key:value_ for key,value_ in item.items() if key != "cells"})
                _fts_replace(con,"data_fts",sid+"_"+item["id"],value)
            _ginsert(con,"data_structure_signature",sid,"WORKBOOK_SIGNATURE",path_s,f"sheets={len(workbook['sheets'])} tables={len(workbook['tables'])} charts={len(workbook['charts'])} formulas={len(workbook['formulas'])} dependencies={len(workbook['formula_dependencies'])} named_ranges={len(workbook['named_ranges'])}",{"parser":workbook["parser"]})
        else:
            value=json.dumps(payload,ensure_ascii=False,default=str)
            _ginsert(con,"data_chunk",sid,source_kind,path_s,value,{"source_kind":source_kind})
            _fts_replace(con,"data_fts",sid+"_"+source_kind,value[:MAX_ARTIFACT_TEXT_EXTRACT])
            _ginsert(con,"data_structure_signature",sid,source_kind.upper()+"_SIGNATURE",path_s,f"bytes={len(value.encode('utf-8'))}")
    elif lane_key == "ppt":
        if ext != ".pptx":
            raise RuntimeError(f"PPT_FORMAT_UNSUPPORTED_CONVERTER_REQUIRED:{ext}")
        presentation=parse_pptx(Path(path_s))
        _ginsert(con,"ppt_file",sid,display,path_s,digest,{**presentation["metadata"],"slide_count":len(presentation["slides"])})
        shape_count=table_count=notes_count=image_count=relationship_count=0
        for slide in presentation["slides"]:
            slide_meta={key:value for key,value in slide.items() if key not in {"shapes","text_blocks","tables","images","notes","relationships"}}
            slide_text="\n".join(item["text"] for item in slide["text_blocks"])
            _ginsert(con,"ppt_slide",sid,slide["id"],path_s,slide_text,slide_meta)
            for item in slide["shapes"]:
                _ginsert(con,"ppt_shape",sid,item["id"],path_s,item.get("text") or item["name"],item)
                shape_count+=1
            for item in slide["text_blocks"]:
                _ginsert(con,"ppt_text_block",sid,item["id"],path_s,item["text"],item)
            if slide["notes"].strip():
                _ginsert(con,"ppt_notes",sid,slide["id"]+"_notes",path_s,slide["notes"],{"slide_id":slide["id"]})
                notes_count+=1
            for item in slide["tables"]:
                _ginsert(con,"ppt_table",sid,item["id"],path_s,json.dumps(item["rows"],ensure_ascii=False),{key:value for key,value in item.items() if key != "rows"})
                table_count+=1
            for item in slide["images"]:
                _ginsert(con,"ppt_image_reference",sid,item["id"],path_s,item["sha256"],item)
                image_count+=1
            for item in slide["relationships"]:
                relationship_id=stable_record_id("ppt_rel",slide["id"],item["relationship_id"],item["target"])
                _ginsert(con,"ppt_slide_relationship",sid,relationship_id,path_s,item["target"],{**item,"slide_id":slide["id"]})
                relationship_count+=1
        for item in presentation["chunks"]:
            _ginsert(con,"ppt_chunk",sid,item["id"],path_s,item["text"],item)
            _fts_replace(con,"ppt_fts",sid+"_"+item["id"],item["text"])
        _ginsert(con,"ppt_structure_signature",sid,"PPTX_STRUCTURE",path_s,f"slides={len(presentation['slides'])} shapes={shape_count} tables={table_count} notes={notes_count} images={image_count} relationships={relationship_count}",{"parser":presentation["parser"]})
    elif lane_key == "pdf_ocr":
        parsed_pdf = parse_pdf_ocr(Path(path_s))
        _ginsert(con, "pdf_file", sid, display, path_s, digest, parsed_pdf["metadata"])
        for page in parsed_pdf["pages"]:
            page_id = page["id"]
            _ginsert(con, "pdf_page", sid, page_id, path_s, page["classification"], {
                key: value for key, value in page.items() if key not in {"text_blocks", "image_regions", "ocr"}
            })
            fts_parts = []
            for item in page["text_blocks"]:
                _ginsert(con, "pdf_text_block", sid, item["id"], path_s, item["text"], {**item, "page_id": page_id})
                fts_parts.append(item["text"])
            for item in page["image_regions"]:
                _ginsert(con, "pdf_image_block", sid, item["id"], path_s, item["extension"], {**item, "page_id": page_id})
            ocr = page["ocr"]
            _ginsert(con, "pdf_ocr_run", sid, page_id + "_ocr", path_s, ocr["status"], {
                "page_id": page_id, "confidence": ocr["confidence"],
                "review_required": ocr["review_required"], "error": ocr["error"],
            })
            for item in ocr["blocks"]:
                _ginsert(con, "pdf_ocr_block", sid, item["id"], path_s, item["text"], {**item, "page_id": page_id})
                fts_parts.append(item["text"])
            for item in ocr["lines"]:
                _ginsert(con, "pdf_ocr_line", sid, item["id"], path_s, item["text"], {**item, "page_id": page_id})
            for item in ocr["regions"]:
                _ginsert(con, "pdf_review_region", sid, item["id"], path_s, item["reason"], {**item, "page_id": page_id})
            if ocr["review_required"] and not ocr["regions"]:
                _ginsert(con, "pdf_review_region", sid, page_id+"_review", path_s, "OCR_EMPTY_OR_ENGINE_UNAVAILABLE", {"page_id": page_id, "review_required": True})
            if fts_parts:
                _fts_replace(con, "pdf_fts", sid+"_"+page_id, "\n".join(fts_parts))
        _ginsert(con, "pdf_structure_signature", sid, "PDF_SIGNATURE", path_s, json.dumps(parsed_pdf["metadata"], sort_keys=True), {"parser": parsed_pdf["parser"], "review_required": parsed_pdf["review_required"]})
    elif lane_key == "images_ocr":
        parsed_image = parse_image_ocr(Path(path_s))
        _ginsert(con, "image_file", sid, display, path_s, digest, parsed_image["metadata"])
        _ginsert(con, "image_metadata", sid, display, path_s, json.dumps(parsed_image["metadata"], ensure_ascii=False, sort_keys=True))
        _ginsert(con, "image_ocr_run", sid, parsed_image["status"], path_s, parsed_image.get("error") or "", {
            "confidence": parsed_image["confidence"], "review_required": parsed_image["review_required"]
        })
        fts_parts=[]
        for item in parsed_image["blocks"]:
            _ginsert(con, "image_ocr_block", sid, item["id"], path_s, item["text"], item)
            fts_parts.append(item["text"])
        for item in parsed_image["lines"]:
            _ginsert(con, "image_ocr_line", sid, item["id"], path_s, item["text"], item)
        for item in parsed_image["regions"]:
            _ginsert(con, "image_review_region", sid, item["id"], path_s, item["reason"], item)
        if parsed_image["review_required"] and not parsed_image["regions"]:
            _ginsert(con, "image_review_region", sid, "whole_image_review", path_s, "OCR_EMPTY_OR_ENGINE_UNAVAILABLE", {"review_required": True})
        if fts_parts:
            _fts_replace(con, "image_ocr_fts", sid, "\n".join(fts_parts))
    elif lane_key == "artifacts":
        size = Path(path_s).stat().st_size if path_s else len(text.encode())
        _ginsert(con, "project_artifact", sid, display, path_s, digest, {"size": size, "status":"READ_ONLY_ARTIFACT"})
        _ginsert(con, "artifact_metadata", sid, display, path_s, f"size={size} ext={ext}")
        _ginsert(
            con,
            "artifact_relation_edge",
            sid,
            "BELONGS_TO_REGISTERED_SOURCE",
            path_s,
            sid,
            {
                "from_entity": sid,
                "to_entity": sid,
                "relation_type": "BELONGS_TO_REGISTERED_SOURCE",
                "confidence": "deterministic",
            },
        )
        if ext in TEXT_ARTIFACT_EXTS and path_s:
            tx = Path(path_s).read_text(encoding="utf-8", errors="replace")[:MAX_ARTIFACT_TEXT_EXTRACT]
            _ginsert(con, "artifact_text_extract", sid, display, path_s, tx)
            _fts_replace(con, "artifact_fts", sid, tx)
    else:
        # Missing/specialized lanes still write only to their own canonical
        # schema.  They must never silently fall back into the Custom sector.
        lane_contract = LANE_DEFS[lane_key]
        item_table = lane_contract["schema"][1] if len(lane_contract["schema"]) > 1 else lane_contract["schema"][0]
        val = text or (Path(path_s).read_text(encoding="utf-8", errors="replace") if path_s and Path(path_s).exists() and Path(path_s).suffix.lower() in TEXT_ARTIFACT_EXTS else "")
        if val:
            _ginsert(con, item_table, sid, display, path_s, val)
            _fts_replace(con, lane_contract["fts_table"], sid, val[:MAX_ARTIFACT_TEXT_EXTRACT])
    con.execute(
        "INSERT OR REPLACE INTO ingestion_state(source_id,source_hash,parser_id,chunker_version,completed_at) VALUES(?,?,?,?,?)",
        (sid,digest,lane_contract["parser_id"],lane_contract["chunker_version"],now()),
    )
    con.commit(); vacuum_close(con)
    emit(progress, "build", f"{lane_key} lane ingested", display, 75, 1, 1, started)


def brain_has_loaded_code_lanes(brain_root: str | Path) -> bool:
    """Return true only when an active GitHub or Local Code source is indexed."""
    root = Path(brain_root)
    for lane_id in ("github_code", "local_code"):
        try:
            _, database = resolve_env15_sector(root, lane_id)
            connection = sqlite3.connect(f"file:{database.as_posix()}?mode=ro", uri=True)
            try:
                row = connection.execute(
                    "SELECT 1 FROM source_registry WHERE availability_state='AVAILABLE' LIMIT 1"
                ).fetchone()
            finally:
                connection.close()
            if row:
                return True
        except (OSError, sqlite3.Error):
            continue
    return False


def materialize_schema_only_package_topology(brain_root: str | Path) -> dict[str, object]:
    """Materialize the live governed Project topology without firing a source lane.

    Env and UOP remain locked read authorities, but there is no third locked
    Project template.  This path is used only when no active code lane exists,
    so package validators can carry a topology derived from the live router and
    sector databases while every unloaded lane remains unexecuted.
    """

    root = Path(brain_root).resolve()
    rendered = render_env15_topologies(root)
    project = rendered["PROJECT"]
    outputs = {
        "mmd": str(project.mmd_path),
        "svg": str(project.svg_path or ""),
        "png": str(project.png_path or ""),
        "hd_png": str(project.hd_png_path or ""),
    }
    missing = [name for name, path in outputs.items() if not path or not Path(path).is_file()]
    if missing:
        raise RuntimeError("LIVE_PROJECT_TOPOLOGY_OUTPUT_MISSING:" + ",".join(missing))
    receipt = root / "receipts" / "schema_only_package_topology.json"
    receipt.parent.mkdir(parents=True, exist_ok=True)
    receipt.write_text(
        json.dumps(
            {
                "contract": "T023_SCHEMA_ONLY_PACKAGE_TOPOLOGY_V1",
                "status": "LIVE_PROJECT_TOPOLOGY_MATERIALIZED",
                "source_lane_execution": "NONE",
                "purpose": "PACKAGE_WHOLE_ENVIRONMENT_WITH_UNLOADED_LANES_UNFIRED",
                "files": outputs,
                "project_topology_authority": "LIVE_ROUTER_AND_SECTOR_SQLITE",
                "project_topology_template_lock": False,
                "created_at": now(),
            },
            indent=2,
            sort_keys=True,
        )
        + "\n",
        encoding="utf-8",
    )
    return {
        "status": "LIVE_PROJECT_TOPOLOGY_MATERIALIZED",
        "source_lane_execution": "NONE",
        "mmd_files": [outputs["mmd"]],
        "files": list(outputs.values()),
        "rendered": [
            {
                "source": outputs["mmd"],
                "svg": outputs["svg"],
                "png": outputs["png"],
                "hd_png": outputs["hd_png"],
                "status": "MMD_TO_SVG_TO_PNG_TO_HD_PNG_PASS",
                "render_origin": "LIVE_ROUTER_AND_SECTOR_SQLITE",
            }
        ],
        "receipt": str(receipt),
    }


def _clean_code_sector_seed(lane_id: str) -> dict[str, dict[str, list]]:
    resource_root = find_env15_resource_root()
    if not resource_root.is_dir():
        raise RuntimeError("ENV15_CODE_SECTOR_BASELINE_DIRECTORY_REQUIRED")
    baseline = (
        resource_root
        / "project"
        / "sectors"
        / lane_id
        / f"{lane_id}_sector_v001.sqlite"
    )
    if not baseline.is_file():
        raise RuntimeError(f"ENV15_CODE_SECTOR_BASELINE_MISSING:{lane_id}:{baseline}")
    connection = sqlite3.connect(f"file:{baseline.as_posix()}?mode=ro", uri=True)
    try:
        virtual_tables = {
            str(row[0])
            for row in connection.execute(
                "SELECT name FROM sqlite_schema WHERE type='table' "
                "AND UPPER(COALESCE(sql,'')) LIKE 'CREATE VIRTUAL TABLE%'"
            )
        }
        shadow_tables = {
            virtual + suffix
            for virtual in virtual_tables
            for suffix in _FTS_SHADOW_SUFFIXES
        }
        seeds: dict[str, dict[str, list]] = {}
        for (table,) in connection.execute(
            "SELECT name FROM sqlite_schema WHERE type='table' "
            "AND name NOT LIKE 'sqlite_%' ORDER BY name"
        ):
            table = str(table)
            if table in shadow_tables or table in virtual_tables:
                continue
            columns = [
                str(row[1])
                for row in connection.execute(f'PRAGMA table_info("{table}")')
            ]
            rows = [list(row) for row in connection.execute(f'SELECT * FROM "{table}"')]
            if rows:
                seeds[table] = {"columns": columns, "rows": rows}
        return seeds
    finally:
        connection.close()


def _describe_skipped_unloaded_lane(brain_root: Path, lane_id: str) -> dict:
    """Describe an omitted lane without opening it for write or changing bytes."""

    physical_sector, database = resolve_env15_sector(brain_root, lane_id)
    before_hash = sha256_file(database)
    read_only = sqlite3.connect(f"file:{database.as_posix()}?mode=ro", uri=True)
    try:
        tables = {
            str(row[0])
            for row in read_only.execute(
                "SELECT name FROM sqlite_schema WHERE type='table' AND name NOT LIKE 'sqlite_%'"
            )
        }
        registered_sources = 0
        if "source_registry" in tables:
            columns = {
                str(row[1])
                for row in read_only.execute("PRAGMA table_info(source_registry)")
            }
            if "lane_id" in columns:
                registered_sources = int(
                    read_only.execute(
                        "SELECT COUNT(*) FROM source_registry WHERE lane_id=?", (lane_id,)
                    ).fetchone()[0]
                )
            elif "lane_key" in columns:
                registered_sources = int(
                    read_only.execute(
                        "SELECT COUNT(*) FROM source_registry WHERE lane_key=?", (lane_id,)
                    ).fetchone()[0]
                )
            elif lane_id == physical_sector:
                registered_sources = int(
                    read_only.execute("SELECT COUNT(*) FROM source_registry").fetchone()[0]
                )
        integrity = tuple(row[0] for row in read_only.execute("PRAGMA integrity_check"))
    finally:
        read_only.close()
    after_hash = sha256_file(database)
    if before_hash != after_hash:
        raise RuntimeError(f"SKIPPED_LANE_BYTES_CHANGED:{lane_id}")
    return {
        "contract": "T023_SKIPPED_NO_SOURCE_PRESERVE_BYTES_V1",
        "status": "SKIPPED_NO_SOURCE",
        "lane_id": lane_id,
        "physical_sector_id": physical_sector,
        "database": str(database),
        "database_sha256": before_hash,
        "database_size_bytes": database.stat().st_size,
        "registered_source_count": registered_sources,
        "preservation_state": (
            "PRESERVED_PRIOR_DATA" if registered_sources else "PRESERVED_SCHEMA_READY_BYTES"
        ),
        "ingestion_lane_fired": False,
        "mutation_receipt_created": False,
        "integrity_check": list(integrity),
    }


def build_brain(
    workspace_dir: str,
    brain_name: str,
    sources: list[dict],
    progress: ProgressCallback=None,
    pipeline_progress: ProgressCallback=None,
    generate_mmd: bool=True,
):
    started=time.time()
    active_sources=[s for s in sources if s.get("active", True)]
    active_lane_id_set: set[str] = set()
    for source in active_sources:
        requested_lane = source.get("lane_key") or "custom"
        try:
            active_lane_id_set.add(resolve_lane_id(requested_lane, scope="any"))
        except UnknownLaneAliasError as exc:
            raise RuntimeError(f"UNKNOWN_LANE_ALIAS:{requested_lane}") from exc
    active_lane_ids=sorted(active_lane_id_set)
    validate_primary_code_mode(active_lane_ids)
    if pipeline_progress:
        pipeline_progress({"stage_id":"sqlite_project_router_sector_creation","stage_percent":0,"active_command":"initialize canonical project/router/sectors"})
    root = init_brain_layout(workspace_dir, brain_name)
    delta_sector_state = ensure_plan_delta_sector(root)
    research_append_schema_state = ensure_env15_research_append_schema(root)
    if pipeline_progress:
        pipeline_progress({"stage_id":"sqlite_project_router_sector_creation","stage_percent":100,"active_command":"canonical project/router/sectors created"})
    project = root / "project"
    router = project / "project_router.sqlite"
    unloaded_lane_ids=[lane_id for lane_id in LANE_DEFS if lane_id not in active_lane_ids]
    prior_activation_path = root / "receipts" / "lane_activation_receipt.json"
    prior_active_lane_ids: list[str] | None = None
    if prior_activation_path.is_file():
        try:
            prior_activation = json.loads(prior_activation_path.read_text(encoding="utf-8"))
            prior_active_lane_ids = sorted(
                str(lane_id) for lane_id in prior_activation.get("active_lane_ids", [])
            )
        except (OSError, json.JSONDecodeError, TypeError):
            prior_active_lane_ids = None
    selection_changed = prior_active_lane_ids is None or prior_active_lane_ids != active_lane_ids
    unloaded_lane_skips = [
        _describe_skipped_unloaded_lane(root, lane_id)
        for lane_id in unloaded_lane_ids
    ]
    # Compatibility field retained only as proof that the retired destructive
    # omitted-lane reconciler did not run.
    unloaded_code_lane_resets: list[dict] = []
    total=max(1,len(active_sources))
    ingestion_results=[]
    plan_delta_results=[]
    sources_reused=0
    sources_indexed=0
    emit(progress, "build", "initializing brain sectors", str(root), 2, 0, total, started)
    if pipeline_progress:
        pipeline_progress({"stage_id":"per_lane_parsing_chunking_indexing","stage_percent":0,"files_done":0,"files_total":len(active_sources),"active_command":"parse chunk and index canonical lanes"})
    for idx, source in enumerate(active_sources, start=1):
        requested_lane = source.get("lane_key") or "custom"
        try:
            lane_key = resolve_lane_id(requested_lane, scope="any")
        except UnknownLaneAliasError as exc:
            raise RuntimeError(f"UNKNOWN_LANE_ALIAS:{requested_lane}") from exc
        lane = LANE_DEFS[lane_key]
        physical_sector, db = resolve_env15_sector(root, lane_key)
        emit(progress, "build", f"building {lane['label']} sector", source.get("display_name") or source.get("path") or "", int(idx*10/total), idx, total, started)
        if str(source.get("canonical_payload_mode") or "") == "REFERENCE_ONLY":
            insert_source_registry(
                db,
                source,
                lane_key,
                str(source.get("canonical_payload_id") or ""),
            )
            ingestion_result = {
                "status": "SKIPPED_CANONICAL_PAYLOAD_REFERENCE_ONLY",
                "canonical_payload_id": str(source.get("canonical_payload_id") or ""),
                "canonical_payload_owner_source_id": str(
                    source.get("canonical_payload_owner_source_id") or ""
                ),
                "full_corpus_copy_created": False,
            }
        elif lane_key == "research":
            ingestion_result = append_env15_research_source(root, source)
        elif lane_key in STRUCTURAL_LANE_INGESTERS:
            ingestion_result = ingest_env15_structural_source(root, lane_key, source)
        elif lane_key in UNIVERSAL_INGESTION_LANES:
            ingestion_result = ingest_env15_universal_with_projection(root, lane_key, source)
        elif lane_key in {"local_code", "github_code"}:
            ingestion_result = ingest_env15_code_source(
                root,
                lane_key,
                source,
                legacy_builder=build_code_sector,
                progress=progress,
                started=started,
            )
        elif lane_key == "chat_lineage":
            ingestion_result = append_env15_chat_lineage_source(root, source)
        else:
            raise RuntimeError(f"ENV15_LANE_INGESTION_UNMAPPED:{lane_key}:{physical_sector}:{db}")
        plan_delta_result = None
        if lane_key == "plan":
            plan_delta_result = register_initial_plan_delta(root, source, ingestion_result)
            plan_delta_results.append(plan_delta_result)
        result_status = str(
            getattr(ingestion_result, "status", None)
            or (ingestion_result.get("status") if isinstance(ingestion_result, dict) else "PASS")
            or "PASS"
        )
        reused = result_status.startswith("SKIPPED") or result_status in {"REPLAY", "IDEMPOTENT_REPLAY"}
        sources_reused += int(reused)
        sources_indexed += int(not reused)
        ingestion_results.append({
            "lane_id": lane_key,
            "status": result_status,
            "reused": reused,
            **(
                {
                    "plan_delta_id": plan_delta_result["delta_id"],
                    "plan_delta_status": plan_delta_result["status"],
                }
                if plan_delta_result
                else {}
            ),
        })
        if pipeline_progress:
            pipeline_progress({
                "stage_id":"per_lane_parsing_chunking_indexing",
                "stage_percent":int(idx*100/total),
                "active_lane":lane_key,
                "active_source_id":str(source.get("source_id") or ""),
                "active_sector_id":physical_sector,
                "active_file":source.get("path") or source.get("display_name") or "",
                "files_done":idx,
                "files_total":len(active_sources),
                "active_command":f"{'reused unchanged' if reused else 'indexed'} {lane_key}",
            })
    if pipeline_progress:
        pipeline_progress({"stage_id":"pointer_router_hash_finalization","stage_percent":0,"files_done":len(active_sources),"files_total":len(active_sources),"active_command":"finalize pointers router and hashes"})
    state_changed = bool(sources_indexed or selection_changed)
    pointer_authority_missing = not (root / "project" / "pointers" / "INDEX.json").is_file()
    pointer_authority = None
    if state_changed or pointer_authority_missing:
        pointer_authority = finalize_project_pointers_and_hashes(
            root,
            active_lane_ids=active_lane_ids,
            ingestion_results=ingestion_results,
            brain_name=brain_name,
        )
    if pipeline_progress:
        pipeline_progress({"stage_id":"pointer_router_hash_finalization","stage_percent":100,"files_done":len(active_sources),"files_total":len(active_sources),"active_command":"pointers router and hashes finalized" if (state_changed or pointer_authority_missing) else "unchanged pointers and hashes reused"})
    code_lanes_loaded = brain_has_loaded_code_lanes(root)
    if generate_mmd and state_changed and code_lanes_loaded:
        generate_env15_topologies(root)
    (root/"receipts"/"build_receipt.md").write_text(f"# Build Receipt\n\nstatus=PASS\nbrain={brain_name}\nbuilder={APP_VERSION}\nsources={len(active_sources)}\ncreated={now()}\n", encoding="utf-8")
    lane_activation_receipt = root / "receipts" / "lane_activation_receipt.json"
    lane_activation_receipt.write_text(
        json.dumps(
            {
                "contract": "T023_SOURCE_LANE_ACTIVATION_V1",
                "registered_source_count": len(sources),
                "active_source_count": len(active_sources),
                "inactive_source_count": len(sources) - len(active_sources),
                "zero_lane_mode": not active_sources,
                "active_lane_ids": active_lane_ids,
                "unloaded_lane_ids": unloaded_lane_ids,
                "lane_states": {
                    lane_id: (
                        "LOADED_CURRENT_BUILD"
                        if lane_id in active_lane_ids
                        else "SYSTEM_APPEND_READY"
                        if lane_id in {"chat_lineage", "research"}
                        else "SKIPPED_NO_SOURCE"
                    )
                    for lane_id in LANE_DEFS
                },
                "canonical_payloads": [
                    {
                        "canonical_payload_id": payload_id,
                        "owner_source_id": next(
                            (
                                str(source.get("canonical_payload_owner_source_id") or "")
                                for source in active_sources
                                if str(source.get("canonical_payload_id") or "") == payload_id
                            ),
                            "",
                        ),
                        "source_ids": sorted(
                            str(source.get("source_id") or "")
                            for source in active_sources
                            if str(source.get("canonical_payload_id") or "") == payload_id
                        ),
                        "reference_only_source_ids": sorted(
                            str(source.get("source_id") or "")
                            for source in active_sources
                            if str(source.get("canonical_payload_id") or "") == payload_id
                            and str(source.get("canonical_payload_mode") or "") == "REFERENCE_ONLY"
                        ),
                    }
                    for payload_id in sorted(
                        {
                            str(source.get("canonical_payload_id") or "")
                            for source in active_sources
                            if str(source.get("canonical_payload_id") or "")
                        }
                    )
                ],
                "duplicate_full_corpus_count": 0,
                "unloaded_code_lane_resets": unloaded_code_lane_resets,
                "unloaded_lane_skips": unloaded_lane_skips,
                "omitted_lane_mutation_count": 0,
                "selection_changed": selection_changed,
                "universal_lane_pointer_index": "project/pointers/INDEX.json",
                "created_at": now(),
            },
            indent=2,
            sort_keys=True,
        )
        + "\n",
        encoding="utf-8",
    )
    emit(progress, "done", "brain build complete", str(root), 100, total, total, started)
    return {
        "brain_root": str(root),
        "router": str(router),
        "incremental": {
            "sources_total": len(active_sources),
            "registered_sources_total": len(sources),
            "zero_lane_mode": not active_sources,
            "active_lane_ids": active_lane_ids,
            "unloaded_lane_ids": unloaded_lane_ids,
            "lane_activation_receipt": str(lane_activation_receipt),
            "sources_indexed": sources_indexed,
            "sources_reused": sources_reused,
            "all_sources_unchanged": (
                bool(active_sources)
                and sources_reused == len(active_sources)
                and not selection_changed
            ),
            "ingestion_results": ingestion_results,
            "unloaded_code_lane_resets": unloaded_code_lane_resets,
            "unloaded_lane_skips": unloaded_lane_skips,
            "omitted_lane_mutation_count": 0,
            "selection_changed": selection_changed,
            "canonical_payload_reference_count": sum(
                str(source.get("canonical_payload_mode") or "") == "REFERENCE_ONLY"
                for source in active_sources
            ),
            "duplicate_full_corpus_count": 0,
            "pointer_hashes_reused": not state_changed and not pointer_authority_missing,
            "pointer_authority": pointer_authority,
            "code_lanes_loaded": code_lanes_loaded,
            "topology_status": (
                "GENERATED" if generate_mmd and state_changed and code_lanes_loaded
                else "REUSED" if code_lanes_loaded and not state_changed
                else "DEFERRED" if code_lanes_loaded and not generate_mmd
                else "SKIPPED_NO_CODE_LANES"
            ),
            "topologies_reused": code_lanes_loaded and not state_changed,
            "plan_delta": {
                "sector_status": delta_sector_state["status"],
                "initial_delta_count": len(plan_delta_results),
                "goal_prompt_status": (
                    plan_delta_results[-1]["goal_prompt"]["status"]
                    if plan_delta_results
                    else "NO_PLAN_SOURCE"
                ),
                "goal_prompt_path": (
                    plan_delta_results[-1]["goal_prompt"]["path"]
                    if plan_delta_results
                    else None
                ),
                "refresh_ledger_relationship": "DISTINCT_DO_NOT_MIX",
            },
            "research_append_schema": research_append_schema_state,
        },
    }


def mmd_safe(s: str, limit=80):
    s=str(s or "").replace('"',"'").replace("["," ").replace("]"," ").replace("{"," ").replace("}"," ").replace("|","/")
    s=re.sub(r"\s+"," ",s).strip()
    return ("..."+s[-limit:]) if len(s)>limit else (s or "none")


def rows(con, sql, args=()):
    try: return con.execute(sql,args).fetchall()
    except Exception: return []


def table_exists(con: sqlite3.Connection, table: str) -> bool:
    return bool(rows(con, "SELECT 1 FROM sqlite_master WHERE type='table' AND name=? LIMIT 1", (table,)))


def table_count(con: sqlite3.Connection, table: str) -> int:
    if not table_exists(con, table):
        return 0
    try:
        return int(con.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0])
    except Exception:
        return 0


def mmd_label(*parts, limit=90) -> str:
    return "\\n".join(mmd_safe(part, limit) for part in parts if str(part or "").strip())


def project_code_sector_dbs(root: Path) -> list[Path]:
    sectors = root / "project" / "sectors"
    candidates = [
        sectors / "local_code" / "local_code_sector_v001.sqlite",
        sectors / "github_code" / "github_code_sector_v001.sqlite",
    ]
    if sectors.exists():
        candidates.extend(sorted(sectors.rglob("*code*_sector_v001.sqlite")))
        candidates.extend(sorted(sectors.rglob("*github*_sector_v001.sqlite")))
    seen = set()
    out = []
    for db in candidates:
        key = str(db.resolve()) if db.exists() else str(db)
        if db.exists() and key not in seen:
            seen.add(key)
            out.append(db)
    return out


def collect_code_topology(root: Path) -> dict:
    totals = {
        "source_file": 0,
        "code_file": 0,
        "code_symbol": 0,
        "app_route": 0,
        "code_import_edge": 0,
        "dependency_item": 0,
        "project_artifact": 0,
        "git_commit": 0,
        "code_chunk": 0,
    }
    routes_sample = []
    files_sample = []
    artifact_sample = []
    dependency_sample = []
    file_lookup = {}
    symbol_lookup = {}
    import_lookup = {}
    dbs = project_code_sector_dbs(root)
    for db in dbs:
        con = sqlite3.connect(db)
        try:
            for table in totals:
                totals[table] += table_count(con, table)
            for fid, path in rows(con, "SELECT file_id, canonical_path FROM code_file ORDER BY canonical_path LIMIT 200"):
                file_lookup.setdefault(fid, path)
            if len(routes_sample) < 14:
                for route_path, route_type, file_id in rows(con, "SELECT route_path, route_type, file_id FROM app_route ORDER BY route_path LIMIT 14"):
                    routes_sample.append((route_path, route_type, file_id, db.name))
                    if len(routes_sample) >= 14:
                        break
            if len(files_sample) < 12:
                file_rows = rows(con, """
                    SELECT cf.file_id, cf.canonical_path, cf.language,
                           COUNT(DISTINCT cs.symbol_id) + COUNT(DISTINCT ci.edge_id) AS graph_degree
                    FROM code_file cf
                    LEFT JOIN code_symbol cs ON cs.file_id = cf.file_id
                    LEFT JOIN code_import_edge ci ON ci.from_file_id = cf.file_id
                    GROUP BY cf.file_id, cf.canonical_path, cf.language
                    ORDER BY graph_degree DESC, cf.canonical_path
                    LIMIT 12
                """)
                if not file_rows:
                    file_rows = [(file_id, canonical_path, language, 0) for file_id, canonical_path, language in rows(con, "SELECT file_id, canonical_path, language FROM code_file ORDER BY canonical_path LIMIT 12")]
                for file_id, canonical_path, language, _graph_degree in file_rows:
                    files_sample.append((file_id, canonical_path, language, db.name))
                    if len(files_sample) >= 12:
                        break
            if len(artifact_sample) < 8:
                for artifact_type, count, size_bytes in rows(con, "SELECT artifact_type,COUNT(*),SUM(size_bytes) FROM project_artifact GROUP BY artifact_type ORDER BY COUNT(*) DESC LIMIT 8"):
                    artifact_sample.append((artifact_type, count, size_bytes or 0, db.name))
                    if len(artifact_sample) >= 8:
                        break
            if len(dependency_sample) < 8:
                for package_name, version_spec in rows(con, "SELECT package_name,version_spec FROM dependency_item ORDER BY package_name LIMIT 8"):
                    dependency_sample.append((package_name, version_spec, db.name))
                    if len(dependency_sample) >= 8:
                        break
            sample_file_ids = [item[2] for item in routes_sample if item[2]] + [item[0] for item in files_sample if item[0]]
            for file_id in sample_file_ids[:20]:
                if file_id not in symbol_lookup:
                    symbol_lookup[file_id] = rows(con, "SELECT symbol_name,symbol_type FROM code_symbol WHERE file_id=? ORDER BY start_line LIMIT 3", (file_id,))
                if file_id not in import_lookup:
                    import_lookup[file_id] = rows(con, "SELECT import_target,import_type FROM code_import_edge WHERE from_file_id=? LIMIT 3", (file_id,))
        finally:
            con.close()
    return {
        "dbs": dbs,
        "totals": totals,
        "routes_sample": routes_sample,
        "files_sample": files_sample,
        "artifact_sample": artifact_sample,
        "dependency_sample": dependency_sample,
        "file_lookup": file_lookup,
        "symbol_lookup": symbol_lookup,
        "import_lookup": import_lookup,
    }


def remove_legacy_local_code_topology(root: Path) -> None:
    topology = root / "project" / "topology"
    for name in (
        "local_code_lane.mmd",
        "local_code_lane.svg",
        "local_code_lane.png",
        "local_code_lane_CRYSTAL.png",
    ):
        path = topology / name
        if path.exists():
            path.unlink()


def write_code_mmd(root: Path):
    remove_legacy_local_code_topology(root)
    return
    db = root/"project"/"sectors"/"local_code"/"local_code_sector_v001.sqlite"
    if not db.exists(): return
    con=sqlite3.connect(db)
    topology = root/"project"/"topology"; topology.mkdir(parents=True, exist_ok=True)
    routes=rows(con,"SELECT route_path, route_type, file_id FROM app_route ORDER BY route_path LIMIT 80")
    deps=rows(con,"SELECT package_name,version_spec FROM dependency_item ORDER BY package_name LIMIT 25")
    arts=rows(con,"SELECT artifact_type,COUNT(*),SUM(size_bytes) FROM project_artifact GROUP BY artifact_type ORDER BY COUNT(*) DESC LIMIT 20")
    lines=["flowchart LR", "  classDef route fill:#fff4d6,stroke:#8a5a00,color:#111;", "  classDef code fill:#eaffea,stroke:#275,color:#111;", "  classDef dep fill:#fef3c7,stroke:#92400e,color:#111;", "  classDef artifact fill:#f3e8ff,stroke:#635,color:#111;", "  classDef root fill:#111827,stroke:#111827,color:#fff;", '  ROOT["coded project workflow"]:::root', '  ROUTES["routes / pages"]:::route', '  CODE["code files / line snapshots / chunks"]:::code', '  DEPS["tools + dependencies"]:::dep', '  ARTS["read-only artifacts"]:::artifact', "  ROOT --> ROUTES", "  ROUTES --> CODE", "  CODE --> DEPS", "  CODE --> ARTS", "  subgraph R[route/page to code path map]", "    direction TB"]
    prev="ROUTES"
    for i,(rp,rt,fid) in enumerate(routes):
        fp=rows(con,"SELECT canonical_path FROM code_file WHERE file_id=?",(fid,))
        path=fp[0][0] if fp else fid
        rn=f"R{i}"; fn=f"F{i}"
        lines += [f'    {rn}["{mmd_safe(rp,70)}"]:::route', f'    {fn}["{mmd_safe(path,90)}"]:::code', f"    {prev} --> {rn}", f"    {rn} --> {fn}"]
        imports=rows(con,"SELECT import_target FROM code_import_edge WHERE from_file_id=? LIMIT 4",(fid,))
        for j,(imp,) in enumerate(imports):
            inn=f"I{i}_{j}"; lines += [f'    {inn}["{mmd_safe(imp,50)}"]:::code', f"    {fn} --> {inn}"]
        prev=rn
    lines += ["  end", "  subgraph D[dependency map]", "    direction TB"]
    prev="DEPS"
    for i,(name,ver) in enumerate(deps):
        dn=f"D{i}"; lines += [f'    {dn}["{mmd_safe(name,60)}\\n{mmd_safe(ver,30)}"]:::dep', f"    {prev} --> {dn}"]; prev=dn
    lines += ["  end", "  subgraph A[artifact payload groups]", "    direction TB"]
    prev="ARTS"
    for i,(typ,cnt,sz) in enumerate(arts):
        an=f"A{i}"; lines += [f'    {an}["{mmd_safe(typ,50)}\\ncount={cnt} bytes={sz or 0}"]:::artifact', f"    {prev} --> {an}"]; prev=an
    lines += ["  end"]
    (topology/"local_code_lane.mmd").write_text("\n".join(lines)+"\n", encoding="utf-8")
    con.close()


def write_project_master_mmd(root: Path):
    router=root/"project"/"project_router.sqlite"
    if not router.exists(): return
    con=sqlite3.connect(router)
    try:
        ensure_project_mutation_schema(con)
        sectors=rows(con,"SELECT lane_key,lane_label,sector_db_path FROM sector_registry ORDER BY lane_label")
        sources=rows(con,"SELECT lane_key,display_name,source_type,path FROM source_registry ORDER BY created_at LIMIT 18")
        mutation_lanes=rows(con,"SELECT lane_id,lane_label,target_table FROM project_mutation_lane_registry ORDER BY lane_id")
        router_counts={
            "sources": table_count(con, "source_registry"),
            "sectors": table_count(con, "sector_registry"),
            "mutation_lanes": table_count(con, "project_mutation_lane_registry"),
            "write_contracts": table_count(con, "project_public_ai_write_contract"),
        }
        con.commit()
    finally:
        con.close()
    code=collect_code_topology(root)
    totals=code["totals"]
    topology=root/"project"/"topology"; topology.mkdir(parents=True, exist_ok=True)
    code_db_names=", ".join(db.name for db in code["dbs"]) or "none"
    lines=[
        "flowchart TD",
        "  classDef root fill:#111827,stroke:#111827,color:#fff;",
        "  classDef locked fill:#ffe8e8,stroke:#922,color:#111;",
        "  classDef router fill:#e0f2fe,stroke:#0369a1,color:#111;",
        "  classDef db fill:#dcfce7,stroke:#166534,color:#111;",
        "  classDef route fill:#fef3c7,stroke:#92400e,color:#111;",
        "  classDef code fill:#ede9fe,stroke:#6d28d9,color:#111;",
        "  classDef artifact fill:#fce7f3,stroke:#9d174d,color:#111;",
        "  classDef mutate fill:#ffedd5,stroke:#c2410c,color:#111;",
        "  classDef export fill:#f8fafc,stroke:#475569,color:#111;",
        f'  START["{mmd_label("project package", "sources="+str(router_counts["sources"]), "sectors="+str(router_counts["sectors"]))}"]:::root',
        '  LAW["locked Env + UOP + public template\\nread-only governance"]:::locked',
        f'  ROUTER["{mmd_label("project_router.sqlite", "generated project truth", "mutation_lanes="+str(router_counts["mutation_lanes"]))}"]:::router',
        f'  CODEDB["{mmd_label("code sector SQLite", code_db_names, "files="+str(totals["code_file"])+" symbols="+str(totals["code_symbol"]))}"]:::db',
        '  START --> LAW',
        '  START --> ROUTER',
        '  LAW -. policy boundary .-> ROUTER',
        '  ROUTER --> CODEDB',
        "",
        "  subgraph SOURCE_INTAKE[1. source intake to router]",
        "    direction TB",
        f'    SRCREG["{mmd_label("source_registry", "rows="+str(router_counts["sources"]))}"]:::router',
        f'    SECREG["{mmd_label("sector_registry", "rows="+str(router_counts["sectors"]))}"]:::router',
        "    ROUTER --> SRCREG",
        "    SRCREG --> SECREG",
    ]
    prev="SRCREG"
    if sources:
        for i,(lane_key,display_name,source_type,path) in enumerate(sources):
            node=f"SRC{i}"
            label=mmd_label(display_name or path or lane_key, lane_key, source_type or "source", limit=72)
            lines += [f'    {node}["{label}"]:::router', f"    {prev} --> {node}"]
            prev=node
    else:
        lines += ['    SRC_EMPTY["no selected sources recorded yet"]:::router', "    SRCREG --> SRC_EMPTY"]
    lines += [
        "  end",
        "",
        "  subgraph PROJECT_LANES[2. generated project DB lanes]",
        "    direction TB",
        "    SECREG --> LANE_HEAD[\"sector DB + pointer index\"]:::db",
    ]
    prev="LANE_HEAD"
    if sectors:
        for i,(lane_key,lane_label,sector_db_path) in enumerate(sectors):
            node=f"LANE{i}"
            extra="code topology source" if lane_key in {"local_code","github_code"} else "schema-ready lane"
            label=mmd_label(lane_label or lane_key, lane_key, extra, limit=78)
            lines += [f'    {node}["{label}"]:::db', f"    {prev} --> {node}"]
            prev=node
    else:
        lines += ['    LANE_EMPTY["no sector rows yet"]:::db', "    LANE_HEAD --> LANE_EMPTY"]
    lines += [
        "  end",
        "",
        "  subgraph CODE_GRAPH[3. route/file/symbol/import graph from SQLite]",
        "    direction TB",
        f'    CODE_SUM["{mmd_label("local code sector", "source_files="+str(totals["source_file"])+" chunks="+str(totals["code_chunk"]), "routes="+str(totals["app_route"])+" imports="+str(totals["code_import_edge"]))}"]:::code',
        "    CODEDB --> CODE_SUM",
    ]
    route_prev="CODE_SUM"
    if code["routes_sample"]:
        for i,(route_path,route_type,file_id,db_name) in enumerate(code["routes_sample"]):
            route_node=f"ROUTE{i}"
            file_node=f"RFILE{i}"
            file_path=code["file_lookup"].get(file_id, file_id)
            lines += [
                f'    {route_node}["{mmd_label(route_path, route_type, limit=74)}"]:::route',
                f'    {file_node}["{mmd_label(file_path, db_name, limit=84)}"]:::code',
                f"    {route_prev} --> {route_node}",
                f"    {route_node} --> {file_node}",
            ]
            for j,(sym_name,sym_type) in enumerate(code["symbol_lookup"].get(file_id, [])[:2]):
                sym_node=f"RSYM{i}_{j}"
                lines += [f'    {sym_node}["{mmd_label(sym_name, sym_type, limit=60)}"]:::code', f"    {file_node} --> {sym_node}"]
            for j,(import_target,import_type) in enumerate(code["import_lookup"].get(file_id, [])[:2]):
                imp_node=f"RIMP{i}_{j}"
                lines += [f'    {imp_node}["{mmd_label(import_target, import_type, limit=60)}"]:::code', f"    {file_node} --> {imp_node}"]
            route_prev=route_node
    elif code["files_sample"]:
        for i,(file_id,canonical_path,language,db_name) in enumerate(code["files_sample"]):
            file_node=f"FILE{i}"
            lines += [f'    {file_node}["{mmd_label(canonical_path, language, limit=84)}"]:::code', f"    {route_prev} --> {file_node}"]
            for j,(sym_name,sym_type) in enumerate(code["symbol_lookup"].get(file_id, [])[:2]):
                sym_node=f"FSYM{i}_{j}"
                lines += [f'    {sym_node}["{mmd_label(sym_name, sym_type, limit=60)}"]:::code', f"    {file_node} --> {sym_node}"]
            for j,(import_target,import_type) in enumerate(code["import_lookup"].get(file_id, [])[:2]):
                imp_node=f"FIMP{i}_{j}"
                lines += [f'    {imp_node}["{mmd_label(import_target, import_type, limit=60)}"]:::code', f"    {file_node} --> {imp_node}"]
            route_prev=file_node
    else:
        lines += ['    CODE_EMPTY["code sector not populated yet"]:::code', "    CODE_SUM --> CODE_EMPTY"]
    lines += [
        "  end",
        "",
        "  subgraph DATA_ARTIFACTS[4. project artifacts and dependencies]",
        "    direction TB",
        f'    ART_SUM["{mmd_label("project_artifact", "rows="+str(totals["project_artifact"]))}"]:::artifact',
        f'    DEP_SUM["{mmd_label("dependency/import surface", "dependency_items="+str(totals["dependency_item"]), "import_edges="+str(totals["code_import_edge"]))}"]:::artifact',
        f'    GIT_SUM["{mmd_label("git lineage", "commits="+str(totals["git_commit"]))}"]:::artifact',
        "    CODE_SUM --> ART_SUM",
        "    CODE_SUM --> DEP_SUM",
        "    CODE_SUM --> GIT_SUM",
    ]
    prev="ART_SUM"
    for i,(artifact_type,count,size_bytes,db_name) in enumerate(code["artifact_sample"]):
        node=f"ART{i}"
        lines += [f'    {node}["{mmd_label(artifact_type, "count="+str(count)+" bytes="+str(size_bytes), limit=66)}"]:::artifact', f"    {prev} --> {node}"]
        prev=node
    prev="DEP_SUM"
    for i,(package_name,version_spec,db_name) in enumerate(code["dependency_sample"]):
        node=f"DEP{i}"
        lines += [f'    {node}["{mmd_label(package_name, version_spec or "declared", limit=66)}"]:::artifact', f"    {prev} --> {node}"]
        prev=node
    lines += [
        "  end",
        "",
        "  subgraph MUTATION_LANES[5. generated project mutation lanes for public models]",
        "    direction TB",
        '    PUBLIC_AI["public model project write request"]:::mutate',
        f'    CONTRACT["{mmd_label("project_public_ai_write_contract", "rows="+str(router_counts["write_contracts"]), "Env/UOP/template denied")}"]:::mutate',
        "    ROUTER --> CONTRACT",
        "    PUBLIC_AI --> CONTRACT",
    ]
    prev="CONTRACT"
    for i,(lane_id,lane_label,target_table) in enumerate(mutation_lanes):
        node=f"MUT{i}"
        label=mmd_label(lane_label, target_table, "append/insert lane", limit=72)
        lines += [f'    {node}["{label}"]:::mutate', f"    {prev} --> {node}"]
        prev=node
    lines += [
        "    MUT_DONE[\"validation receipts + model writeback\"]:::mutate",
        f"    {prev} --> MUT_DONE",
        "    MUT_DONE --> ROUTER",
        "  end",
        "",
        "  subgraph OUTPUTS[6. topology proof and one-click exports]",
        "    direction TB",
        '    MMD["project/topology/project_master_topology.mmd\\nsingle project truth"]:::export',
        '    RENDER["SVG/PNG render manifest"]:::export',
        '    CHATGPT["ChatGPT_LocalAI_<brain>_Sqlite_brain.zip"]:::export',
        '    GEMINI["Gemini_<brain>_Sqlite_brain.zip"]:::export',
        "    CODE_SUM --> MMD",
        "    MMD --> RENDER",
        "    RENDER --> CHATGPT",
        "    RENDER --> GEMINI",
        "  end",
    ]
    (topology/"project_master_topology.mmd").write_text("\n".join(lines)+"\n", encoding="utf-8")


def generate_project_mmd(workspace_dir: str, brain_name: str, progress: ProgressCallback=None):
    root=brain_output_dir(workspace_dir, brain_name)
    if not brain_has_loaded_code_lanes(root):
        emit(progress,"mmd","topology skipped: no active code lanes",str(root),100,0,0,time.time())
        return {
            "status":"SKIPPED_NO_CODE_LANES",
            "reason":"PROJECT_TOPOLOGY_REQUIRES_ACTIVE_GITHUB_OR_LOCAL_CODE_SOURCE",
            "topology":None,
            "mmd_files":[],
        }
    if (root/".uepc_env").is_file():
        generated=generate_env15_topologies(root)
        mmds=[Path(item.mmd_path) for item in generated.values()]
        for idx,mmd in enumerate(mmds, start=1):
            emit(progress,"mmd",f"generated {mmd.name}",str(mmd),int(idx*100/max(1,len(mmds))),idx,len(mmds),time.time())
        return {"topology":str(root),"mmd_files":[str(path) for path in mmds]}
    write_code_mmd(root); write_project_master_mmd(root)
    topology=root/"project"/"topology"
    mmds=sorted(topology.glob("*.mmd"))
    for idx,mmd in enumerate(mmds, start=1):
        emit(progress,"mmd",f"generated {mmd.name}",str(mmd),int(idx*100/max(1,len(mmds))),idx,len(mmds),time.time())
    return {"topology":str(topology),"mmd_files":[str(path) for path in mmds]}


def render_topology(workspace_dir: str, brain_name: str, progress: ProgressCallback=None, generate_mmd: bool=True):
    root=brain_output_dir(workspace_dir, brain_name)
    if not brain_has_loaded_code_lanes(root):
        emit(progress,"render","topology render skipped: no active code lanes",str(root),100,0,0,time.time())
        return {
            "status":"SKIPPED_NO_CODE_LANES",
            "reason":"PROJECT_TOPOLOGY_REQUIRES_ACTIVE_GITHUB_OR_LOCAL_CODE_SOURCE",
            "topology":None,
            "rendered":[],
        }
    if (root/".uepc_env").is_file():
        rendered=render_env15_topologies(root)
        manifest=[]
        for idx,item in enumerate(rendered.values(), start=1):
            row={
                "mmd":item.mmd_path,"svg":item.svg_path,"png":item.png_path,
                "hd_png":item.hd_png_path,"status":item.render_status
            }
            manifest.append(row)
            emit(progress,"render",f"rendered {Path(item.mmd_path).name}",item.mmd_path,int(idx*100/len(rendered)),idx,len(rendered),time.time())
        manifest_path=root/"receipts"/"ENV15_RENDER_MANIFEST.json"
        manifest_path.write_text(json.dumps(manifest,indent=2),encoding="utf-8")
        return {"topology":str(root),"rendered":manifest,"manifest":str(manifest_path)}
    if generate_mmd:
        generate_project_mmd(workspace_dir, brain_name, progress)
    topology=root/"project"/"topology"
    mmds=list(topology.glob("*.mmd"))
    # Prefer the Windows command shim; PowerShell's mmdc.ps1 cannot be started
    # directly by subprocess with shell=False.
    mmdc=shutil.which("mmdc.cmd") or shutil.which("mmdc.exe") or shutil.which("mmdc")
    manifest=[]
    for idx,mmd in enumerate(mmds, start=1):
        item={"mmd":str(mmd), "status":"MMD_ONLY_RENDERER_NOT_FOUND"}
        if mmdc:
            svg=mmd.with_suffix(".svg"); png=mmd.with_suffix(".png"); crystal=mmd.with_name(mmd.stem+"_CRYSTAL.png")
            r=run_hidden([mmdc,"-i",str(mmd),"-o",str(svg),"-b","white"],capture_output=True,text=True,timeout=1200)
            if r.returncode==0:
                run_hidden([mmdc,"-i",str(mmd),"-o",str(crystal),"-b","white","-w","7680","-H","4320"],capture_output=True,text=True,timeout=1800)
                if crystal.exists(): shutil.copy2(crystal,png)
                item={"mmd":str(mmd),"svg":str(svg),"png":str(png),"crystal_png":str(crystal),"status":"MMD_TO_SVG_TO_PNG_TO_HD_PNG_PASS"}
            else:
                item={"mmd":str(mmd),"status":"SVG_RENDER_FAILED","stderr":r.stderr[-1000:]}
        manifest.append(item)
        emit(progress,"render",f"rendered {mmd.name}",str(mmd),int(idx*100/max(1,len(mmds))),idx,len(mmds),time.time())
    (topology/"render_manifest.json").write_text(json.dumps(manifest,indent=2),encoding="utf-8")
    return {"topology":str(topology),"rendered":manifest}


def safe_sqlite_backup(src: Path, dst: Path):
    dst.parent.mkdir(parents=True, exist_ok=True)
    try:
        s=sqlite3.connect(str(src), timeout=60)
        d=sqlite3.connect(str(dst), timeout=60)
        s.backup(d)
        d.commit(); d.close(); s.close()
    except Exception:
        shutil.copy2(src,dst)


def copy_tree_safe(src: Path, dst: Path):
    for p in src.rglob("*"):
        if not p.is_file(): continue
        if p.name.endswith(("-wal","-shm")): continue
        rel=p.relative_to(src)
        target=dst/rel
        target.parent.mkdir(parents=True, exist_ok=True)
        exact_env_uop_authority = (
            (p.parent.name.casefold(), p.name.casefold())
            in {
                ("env", "env_sqlite.sqlite"),
                ("uop", "uop_sqlite.sqlite"),
            }
        )
        if exact_env_uop_authority:
            shutil.copy2(p, target)
        elif p.suffix.lower() in {".sqlite", ".db"}:
            safe_sqlite_backup(p,target)
        else:
            shutil.copy2(p,target)


def find_env_source() -> Optional[Path]:
    try:
        source = find_env15_resource_root()
        validation = validate_env15_resource(source)
        if validation.structural_validation_passed:
            return source
    except Exception:
        return None
    return None


def copy_locked_env(stage: Path):
    src=find_env_source()
    if not src: raise RuntimeError("PUBLIC_MODEL_ENV15_RESOURCE_NOT_FOUND_OR_INVALID")
    if src.is_file() and src.suffix.lower()==".zip":
        with zipfile.ZipFile(src) as z: z.extractall(stage)
    else:
        copy_tree_safe(src, stage)


def write_flash_prompt(stage: Path, brain_name: str):
    prompt = LOCKED_FLASH_PROMPT
    (stage/"FLASH_ME_FIRST_SINGLE_PROMPT.txt").write_text(prompt, encoding="utf-8")
    (stage/"prompts").mkdir(parents=True, exist_ok=True)
    (stage/"prompts"/"FLASH_ME_FIRST_SINGLE_PROMPT.txt").write_text(prompt, encoding="utf-8")


def zip_stage(stage: Path, zip_path: Path):
    if zip_path.exists(): zip_path.unlink()
    with zipfile.ZipFile(zip_path,"w",compression=zipfile.ZIP_DEFLATED,compresslevel=9) as z:
        for p in sorted(stage.rglob("*")):
            if p.is_file():
                comp=zipfile.ZIP_STORED if p.suffix.lower()==".png" else zipfile.ZIP_DEFLATED
                z.write(
                    p,
                    p.relative_to(stage).as_posix(),
                    compress_type=comp,
                    compresslevel=None if comp == zipfile.ZIP_STORED else 9,
                )
    return sha256_file(zip_path)


def export_one_upload_package(
    workspace_dir: str,
    brain_name: str,
    progress: ProgressCallback = None,
    *,
    package_use_mode: str = "CANONICAL_FLASHABLE",
):
    started=time.time()
    normalized_use_mode = str(package_use_mode or "CANONICAL_FLASHABLE").strip().upper()
    if normalized_use_mode not in {"CANONICAL_FLASHABLE", "READ_ONLY_STRESS_RESULT"}:
        raise RuntimeError(f"PACKAGE_USE_MODE_INVALID:{normalized_use_mode}")
    root=brain_output_dir(workspace_dir, brain_name)
    project=root/"project"
    if not (project/"project_router.sqlite").exists(): raise RuntimeError("BRAIN_NOT_BUILT_YET")
    topology=project/"topology"
    required_topology=[
        topology/"project_master_topology.mmd",
        topology/"project_master_topology.svg",
        topology/"project_master_topology.png",
    ]
    missing=[str(path) for path in required_topology if not path.is_file() or path.stat().st_size == 0]
    if missing:
        raise RuntimeError("CHATGPT_TOPOLOGY_RENDER_INCOMPLETE: " + ";".join(missing))
    packages=root/"packages"; packages.mkdir(parents=True, exist_ok=True)
    retired_provider_outputs = retire_provider_outputs(packages, "CHATGPT")
    clear_provider_skip_marker(packages, "CHATGPT")
    live_router_authority = reconcile_project_router_authority(root)
    slug=slugify_name(brain_name)
    variant_suffix = (
        "_read_only_stress_result"
        if normalized_use_mode == "READ_ONLY_STRESS_RESULT"
        else ""
    )
    stage=packages/f"{slug}_one_upload_package{variant_suffix}_v001"
    if stage.exists(): shutil.rmtree(stage)
    stage.mkdir(parents=True)
    if (root/".uepc_env").is_file():
        emit(progress,"export","copying live Env15 governed brain state",str(stage),10,0,1,started)
        for name in (".uepc_env", ".uepc_profile", ".uepc_project", "FLASH_ME_FIRST_SINGLE_PROMPT.txt", "README_NEXT_PROMPT.txt"):
            source_path=root/name
            if source_path.is_file():
                shutil.copy2(source_path,stage/name)
        for name in ("env","uop","project","manifests","recovery","receipts"):
            source_path=root/name
            if source_path.is_dir():
                copy_tree_safe(source_path,stage/name)
    else:
        emit(progress,"export","copying governed Env/UOP authority",str(stage),10,0,1,started)
        copy_locked_env(stage)
        emit(progress,"export","overlaying generated project sector DBs safely",str(project),45,0,1,started)
        copy_tree_safe(project, stage/"project")
    no_third_locked_container = prune_provider_project_lock_artifacts(stage)
    (stage/"manifests").mkdir(exist_ok=True)
    (stage/"receipts").mkdir(exist_ok=True)
    write_flash_prompt(stage, brain_name)
    state_travel_patch = patch_package_folder(stage)
    provider_projection = compact_chatgpt_package_stage(stage)
    provider_readability = materialize_provider_readable_tree(
        root,
        stage / "provider_readable",
    )
    file_count=sum(1 for p in stage.rglob("*") if p.is_file())
    (stage/"receipts"/"export_receipt.md").write_text(
        f"# Export Receipt\n\nstatus=PASS\ncreated={now()}\nfiles={file_count}\n"
        f"provider_size_contract_bytes={CHATGPT_PACKAGE_MAX_BYTES}\n"
        f"provider_projection={provider_projection['contract']}\n",
        encoding="utf-8",
    )
    authority_reseal = reseal_chatgpt_package_authority(
        stage,
        brain_name=brain_name,
        package_use_mode=normalized_use_mode,
    )
    project_package_parity = validate_active_project_package_parity(root, stage)
    if project_package_parity.get("status") != "PASS":
        raise RuntimeError(
            "ACTIVE_PROJECT_PACKAGE_PARITY_FAILED:"
            + ";".join(project_package_parity.get("errors") or [])
        )
    forbidden_project_lock_references = find_forbidden_provider_project_lock_references(stage)
    if forbidden_project_lock_references:
        raise RuntimeError(
            "PROVIDER_PROJECT_LOCK_REFERENCE_FORBIDDEN:"
            + ";".join(forbidden_project_lock_references)
        )
    zip_name_suffix = (
        "_Read_Only_Stress_Result"
        if normalized_use_mode == "READ_ONLY_STRESS_RESULT"
        else ""
    )
    zip_path=packages/f"ChatGPT_LocalAI_{safe_export_name(brain_name)}{zip_name_suffix}_Sqlite_brain.zip"
    digest=zip_stage(stage,zip_path)
    observed_package_bytes = zip_path.stat().st_size
    if observed_package_bytes > CHATGPT_PACKAGE_MAX_BYTES:
        retire_provider_outputs(packages, "CHATGPT")
        retire_provider_outputs(packages, "GEMINI")
        downstream_skip = write_provider_skip_marker(
            packages,
            "GEMINI",
            observed_bytes=None,
            maximum_bytes=GEMINI_PACKAGE_MAX_BYTES,
            reason="GEMINI_SKIPPED_BECAUSE_CHATGPT_512000000_BYTE_PARENT_LIMIT_WAS_EXCEEDED",
        )
        skipped = write_provider_skip_marker(
            packages,
            "CHATGPT",
            observed_bytes=observed_package_bytes,
            maximum_bytes=CHATGPT_PACKAGE_MAX_BYTES,
            reason="CHATGPT_ARCHIVE_EXCEEDS_EXACT_512000000_BYTE_LIMIT",
        )
        return {
            **skipped,
            "package_folder": "",
            "package_zip": "",
            "sha256": "",
            "package_byte_size": observed_package_bytes,
            "package_use_mode": normalized_use_mode,
            "provider_projection": provider_projection,
            "provider_readability": provider_readability,
            "project_package_parity": project_package_parity,
            "live_router_authority": live_router_authority,
            "public_model_state_travel_patch": state_travel_patch,
            "skip_downstream_gemini": True,
            "downstream_gemini_skip": downstream_skip,
            "retired_provider_outputs": retired_provider_outputs,
            "staging_removed": True,
            "staging_path": str(stage),
        }
    package_byte_size=enforce_provider_package_size(
        zip_path,
        CHATGPT_PACKAGE_MAX_BYTES,
        "ChatGPT_LocalAI",
    )
    validation=validate_chatgpt_package(zip_path)
    if validation["status"] != "PASS":
        raise RuntimeError("CHATGPT_PACKAGE_VALIDATION_FAILED: " + ";".join(validation["errors"]))
    remove_provider_stage(stage)
    emit(progress,"done","one-upload package ready",str(zip_path),100,1,1,started)
    return {
        "package_folder":"",
        "package_zip":str(zip_path),
        "sha256":digest,
        "package_byte_size":package_byte_size,
        "provider_size_limit_bytes":CHATGPT_PACKAGE_MAX_BYTES,
        "package_use_mode": normalized_use_mode,
        "provider_projection":provider_projection,
        "provider_readability":provider_readability,
        "project_package_parity":project_package_parity,
        "authority_reseal": authority_reseal,
        "no_third_locked_container": no_third_locked_container,
        "live_router_authority": live_router_authority,
        "validation":validation,
        "public_model_state_travel_patch":state_travel_patch,
        "retired_provider_outputs":retired_provider_outputs,
        "staging_removed": True,
        "staging_path": str(stage),
    }


def export_gemini_exact10(
    workspace_dir: str,
    brain_name: str,
    progress: ProgressCallback = None,
    *,
    chatgpt_package: str | Path | None = None,
    package_use_mode: str = "CANONICAL_FLASHABLE",
):
    root=brain_output_dir(workspace_dir, brain_name)
    started=time.time()
    normalized_use_mode = str(package_use_mode or "CANONICAL_FLASHABLE").strip().upper()
    if normalized_use_mode == "READ_ONLY_STRESS_RESULT":
        provider_prompt = (
            "T023 READ-ONLY STRESS RESULT\n\n"
            "This exact-10 package is evaluation evidence only. Do not flash, mutate, "
            "relock, promote, or treat it as canonical. Query the populated Project payload "
            "read-only and report observations.\n"
        )
    elif normalized_use_mode == "CANONICAL_FLASHABLE":
        provider_prompt = LOCKED_FLASH_PROMPT.rstrip() + "\n\n" + GEMINI_FLASH_APPEND.strip() + "\n"
    else:
        raise RuntimeError(f"PACKAGE_USE_MODE_INVALID:{normalized_use_mode}")
    result=export_gemini_provider_readable(
        root,
        brain_name,
        provider_prompt,
        chatgpt_package=chatgpt_package,
        package_use_mode=normalized_use_mode,
        maximum_bytes=GEMINI_PACKAGE_MAX_BYTES,
    )
    result["package_use_mode"] = normalized_use_mode
    package_path = result.get("gemini_package_zip") or result.get("skip_marker") or ""
    emit(progress,"done","Gemini provider-readable package validated",package_path,100,1,1,started)
    return result


def zip_folder(folder: Path, zip_path: Path):
    with zipfile.ZipFile(zip_path,"w",zipfile.ZIP_DEFLATED) as z:
        if folder.exists():
            for p in sorted(folder.rglob("*")):
                if p.is_file(): z.write(p,p.relative_to(folder).as_posix())


def make_conjoined_db(project: Path, out: Path):
    if out.exists(): out.unlink()
    dest=sqlite3.connect(out)
    dest.execute("CREATE TABLE manifest(source_db TEXT, table_name TEXT, row_count INTEGER)")
    dbs=[project/"project_router.sqlite"]+list((project/"sectors").rglob("*.sqlite"))
    for db in dbs:
        if not db.exists(): continue
        src=sqlite3.connect(db); src.row_factory=sqlite3.Row
        prefix=re.sub(r"[^A-Za-z0-9_]+","_", db.relative_to(project).as_posix()).strip("_").replace("_sqlite","")
        for (t,) in src.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE '%_idx' AND name NOT LIKE '%_data' AND name NOT LIKE '%_config' AND name NOT LIKE '%_docsize' AND name NOT LIKE '%_content'"):
            try:
                rows_=src.execute(f'SELECT * FROM "{t}"').fetchall()
                cols=[d[0] for d in src.execute(f'SELECT * FROM "{t}" LIMIT 1').description or []]
                nt=f"{prefix}__{t}"[:120]
                dest.execute(f'CREATE TABLE IF NOT EXISTS "{nt}"({",".join(["\""+c.replace("\"","_")+"\" TEXT" for c in cols])})')
                for r in rows_:
                    dest.execute(f'INSERT INTO "{nt}" VALUES({",".join(["?" for _ in cols])})', [str(r[c]) if r[c] is not None else None for c in cols])
                dest.execute("INSERT INTO manifest VALUES(?,?,?)", (str(db.relative_to(project)), t, len(rows_)))
            except Exception:
                pass
        src.close()
    dest.commit(); dest.close()


def scan_tools(progress: ProgressCallback=None):
    tools = {
        "git": shutil.which("git"),
        "mmdc": shutil.which("mmdc"),
        "mmdc.cmd": shutil.which("mmdc.cmd"),
        "mmdc.exe": shutil.which("mmdc.exe"),
        "tesseract": shutil.which("tesseract"),
        "python_executable": sys.executable,
        "python_version": sys.version.split()[0],
        "node": shutil.which("node"),
        "npm": shutil.which("npm") or shutil.which("npm.cmd"),
        "pyinstaller": shutil.which("pyinstaller"),
        "inno_iscc": shutil.which("ISCC") or shutil.which("ISCC.exe"),
    }
    py = {}
    for mod in ["openpyxl", "docx", "pptx", "pypdf", "fitz", "PIL", "pytesseract", "sqlite3", "zipfile", "tkinter"]:
        try:
            __import__(mod)
            py[mod] = "OK"
        except Exception as exc:
            py[mod] = "MISSING: " + str(exc)[:120]
    return {"tools":tools, "python_modules":py, "created_at":now(), "install_policy":"same-process hidden pip only; no app relaunch"}


def install_missing_dependencies(progress: ProgressCallback=None):
    mods=["openpyxl","python-docx","python-pptx","pypdf","PyMuPDF","Pillow","pytesseract"]
    run_hidden([sys.executable,"-m","pip","install",*mods], check=False)
    return scan_tools(progress)
