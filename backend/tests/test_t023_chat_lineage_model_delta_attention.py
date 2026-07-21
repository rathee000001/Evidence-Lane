from __future__ import annotations

import hashlib
import json
import sqlite3
import stat
from pathlib import Path

from docx import Document

from sqlite_brain_builder import ipc_worker
from sqlite_brain_builder.brain_versions import capture_brain_version
from sqlite_brain_builder.runtime.env15_chat_lineage import append_env15_chat_lineage_source
from sqlite_brain_builder.runtime.env15_project_schema import resolve_env15_sector
from sqlite_brain_builder.runtime.env15_resource import install_env15_resource
from sqlite_brain_builder.runtime.stable_runtime_v53 import build_brain
from sqlite_brain_builder.workspace.model_connectors import upsert_model_connector
from sqlite_brain_builder.workspace.model_execution import execute_model_connector


def _docx_fixture(path: Path) -> None:
    document = Document()
    document.add_heading("Deterministic Heading", level=1)
    document.add_paragraph("First item", style="List Bullet")
    document.add_paragraph("Second item", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Key"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Alpha"
    table.cell(1, 1).text = "One"
    document.save(path)


def _verified_brain(workspace: Path, brain_name: str) -> None:
    source = workspace.parent / f"{brain_name}.md"
    source.write_text("# Verified baseline\n", encoding="utf-8")
    build_brain(
        str(workspace),
        brain_name,
        [{"source_id": f"source-{brain_name}", "lane_key": "docs", "path": str(source), "active": True}],
        generate_mmd=False,
    )
    capture_brain_version(
        workspace,
        brain_name,
        actor_type="app",
        actor_name="selected-brain contract test",
        reason="verified model-delta baseline",
    )


def _profile(*, output: bool) -> dict:
    response_mapping = {
        "text": "choices.0.message.content",
        "provider_response_id": "id",
    }
    if output:
        response_mapping["output"] = "choices.0.message.output"
    return {
        "endpoint_id": "bounded-model",
        "display_name": "Bounded Model",
        "endpoint_type": "OPENAI_COMPATIBLE",
        "base_url": "https://models.example.test",
        "model_id": "model-1",
        "health_check_path": "/health",
        "request_path": "/v1/chat/completions",
        "model_list_path": "/v1/models",
        "authentication_type": "NONE",
        "authentication_reference": None,
        "streaming_supported": False,
        "tool_call_supported": False,
        "file_supported": False,
        "context_limit": 32768,
        "request_mapping": {"model": "model", "messages": "messages"},
        "response_mapping": response_mapping,
        "usage_mapping": {
            "input_tokens": "usage.prompt_tokens",
            "output_tokens": "usage.completion_tokens",
            "total_tokens": "usage.total_tokens",
        },
        "timeout": 45,
        "local_or_remote": "REMOTE",
        "privacy_classification": "BOUNDED_REMOTE",
        "enabled": True,
        "active_prompt_bar": True,
    }


def _select(workspace: Path, brain_name: str) -> dict:
    return ipc_worker.handle(
        {
            "command": "brain.select",
            "payload": {"workspace_dir": str(workspace), "brain_name": brain_name},
        }
    )


def test_docx_chat_lineage_is_deterministic_markdown_before_sqlite_with_original_provenance(
    tmp_path: Path,
) -> None:
    brain_root = tmp_path / "brain"
    install_env15_resource(brain_root)
    source = tmp_path / "lineage.docx"
    _docx_fixture(source)
    original = source.read_bytes()
    original_hash = hashlib.sha256(original).hexdigest()
    payload = {
        "source_id": "source_chat_docx",
        "lane_key": "chat_lineage",
        "path": str(source),
        "assistant_response": "Imported response",
        "active": True,
    }

    first = append_env15_chat_lineage_source(brain_root, payload)
    _, database = resolve_env15_sector(brain_root, "chat_lineage")
    with sqlite3.connect(database) as connection:
        prompt = connection.execute(
            "SELECT content,content_sha256,storage_status FROM prompt_raw_exact WHERE turn_id=?",
            (first["turn_id"],),
        ).fetchone()
        link = connection.execute(
            "SELECT stable_file_id,external_display_uri,mime_type,size_bytes,sha256,hash_status "
            "FROM file_link_registry WHERE turn_id=?",
            (first["turn_id"],),
        ).fetchone()
        normalization = connection.execute(
            "SELECT source_id,source_format,normalized_format,original_sha256,normalized_sha256,"
            "parser_id,original_size_bytes,normalized_size_bytes,receipt_json "
            "FROM source_normalization_receipt WHERE turn_id=?",
            (first["turn_id"],),
        ).fetchone()
        counts_before = {
            table: connection.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0]
            for table in (
                "turn_prepare",
                "prompt_raw_exact",
                "file_link_registry",
                "source_normalization_receipt",
                "turn_fts",
            )
        }

    markdown = prompt[0]
    assert markdown == (
        "# Deterministic Heading\n\n"
        "- First item\n\n"
        "- Second item\n\n"
        "| Key | Value |\n"
        "| --- | --- |\n"
        "| Alpha | One |\n"
    )
    markdown_hash = hashlib.sha256(markdown.encode("utf-8")).hexdigest()
    assert prompt[1:] == (markdown_hash, "FULL_EXACT")
    assert link == (
        "source_chat_docx",
        str(source.resolve()),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        len(original),
        original_hash,
        "VERIFIED",
    )
    assert normalization[:8] == (
        "source_chat_docx",
        "DOCX",
        "MARKDOWN",
        original_hash,
        markdown_hash,
        "docx_structural_v1_to_deterministic_markdown_v1",
        len(original),
        len(markdown.encode("utf-8")),
    )
    normalization_receipt = json.loads(normalization[8])
    assert normalization_receipt["original_path"] == str(source.resolve())
    assert normalization_receipt["original_sha256"] == original_hash
    assert normalization_receipt["normalized_sha256"] == markdown_hash
    assert first["normalization"]["normalized_sha256"] == markdown_hash
    persisted_receipt = json.loads(Path(first["receipt_path"]).read_text(encoding="utf-8"))
    assert persisted_receipt["normalization"] == first["normalization"]

    replay = append_env15_chat_lineage_source(brain_root, payload)
    with sqlite3.connect(database) as connection:
        counts_after = {
            table: connection.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0]
            for table in counts_before
        }
    assert replay["status"] == "SKIPPED_UNCHANGED"
    assert replay["turn_id"] == first["turn_id"]
    assert replay["normalization"] == first["normalization"]
    assert counts_after == counts_before


def test_governed_model_execution_indexes_immutable_output_without_arming_project_refresh(
    tmp_path: Path,
) -> None:
    workspace = tmp_path / "workspace"
    _verified_brain(workspace, "Delta Brain")
    _verified_brain(workspace, "Other Brain")
    upsert_model_connector(workspace, _profile(output=True))
    content = "# Proposed patch output\n\n- bounded change\n"
    digest = hashlib.sha256(content.encode("utf-8")).hexdigest()

    def transport(*_args) -> dict:
        return {
            "id": "provider-response-delta",
            "choices": [
                {
                    "message": {
                        "content": "Candidate response with governed output.",
                        "output": {
                            "content": content,
                            "sha256": digest,
                            "media_type": "text/markdown",
                            "target_path": "src/example.py",
                            "language": "python",
                            "patch_type": "PROPOSED_PATCH",
                        },
                    }
                }
            ],
            "usage": {"prompt_tokens": 7, "completion_tokens": 4, "total_tokens": 11},
        }

    result = execute_model_connector(
        workspace,
        "Delta Brain",
        "bounded-model",
        "Produce the bounded governed delta.",
        remote_consent=True,
        transport=transport,
    )

    assert result["status"] == "CANDIDATE_RESPONSE"
    assert result["output_id"] == f"model_output_{digest[:24]}"
    assert result["output_sha256"] == digest
    output_path = Path(result["output_path"])
    assert output_path.read_text(encoding="utf-8") == content
    assert output_path.stat().st_mode & stat.S_IWUSR == 0
    assert result["output_application_status"] == "PROPOSED_NOT_APPLIED"
    assert result["candidate_fused"] is False
    output_context = _select(workspace, "Delta Brain")
    other_context = _select(workspace, "Other Brain")
    assert output_context["model_output_state"]["latest_output_id"] == result["output_id"]
    assert output_context["model_output_state"]["project_truth_authority"] is False
    assert output_context["refresh_attention"]["self_active"] is False
    assert output_context["refresh_attention"]["reason_ids"] == []
    assert output_context["refresh_attention"]["model_outputs_are_project_change_truth"] is False
    assert output_context["refresh_attention"]["automatic_refresh"] is False
    assert output_context["refresh_attention"]["automatic_fuse"] is False
    assert other_context["model_output_state"]["output_count"] == 0
    assert other_context["refresh_attention"]["self_active"] is False


def test_response_without_governed_delta_mapping_never_becomes_a_delta(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    _verified_brain(workspace, "Response Brain")
    upsert_model_connector(workspace, _profile(output=False))

    result = execute_model_connector(
        workspace,
        "Response Brain",
        "bounded-model",
        "Return only a response candidate.",
        remote_consent=True,
        transport=lambda *_args: {
            "id": "response-only",
            "choices": [{"message": {"content": "Response only"}}],
        },
    )
    context = _select(workspace, "Response Brain")

    assert "output_id" not in result
    assert context["model_output_state"]["output_count"] == 0
    assert context["refresh_attention"]["self_active"] is False
    assert context["refresh_candidate"]["automatic_refresh"] is False
    assert context["refresh_candidate"]["automatic_fuse"] is False
