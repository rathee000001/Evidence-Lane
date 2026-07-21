from __future__ import annotations

import csv
import hashlib
import json
import re
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET


class OfficeIngestError(RuntimeError):
    pass


def stable_record_id(kind: str, *parts: object) -> str:
    material = "\x1f".join([kind, *(str(part) for part in parts)])
    return f"{kind}_{hashlib.sha256(material.encode('utf-8')).hexdigest()[:24]}"


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _xml_text(path: Path, member: str) -> str:
    try:
        with zipfile.ZipFile(path) as archive:
            root = ET.fromstring(archive.read(member))
        return "\n".join(text.strip() for text in root.itertext() if text and text.strip())
    except Exception:
        return ""


def parse_docx(path: str | Path) -> dict[str, Any]:
    source = Path(path)
    if source.suffix.casefold() != ".docx":
        raise OfficeIngestError(f"DOCX_EXTENSION_REQUIRED:{source}")
    from docx import Document

    document = Document(source)
    properties = document.core_properties
    metadata = {
        "title": properties.title or "",
        "subject": properties.subject or "",
        "author": properties.author or "",
        "keywords": properties.keywords or "",
        "category": properties.category or "",
        "comments": properties.comments or "",
        "created": properties.created.isoformat() if properties.created else None,
        "modified": properties.modified.isoformat() if properties.modified else None,
        "last_modified_by": properties.last_modified_by or "",
        "revision": properties.revision,
    }
    sections = []
    headers_footers = []
    for index, section in enumerate(document.sections, start=1):
        section_id = stable_record_id("section", index)
        sections.append(
            {
                "id": section_id,
                "order": index,
                "start_type": str(section.start_type),
                "orientation": str(section.orientation),
                "page_width": int(section.page_width or 0),
                "page_height": int(section.page_height or 0),
                "top_margin": int(section.top_margin or 0),
                "bottom_margin": int(section.bottom_margin or 0),
                "left_margin": int(section.left_margin or 0),
                "right_margin": int(section.right_margin or 0),
            }
        )
        for kind, container in (("header", section.header), ("footer", section.footer)):
            text = "\n".join(paragraph.text for paragraph in container.paragraphs if paragraph.text.strip())
            if text:
                headers_footers.append(
                    {
                        "id": stable_record_id(kind, index, text),
                        "section_id": section_id,
                        "kind": kind,
                        "text": text,
                    }
                )

    relationship_links = {}
    embedded_images = []
    for rel_id, relationship in sorted(document.part.rels.items()):
        reltype = str(relationship.reltype)
        if reltype.endswith("/hyperlink"):
            relationship_links[rel_id] = str(relationship.target_ref)
        if reltype.endswith("/image"):
            blob = relationship.target_part.blob
            embedded_images.append(
                {
                    "id": stable_record_id("doc_image", rel_id, _sha256(blob)),
                    "relationship_id": rel_id,
                    "target": str(relationship.target_ref),
                    "content_type": getattr(relationship.target_part, "content_type", ""),
                    "byte_size": len(blob),
                    "sha256": _sha256(blob),
                }
            )

    paragraphs = []
    headings = []
    hyperlinks = []
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    for order, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text or ""
        style = paragraph.style.name if paragraph.style is not None else ""
        heading_match = re.match(r"Heading\s+(\d+)", style, re.IGNORECASE)
        list_kind = ""
        paragraph_properties = paragraph._p.pPr
        if paragraph_properties is not None and paragraph_properties.numPr is not None:
            list_kind = "numbered_or_bulleted"
        elif style.casefold().startswith("list"):
            list_kind = style
        paragraph_id = stable_record_id("paragraph", order, style, text)
        record = {
            "id": paragraph_id,
            "order": order,
            "text": text,
            "style": style,
            "is_list": bool(list_kind),
            "list_kind": list_kind,
            "heading_level": int(heading_match.group(1)) if heading_match else None,
        }
        paragraphs.append(record)
        if heading_match:
            headings.append(
                {
                    "id": stable_record_id("heading", paragraph_id),
                    "paragraph_id": paragraph_id,
                    "order": order,
                    "level": int(heading_match.group(1)),
                    "text": text,
                }
            )
        for hyperlink in paragraph._p.findall(".//w:hyperlink", namespace):
            rel_id = hyperlink.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            link_text = "".join(node.text or "" for node in hyperlink.findall(".//w:t", namespace))
            hyperlinks.append(
                {
                    "id": stable_record_id("hyperlink", paragraph_id, rel_id, link_text),
                    "paragraph_id": paragraph_id,
                    "relationship_id": rel_id or "",
                    "text": link_text,
                    "target": relationship_links.get(rel_id or "", ""),
                }
            )

    tables = []
    for table_order, table in enumerate(document.tables, start=1):
        table_id = stable_record_id("doc_table", table_order)
        cell_groups: dict[int, list[str]] = {}
        rows = []
        for row_index, row in enumerate(table.rows, start=1):
            cells = []
            for column_index, cell in enumerate(row.cells, start=1):
                address = f"R{row_index}C{column_index}"
                cell_id = stable_record_id("doc_cell", table_id, address)
                cell_groups.setdefault(id(cell._tc), []).append(address)
                cells.append(
                    {
                        "id": cell_id,
                        "address": address,
                        "row": row_index,
                        "column": column_index,
                        "text": cell.text,
                    }
                )
            rows.append({"row": row_index, "cells": cells})
        merged = [addresses for addresses in cell_groups.values() if len(addresses) > 1]
        tables.append(
            {
                "id": table_id,
                "order": table_order,
                "style": table.style.name if table.style is not None else "",
                "row_count": len(table.rows),
                "column_count": len(table.columns),
                "merged_cell_groups": merged,
                "rows": rows,
            }
        )

    comments = _xml_text(source, "word/comments.xml")
    footnotes = _xml_text(source, "word/footnotes.xml")
    chunks = []
    current_heading = ""
    block: list[dict[str, Any]] = []
    for paragraph in paragraphs:
        if paragraph["heading_level"] is not None:
            current_heading = paragraph["text"]
        if paragraph["text"].strip():
            block.append(paragraph)
        if len(block) >= 20:
            text = "\n".join(item["text"] for item in block)
            chunks.append(
                {
                    "id": stable_record_id("doc_chunk", block[0]["id"], block[-1]["id"], _sha256(text.encode())),
                    "heading": current_heading,
                    "start_structure_id": block[0]["id"],
                    "end_structure_id": block[-1]["id"],
                    "text": text,
                }
            )
            block = []
    if block:
        text = "\n".join(item["text"] for item in block)
        chunks.append(
            {
                "id": stable_record_id("doc_chunk", block[0]["id"], block[-1]["id"], _sha256(text.encode())),
                "heading": current_heading,
                "start_structure_id": block[0]["id"],
                "end_structure_id": block[-1]["id"],
                "text": text,
            }
        )
    return {
        "parser": "docx_structural_v1",
        "metadata": metadata,
        "sections": sections,
        "headers_footers": headers_footers,
        "paragraphs": paragraphs,
        "headings": headings,
        "hyperlinks": hyperlinks,
        "tables": tables,
        "comments_text": comments,
        "footnotes_text": footnotes,
        "embedded_images": embedded_images,
        "chunks": chunks,
    }


_FORMULA_REFERENCE = re.compile(
    r"(?:(?:'([^']+)'|([A-Za-z0-9_ ]+))!)?\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?"
)


def _formula_dependencies(formula: str, current_sheet: str) -> list[dict[str, str]]:
    dependencies = []
    for match in _FORMULA_REFERENCE.finditer(formula or ""):
        sheet = (match.group(1) or match.group(2) or current_sheet).strip()
        start = f"{match.group(3)}{match.group(4)}"
        end = f"{match.group(5)}{match.group(6)}" if match.group(5) else start
        dependencies.append({"sheet": sheet, "range": start if start == end else f"{start}:{end}"})
    return dependencies


def parse_workbook(path: str | Path) -> dict[str, Any]:
    source = Path(path)
    if source.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise OfficeIngestError(f"XLSX_OR_XLSM_REQUIRED:{source}")
    import openpyxl

    workbook = openpyxl.load_workbook(source, data_only=False, read_only=False, keep_vba=source.suffix.casefold() == ".xlsm")
    try:
        properties = workbook.properties
        sheets = []
        formulas = []
        dependencies = []
        tables = []
        charts = []
        ranges = []
        merged_cells = []
        hidden_dimensions = []
        validations = []
        chunks = []
        for sheet_order, worksheet in enumerate(workbook.worksheets, start=1):
            used_range = worksheet.calculate_dimension()
            sheet_id = stable_record_id("sheet", sheet_order, worksheet.title)
            sheet_record = {
                "id": sheet_id,
                "order": sheet_order,
                "title": worksheet.title,
                "state": worksheet.sheet_state,
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "used_range": used_range,
            }
            sheets.append(sheet_record)
            ranges.append({"id": stable_record_id("used_range", sheet_id, used_range), "sheet": worksheet.title, "kind": "used", "range": used_range})
            for merged in sorted(str(value) for value in worksheet.merged_cells.ranges):
                merged_cells.append({"id": stable_record_id("merged", sheet_id, merged), "sheet": worksheet.title, "range": merged})
            for index, dimension in worksheet.row_dimensions.items():
                if dimension.hidden:
                    hidden_dimensions.append({"sheet": worksheet.title, "kind": "row", "index": str(index)})
            for index, dimension in worksheet.column_dimensions.items():
                if dimension.hidden:
                    hidden_dimensions.append({"sheet": worksheet.title, "kind": "column", "index": str(index)})
            for data_validation in worksheet.data_validations.dataValidation:
                validations.append(
                    {
                        "id": stable_record_id("validation", sheet_id, str(data_validation.sqref), data_validation.type),
                        "sheet": worksheet.title,
                        "range": str(data_validation.sqref),
                        "type": data_validation.type,
                        "formula1": data_validation.formula1,
                        "formula2": data_validation.formula2,
                    }
                )
            cells = []
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    value = cell.value
                    cell_record = {
                        "id": stable_record_id("cell", sheet_id, cell.coordinate),
                        "sheet": worksheet.title,
                        "coordinate": cell.coordinate,
                        "data_type": cell.data_type,
                        "value": value.isoformat() if hasattr(value, "isoformat") else value,
                        "number_format": cell.number_format,
                    }
                    cells.append(cell_record)
                    if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
                        formula = str(value)
                        formula_id = stable_record_id("formula", sheet_id, cell.coordinate, formula)
                        formulas.append({"id": formula_id, "sheet": worksheet.title, "coordinate": cell.coordinate, "formula": formula})
                        for dependency in _formula_dependencies(formula, worksheet.title):
                            dependencies.append(
                                {
                                    "id": stable_record_id("formula_edge", formula_id, dependency["sheet"], dependency["range"]),
                                    "formula_id": formula_id,
                                    "from": f"{worksheet.title}!{cell.coordinate}",
                                    "to": f"{dependency['sheet']}!{dependency['range']}",
                                }
                            )
            for table_name in sorted(worksheet.tables):
                table = worksheet.tables[table_name]
                tables.append(
                    {
                        "id": stable_record_id("sheet_table", sheet_id, table.name, table.ref),
                        "sheet": worksheet.title,
                        "name": table.name,
                        "display_name": table.displayName,
                        "range": table.ref,
                        "headers": [column.name for column in table.tableColumns],
                    }
                )
            for chart_order, chart in enumerate(worksheet._charts, start=1):
                series_ranges = []
                for series in getattr(chart, "ser", ()):
                    for attribute in ("val", "cat", "xVal", "yVal"):
                        node = getattr(series, attribute, None)
                        for reference_name in ("numRef", "strRef"):
                            reference = getattr(node, reference_name, None) if node is not None else None
                            formula = getattr(reference, "f", None) if reference is not None else None
                            if formula:
                                series_ranges.append(str(formula))
                charts.append(
                    {
                        "id": stable_record_id("chart", sheet_id, chart_order, type(chart).__name__),
                        "sheet": worksheet.title,
                        "order": chart_order,
                        "chart_type": type(chart).__name__,
                        "anchor": str(getattr(chart, "anchor", "")),
                        "source_ranges": sorted(set(series_ranges)),
                    }
                )
            for offset in range(0, len(cells), 250):
                group = cells[offset : offset + 250]
                chunks.append(
                    {
                        "id": stable_record_id("data_chunk", sheet_id, offset, *(item["id"] for item in group)),
                        "sheet": worksheet.title,
                        "start": offset + 1,
                        "end": offset + len(group),
                        "cells": group,
                    }
                )
            sheet_record["cells"] = cells

        named_ranges = []
        try:
            definitions = list(workbook.defined_names.values())
        except AttributeError:
            definitions = list(workbook.defined_names.definedName)
        for definition in definitions:
            named_ranges.append(
                {
                    "id": stable_record_id("named_range", definition.name, definition.attr_text),
                    "name": definition.name,
                    "value": definition.attr_text,
                    "hidden": bool(definition.hidden),
                    "local_sheet_id": definition.localSheetId,
                }
            )
        return {
            "parser": "xlsx_structural_v1",
            "workbook": {
                "title": properties.title or "",
                "creator": properties.creator or "",
                "created": properties.created.isoformat() if properties.created else None,
                "modified": properties.modified.isoformat() if properties.modified else None,
                "calculation_mode": getattr(workbook.calculation, "calcMode", None),
            },
            "sheets": sheets,
            "ranges": ranges,
            "named_ranges": named_ranges,
            "merged_cells": merged_cells,
            "hidden_dimensions": hidden_dimensions,
            "validations": validations,
            "formulas": formulas,
            "formula_dependencies": dependencies,
            "tables": tables,
            "charts": charts,
            "chunks": chunks,
        }
    finally:
        workbook.close()


def parse_data_source(path: str | Path) -> dict[str, Any]:
    source = Path(path)
    extension = source.suffix.casefold()
    if extension in {".xlsx", ".xlsm"}:
        return {"source_kind": "workbook", "status": "SUPPORTED", "payload": parse_workbook(source)}
    if extension == ".xls":
        return {"source_kind": "legacy_xls", "status": "UNSUPPORTED_CONVERTER_REQUIRED", "payload": {}}
    if extension in {".csv", ".tsv"}:
        delimiter = "\t" if extension == ".tsv" else ","
        with source.open("r", encoding="utf-8-sig", errors="replace", newline="") as stream:
            rows = list(csv.reader(stream, delimiter=delimiter))
        return {
            "source_kind": "delimited",
            "status": "SUPPORTED",
            "payload": {"delimiter": delimiter, "header": rows[0] if rows else [], "rows": rows[1:]},
        }
    if extension == ".json":
        return {"source_kind": "json", "status": "SUPPORTED", "payload": json.loads(source.read_text(encoding="utf-8-sig"))}
    if extension == ".jsonl":
        rows = [json.loads(line) for line in source.read_text(encoding="utf-8-sig").splitlines() if line.strip()]
        return {"source_kind": "jsonl", "status": "SUPPORTED", "payload": rows}
    if extension == ".parquet":
        try:
            import pyarrow.parquet as parquet
        except ImportError:
            return {"source_kind": "parquet", "status": "UNSUPPORTED_PYARROW_NOT_INSTALLED", "payload": {}}
        table = parquet.read_table(source)
        chunks = []
        for ordinal, batch in enumerate(table.to_batches(max_chunksize=1000), start=1):
            rows = batch.to_pylist()
            chunks.append(
                {
                    "id": stable_record_id("parquet_chunk", ordinal, _sha256(json.dumps(rows, sort_keys=True, default=str).encode())),
                    "ordinal": ordinal,
                    "rows": rows,
                }
            )
        return {
            "source_kind": "parquet",
            "status": "SUPPORTED",
            "payload": {
                "schema": str(table.schema),
                "columns": [{"name": field.name, "type": str(field.type)} for field in table.schema],
                "row_count": table.num_rows,
                "column_count": table.num_columns,
                "chunks": chunks,
            },
        }
    raise OfficeIngestError(f"DATA_SOURCE_TYPE_UNSUPPORTED:{extension or '<none>'}")


def parse_pptx(path: str | Path) -> dict[str, Any]:
    source = Path(path)
    if source.suffix.casefold() != ".pptx":
        raise OfficeIngestError(f"PPTX_EXTENSION_REQUIRED:{source}")
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(source)
    slides = []
    chunks = []
    for slide_order, slide in enumerate(presentation.slides, start=1):
        slide_id = stable_record_id("slide", slide_order, getattr(slide.slide_layout, "name", ""))
        shapes = []
        tables = []
        images = []
        text_blocks = []
        for shape_order, shape in enumerate(slide.shapes, start=1):
            shape_id = stable_record_id("shape", slide_id, shape_order, shape.name)
            shape_type = str(shape.shape_type)
            shape_record = {
                "id": shape_id,
                "order": shape_order,
                "name": shape.name,
                "shape_type": shape_type,
                "left": int(shape.left),
                "top": int(shape.top),
                "width": int(shape.width),
                "height": int(shape.height),
                "rotation": float(shape.rotation or 0),
            }
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text or ""
                shape_record["text"] = text
                if text.strip():
                    text_blocks.append({"id": stable_record_id("ppt_text", shape_id, text), "shape_id": shape_id, "text": text})
            if getattr(shape, "has_table", False):
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                tables.append(
                    {
                        "id": stable_record_id("ppt_table", shape_id),
                        "shape_id": shape_id,
                        "row_count": len(shape.table.rows),
                        "column_count": len(shape.table.columns),
                        "rows": rows,
                    }
                )
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                blob = shape.image.blob
                images.append(
                    {
                        "id": stable_record_id("ppt_image", shape_id, _sha256(blob)),
                        "shape_id": shape_id,
                        "filename": shape.image.filename,
                        "content_type": shape.image.content_type,
                        "byte_size": len(blob),
                        "sha256": _sha256(blob),
                    }
                )
            shapes.append(shape_record)
        notes = ""
        try:
            notes = slide.notes_slide.notes_text_frame.text or ""
        except Exception:
            notes = ""
        relationships = [
            {
                "relationship_id": rel_id,
                "relationship_type": str(relationship.reltype),
                "target": str(relationship.target_ref),
            }
            for rel_id, relationship in sorted(slide.part.rels.items())
        ]
        slide_record = {
            "id": slide_id,
            "order": slide_order,
            "layout": getattr(slide.slide_layout, "name", ""),
            "shapes": shapes,
            "text_blocks": text_blocks,
            "tables": tables,
            "images": images,
            "notes": notes,
            "relationships": relationships,
        }
        slides.append(slide_record)
        chunk_text = "\n".join([*(item["text"] for item in text_blocks), notes]).strip()
        if chunk_text:
            chunks.append(
                {
                    "id": stable_record_id("ppt_chunk", slide_id, _sha256(chunk_text.encode())),
                    "slide_id": slide_id,
                    "text": chunk_text,
                }
            )
    properties = presentation.core_properties
    return {
        "parser": "pptx_structural_v1",
        "metadata": {
            "title": properties.title or "",
            "subject": properties.subject or "",
            "author": properties.author or "",
            "created": properties.created.isoformat() if properties.created else None,
            "modified": properties.modified.isoformat() if properties.modified else None,
            "slide_width": presentation.slide_width,
            "slide_height": presentation.slide_height,
        },
        "slides": slides,
        "chunks": chunks,
    }


__all__ = [
    "OfficeIngestError",
    "parse_data_source",
    "parse_docx",
    "parse_pptx",
    "parse_workbook",
    "stable_record_id",
]
